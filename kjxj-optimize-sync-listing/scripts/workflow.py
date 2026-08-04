#!/usr/bin/env python3
"""Deterministic job, confirmation, and delivery gates for Listing orchestration."""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib.util
import json
import re
import subprocess
import sys
import unicodedata
import zipfile
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


SCHEMA_VERSION = 1
ASIN_RE = re.compile(r"(?<![A-Z0-9])(B0[A-Z0-9]{8})(?![A-Z0-9])", re.IGNORECASE)
PHASES = {
    "prepared", "keywords_ready", "listing_ready", "title_ready", "sync_planned",
    "awaiting_confirmation", "confirmed", "applied", "verified",
    "ppc_ready", "qa_ready", "completed", "blocked",
}
CONFIRM_PHRASES = {"确认执行", "按此执行", "确认并执行", "执行"}
AUDIT_SECTIONS = (
    "diagnostic", "audit", "keyword_priority", "keyword_coverage", "keyword_gaps",
    "before_after", "issues_fixed", "recommendations", "working_well",
    "selling_point_notes", "uploaded_keyword_file_summary", "competitive_comparison",
)
AUDIT_LABELS = {
    "diagnostic": "诊断信息",
    "audit": "评分",
    "keyword_priority": "关键词优先级",
    "keyword_coverage": "关键词覆盖",
    "keyword_gaps": "关键词缺口",
    "before_after": "修改对比",
    "issues_fixed": "已修复问题",
    "recommendations": "建议",
    "working_well": "原Listing中表现较好的部分",
    "selling_point_notes": "核心卖点融入说明",
    "uploaded_keyword_file_summary": "上传关键词文件摘要",
    "competitive_comparison": "竞品对比",
}
REQUIRED_AUDIT_SECTIONS = (
    "audit", "keyword_coverage", "before_after", "issues_fixed", "recommendations", "working_well",
)
SHA256_RE = re.compile(r"[0-9a-f]{64}")
CHINESE_RE = re.compile(r"[\u4e00-\u9fff]")
FILENAME_UNSAFE_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]')


def safe_filename_component(value: object, fallback: str) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii").lower()
    normalized = re.sub(r"[^a-z0-9]+", "-", normalized).strip("-")
    return normalized[:80] or fallback


def safe_chinese_filename_component(value: object) -> str:
    """Preserve a validated Chinese prefix while removing filesystem-unsafe text."""
    normalized = FILENAME_UNSAFE_RE.sub("", str(value or "").strip())
    normalized = re.sub(r"\s+", "-", normalized)
    normalized = re.sub(r"-+", "-", normalized).strip("-.")
    return normalized[:80]


def confirmation_review_filename(job: dict[str, Any], generated_at: dt.datetime | None = None) -> str:
    timestamp = generated_at or dt.datetime.now(dt.timezone.utc)
    keyword = next((item for item in job.get("core_keywords", []) if str(item).strip()), "keyword")
    chinese_name = safe_chinese_filename_component(job.get("core_keyword_chinese_name"))
    if not chinese_name or not CHINESE_RE.search(chinese_name):
        fail("生成审核包失败：core_keyword_chinese_name 必须是包含中文字符的准确核心词中文名。")
    return "-".join((
        chinese_name,
        "confirmation-review",
        safe_filename_component(job.get("primary_asin"), "asin").upper(),
        safe_filename_component(keyword, "keyword"),
        safe_filename_component(job.get("marketplace"), "site").upper(),
        timestamp.strftime("%Y%m%d"),
    )) + ".md"

MARKETPLACE_ALIASES = {
    "US": {"us", "usa", "amazon us", "united states", "美国", "美国站"},
    "UK": {"uk", "amazon uk", "united kingdom", "great britain", "英国", "英国站"},
    "DE": {"de", "amazon de", "germany", "deutschland", "德国", "德国站"},
    "FR": {"fr", "amazon fr", "france", "法国", "法国站"},
    "IT": {"it", "amazon it", "italy", "italia", "意大利", "意大利站"},
    "ES": {"es", "amazon es", "spain", "españa", "西班牙", "西班牙站"},
    "JP": {"jp", "amazon jp", "japan", "日本", "日本站"},
    "CA": {"ca", "amazon ca", "canada", "加拿大", "加拿大站"},
    "AU": {"au", "amazon au", "australia", "澳大利亚", "澳洲", "澳大利亚站"},
    "IN": {"in", "amazon in", "india", "印度", "印度站"},
    "MX": {"mx", "amazon mx", "mexico", "méxico", "墨西哥", "墨西哥站"},
    "BR": {"br", "amazon br", "brazil", "brasil", "巴西", "巴西站"},
    "IE": {"ie", "amazon ie", "ireland", "爱尔兰", "爱尔兰站"},
}
SELLERSPRITE_MARKETPLACES = {
    "US": "美国站", "UK": "英国站", "DE": "德国站", "FR": "法国站",
    "IT": "意大利站", "ES": "西班牙站", "JP": "日本站", "CA": "加拿大站",
    "AU": "澳大利亚站", "IN": "印度站", "MX": "墨西哥站", "BR": "巴西站",
    "IE": "爱尔兰站",
}
MARKETPLACE_LANGUAGES = {
    "US": "en_US", "UK": "en_GB", "DE": "de_DE", "FR": "fr_FR",
    "IT": "it_IT", "ES": "es_ES", "JP": "ja_JP", "CA": "en_CA",
    "AU": "en_AU", "IN": "en_IN", "MX": "es_MX", "BR": "pt_BR",
    "IE": "en_IE",
}
MARKETPLACE_CURRENCIES = {
    "US": "USD", "UK": "GBP", "DE": "EUR", "FR": "EUR",
    "IT": "EUR", "ES": "EUR", "JP": "JPY", "CA": "CAD",
    "AU": "AUD", "IN": "INR", "MX": "MXN", "BR": "BRL",
    "IE": "EUR",
}
DEFAULT_MONTHLY_BUDGET = 600
DEFAULT_BREAK_EVEN_ACOS = 0.40
DEFAULT_SELLING_PRICE = 40
FX_SOURCE_URLS = (
    "https://api.frankfurter.app/latest?from=USD",
    "https://open.er-api.com/v6/latest/USD",
)
PRODUCT_STAGE_ALIASES = {
    "launch": {"launch", "new", "new product", "新品", "新品期", "启动期"},
    "mature": {"mature", "existing", "established", "成熟", "成熟期", "在售成熟商品"},
}
UNKNOWN_VALUES = {"", "unknown", "n/a", "na", "none", "null", "未知", "不知道", "无"}
PPC_COPY_BLOCKS = (
    "manual_exact_keywords", "manual_broad_keywords",
    "auto_negative_exact_keywords", "broad_negative_exact_keywords",
    "negative_phrase_keywords", "product_targeting_asins",
)
PPC_REQUIRED_PLAN_FIELDS = (
    "financial_framework", "keyword_sources", "campaigns", "copy_blocks",
    "budget_summary", "launch_schedule", "optimization_plan_4_weeks", "risk_notes",
    "data_quality", "bid_guidance",
)
PPC_SOURCE_SKILL = "amazon-ppc-campaign"
PPC_SOURCE_SKILL_PATH = "/Users/apple/.agents/skills/amazon-ppc-campaign/SKILL.md"
PPC_INVOCATION_ID_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9._:-]{11,}")
PPC_STRATEGY_KEYS = {
    "campaigns", "copy_blocks", "bids", "bid", "starting_bid", "default_bid",
    "negative_keywords", "negative_phrase_keywords", "manual_exact_keywords",
    "manual_broad_keywords", "keyword_priority", "budget_allocation",
}
PPC_AMOUNT_BID_KEYS = {"bid", "cpc", "max_cpc", "starting_bid", "default_bid", "bid_amount"}
PPC_PLACEHOLDER_RE = re.compile(r"\[(?:keyword|term|asin|product)[^\]]*\]|\b(?:TBD|TODO|XX(?:\.XX)?)\b|待填写", re.IGNORECASE)
KEYWORD_SOURCES = {
    "user_core", "uploaded_file", "sellersprite", "amazon_autocomplete", "competitor_listing",
}
KEYWORD_RELEVANCE = {"high", "medium"}
KEYWORD_FALLBACK_STATUSES = {"complete", "insufficient", "failed", "not_needed"}
KEYWORD_ID_RE = re.compile(r"KW-\d{3,}")
KEYWORD_COUNT_MINIMUM = 10
CHINESE_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
PPC_COPY_BLOCK_LABELS = (
    "Manual Exact Keywords", "Manual Broad Keywords",
    "Auto Negative Exact Keywords", "Broad Negative Exact Keywords",
    "Negative Phrase Keywords", "Product Targeting ASINs",
)
PPC_REQUIRED_MARKDOWN_SECTIONS = (
    "输入数据质量", "财务依据与默认值影响", "关键词来源与排除边界",
    "竞价依据限制", "Campaign 设计理由",
)
COLOR_ALIASES = {
    "grey": {"grey", "gray", "灰", "灰色"},
    "green": {"green", "绿色", "綠色", "绿"},
    "pink": {"pink", "粉色", "粉红", "粉紅色"},
    "white": {"white", "白", "白色"},
    "black": {"black", "黑", "黑色"},
    "blue": {"blue", "蓝", "蓝色", "藍色"},
    "beige": {"beige", "米色"},
    "brown": {"brown", "棕色", "咖啡色"},
    "red": {"red", "红", "红色", "紅色"},
}
COLOR_LOCALIZATION = {
    "US": {"grey": "Gray", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "UK": {"grey": "Grey", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "IE": {"grey": "Grey", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "CA": {"grey": "Grey", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "AU": {"grey": "Grey", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "IN": {"grey": "Grey", "green": "Green", "pink": "Pink", "white": "White", "black": "Black", "blue": "Blue", "beige": "Beige", "brown": "Brown", "red": "Red"},
    "DE": {"grey": "Grau", "green": "Grün", "pink": "Rosa", "white": "Weiß", "black": "Schwarz", "blue": "Blau", "beige": "Beige", "brown": "Braun", "red": "Rot"},
    "FR": {"grey": "Gris", "green": "Vert", "pink": "Rose", "white": "Blanc", "black": "Noir", "blue": "Bleu", "beige": "Beige", "brown": "Marron", "red": "Rouge"},
    "IT": {"grey": "Grigio", "green": "Verde", "pink": "Rosa", "white": "Bianco", "black": "Nero", "blue": "Blu", "beige": "Beige", "brown": "Marrone", "red": "Rosso"},
    "ES": {"grey": "Gris", "green": "Verde", "pink": "Rosa", "white": "Blanco", "black": "Negro", "blue": "Azul", "beige": "Beige", "brown": "Marrón", "red": "Rojo"},
    "MX": {"grey": "Gris", "green": "Verde", "pink": "Rosa", "white": "Blanco", "black": "Negro", "blue": "Azul", "beige": "Beige", "brown": "Café", "red": "Rojo"},
    "BR": {"grey": "Cinza", "green": "Verde", "pink": "Rosa", "white": "Branco", "black": "Preto", "blue": "Azul", "beige": "Bege", "brown": "Marrom", "red": "Vermelho"},
    "JP": {"grey": "グレー", "green": "グリーン", "pink": "ピンク", "white": "ホワイト", "black": "ブラック", "blue": "ブルー", "beige": "ベージュ", "brown": "ブラウン", "red": "レッド"},
}


def fail(message: str) -> None:
    raise ValueError(message)


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"无法读取 JSON：{path}：{exc}")
    if not isinstance(value, dict):
        fail(f"JSON 顶层必须是对象：{path}")
    return value


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_marketplace(value: Any) -> str:
    normalized = str(value or "").strip().casefold()
    for code, aliases in MARKETPLACE_ALIASES.items():
        if normalized == code.casefold() or normalized in {alias.casefold() for alias in aliases}:
            return code
    fail(f"不支持的 Amazon 站点：{value}")


def normalize_asin(value: Any, field: str) -> str:
    text = str(value or "").strip().upper()
    if not ASIN_RE.fullmatch(text):
        fail(f"{field} 不是有效的 10 位 ASIN：{value}")
    return text


def split_text_values(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，、;；\n]+", str(value)) if item.strip()]


def optional_positive_number(value: Any, field: str) -> float | None:
    if value is None or str(value).strip().casefold() in UNKNOWN_VALUES:
        return None
    if isinstance(value, bool):
        fail(f"{field} 必须是正数。")
    text = str(value).strip().replace(",", "")
    text = re.sub(r"^[^\d.+-]+|[^\d.]+$", "", text)
    try:
        number = float(text)
    except ValueError:
        fail(f"{field} 必须是正数：{value}")
    if number <= 0:
        fail(f"{field} 必须大于 0：{value}")
    return number


def optional_rate(value: Any, field: str) -> float | None:
    if value is None or str(value).strip().casefold() in UNKNOWN_VALUES:
        return None
    text = str(value).strip().replace("%", "")
    try:
        rate = float(text)
    except ValueError:
        fail(f"{field} 必须是百分比或 0-1 小数：{value}")
    if "%" in str(value) or rate > 1:
        rate /= 100
    if not 0 < rate <= 1:
        fail(f"{field} 必须在 0% 到 100% 之间：{value}")
    return rate


def normalize_product_stage(value: Any) -> str:
    normalized = str(value or "launch").strip().casefold()
    for stage, aliases in PRODUCT_STAGE_ALIASES.items():
        if normalized in {item.casefold() for item in aliases}:
            return stage
    fail(f"ppc_campaign.product_stage 必须是 launch 或 mature：{value}")


def fetch_fx_json(url: str) -> dict[str, Any]:
    request = Request(url, headers={"User-Agent": "Codex Amazon Listing Workflow/1.0"})
    try:
        with urlopen(request, timeout=10) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, OSError, json.JSONDecodeError) as exc:
        raise ValueError(str(exc)) from exc
    if not isinstance(payload, dict):
        raise ValueError("汇率响应顶层不是对象")
    return payload


def parse_fx_payload(url: str, payload: dict[str, Any], currency: str) -> tuple[float, str]:
    rates = payload.get("rates")
    if not isinstance(rates, dict) or currency not in rates:
        raise ValueError(f"响应缺少 {currency} 汇率")
    try:
        rate = float(rates[currency])
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{currency} 汇率不是有效数字") from exc
    if rate <= 0:
        raise ValueError(f"{currency} 汇率必须大于 0")

    rate_date = str(payload.get("date") or "").strip()
    if not rate_date:
        timestamp = payload.get("time_last_update_unix")
        if timestamp is not None:
            try:
                rate_date = dt.datetime.fromtimestamp(float(timestamp), dt.timezone.utc).date().isoformat()
            except (TypeError, ValueError, OSError):
                rate_date = ""
    if not rate_date:
        updated = str(payload.get("time_last_update_utc") or "").strip()
        if updated:
            try:
                rate_date = parsedate_to_datetime(updated).date().isoformat()
            except (TypeError, ValueError, OverflowError):
                rate_date = ""
    if not rate_date:
        raise ValueError(f"{url} 响应缺少可核验汇率日期")
    return rate, rate_date


def resolve_default_monthly_budget(currency: str, fetcher=None) -> tuple[float | int, dict[str, Any]]:
    if currency in {"EUR", "USD"}:
        return DEFAULT_MONTHLY_BUDGET, {
            "used": True,
            "base_amount": DEFAULT_MONTHLY_BUDGET,
            "base_currency": currency,
            "exchange_rate": 1.0,
            "rate_date": None,
            "source_url": "skill-default",
            "local_amount": DEFAULT_MONTHLY_BUDGET,
        }

    load = fetcher or fetch_fx_json
    errors: list[str] = []
    for url in FX_SOURCE_URLS:
        try:
            rate, rate_date = parse_fx_payload(url, load(url), currency)
            raw_amount = DEFAULT_MONTHLY_BUDGET * rate
            local_amount: float | int = int(round(raw_amount)) if currency == "JPY" else round(raw_amount, 2)
            return local_amount, {
                "used": True,
                "base_amount": DEFAULT_MONTHLY_BUDGET,
                "base_currency": "USD",
                "exchange_rate": rate,
                "rate_date": rate_date,
                "source_url": url,
                "local_amount": local_amount,
            }
        except ValueError as exc:
            errors.append(f"{url}: {exc}")
    fail("汇率数据获取失败，无法按 USD 600 换算默认广告预算：" + "；".join(errors))


def normalize_ppc_campaign(value: Any, marketplace: str) -> dict[str, Any]:
    if value is None:
        source: dict[str, Any] = {}
    elif isinstance(value, dict):
        source = value
    else:
        fail("ppc_campaign 必须是对象。")
    expected_currency = MARKETPLACE_CURRENCIES[marketplace]
    currency = str(source.get("currency") or expected_currency).strip().upper()
    if currency != expected_currency:
        fail(f"{marketplace} 站广告预算币种必须是 {expected_currency}，不能使用 {currency}。")
    monthly_budget = optional_positive_number(source.get("monthly_ad_budget"), "月度广告预算")
    landed_cost = optional_positive_number(source.get("landed_cost"), "单件到岸成本")
    amazon_fees = optional_positive_number(source.get("amazon_fees"), "Amazon 单件费用")
    break_even_acos = optional_rate(source.get("break_even_acos"), "盈亏平衡 ACoS")
    selling_price = optional_positive_number(source.get("selling_price"), "单件售价")
    conversion_rate = optional_rate(source.get("conversion_rate"), "预期转化率")

    defaults_applied: dict[str, Any] = {
        "monthly_ad_budget": {"used": False},
        "break_even_acos": {"used": False},
        "selling_price": {"used": False},
    }
    if monthly_budget is None:
        monthly_budget, defaults_applied["monthly_ad_budget"] = resolve_default_monthly_budget(currency)

    complete_costs = landed_cost is not None and amazon_fees is not None
    partial_costs = (landed_cost is None) != (amazon_fees is None)
    if partial_costs:
        ignored = {}
        if landed_cost is not None:
            ignored["landed_cost"] = landed_cost
        if amazon_fees is not None:
            ignored["amazon_fees"] = amazon_fees
        landed_cost = None
        amazon_fees = None
        defaults_applied["ignored_incomplete_cost_inputs"] = ignored
    if break_even_acos is None and not complete_costs:
        break_even_acos = DEFAULT_BREAK_EVEN_ACOS
        defaults_applied["break_even_acos"] = {
            "used": True,
            "value": DEFAULT_BREAK_EVEN_ACOS,
            "reason": "未提供完整成本数据或显式盈亏平衡 ACoS",
        }
    if selling_price is None:
        selling_price = DEFAULT_SELLING_PRICE
        defaults_applied["selling_price"] = {
            "used": True,
            "value": DEFAULT_SELLING_PRICE,
            "reason": "售价缺失，按流程默认 40",
        }
    return {
        "enabled": True,
        "mode": "build",
        "primary_asin_only": True,
        "currency": currency,
        "monthly_ad_budget": monthly_budget,
        "selling_price": selling_price,
        "landed_cost": landed_cost,
        "amazon_fees": amazon_fees,
        "break_even_acos": break_even_acos,
        "conversion_rate": conversion_rate,
        "product_stage": normalize_product_stage(source.get("product_stage")),
        "missing_user_inputs": [],
        "defaults_applied": defaults_applied,
        "selling_price_requires_discovery": False,
    }


def color_canonical(value: str) -> str | None:
    normalized = value.strip().casefold()
    for canonical, aliases in COLOR_ALIASES.items():
        if normalized in {alias.casefold() for alias in aliases}:
            return canonical
    return None


def normalized_attribute(raw: dict[str, Any], marketplace: str) -> dict[str, str]:
    attr_type = str(raw.get("type", "other")).strip() or "other"
    source = str(raw.get("source_value", raw.get("value", ""))).strip()
    canonical = str(raw.get("canonical_value", "")).strip()
    localized = str(raw.get("marketplace_value", "")).strip()
    if attr_type == "color":
        canonical = canonical or color_canonical(source) or ""
        if canonical and not localized:
            localized = COLOR_LOCALIZATION.get(marketplace, {}).get(canonical, "")
    if attr_type == "size":
        canonical = canonical or source.upper()
        localized = localized or source.upper()
    return {
        "type": attr_type,
        "source_value": source,
        "canonical_value": canonical,
        "marketplace_value": localized,
    }


def parse_variant_string(raw: str, marketplace: str) -> dict[str, Any]:
    match = ASIN_RE.search(raw)
    if not match:
        fail(f"变体信息缺少有效 ASIN：{raw}")
    asin = match.group(1).upper()
    remainder = (raw[:match.start()] + raw[match.end():]).strip(" \t-–—|｜:：")
    notes = [item.strip() for item in re.findall(r"[（(]([^()（）]+)[）)]", remainder) if item.strip()]
    plain = re.sub(r"[（(][^()（）]+[）)]", "", remainder)
    segments = [item.strip() for item in re.split(r"\s*[-–—|｜]\s*", plain) if item.strip()]
    attributes: list[dict[str, str]] = []
    unresolved: list[str] = []
    for segment in segments:
        canonical = color_canonical(segment)
        if canonical:
            attributes.append(normalized_attribute({"type": "color", "source_value": segment, "canonical_value": canonical}, marketplace))
        elif re.fullmatch(r"(?:XXS|XS|S|M|L|XL|XXL|XXXL|\d+XL)", segment, re.IGNORECASE):
            attributes.append(normalized_attribute({"type": "size", "source_value": segment}, marketplace))
        elif re.search(r"\d", segment) and re.search(r"(?:cm|mm|m|in|inch|英寸|厘米|毫米)", segment, re.IGNORECASE):
            attributes.append(normalized_attribute({"type": "dimension", "source_value": segment, "canonical_value": segment, "marketplace_value": segment}, marketplace))
        else:
            unresolved.append(segment)
    for note in notes:
        attributes.append(normalized_attribute({"type": "audience_note", "source_value": note}, marketplace))
        unresolved.append(note)
    expected = [item["marketplace_value"] for item in attributes if item["type"] in {"color", "size", "model"} and item["marketplace_value"]]
    return {"asin": asin, "raw": raw, "attributes": attributes, "expected_title_terms": expected, "unresolved_segments": unresolved}


def normalize_variant(value: Any, marketplace: str) -> dict[str, Any]:
    if isinstance(value, str):
        return parse_variant_string(value, marketplace)
    if not isinstance(value, dict):
        fail(f"变体必须是字符串或对象：{value}")
    raw = str(value.get("raw", "")).strip()
    asin_value = value.get("asin")
    if not asin_value and raw:
        return parse_variant_string(raw, marketplace)
    asin = normalize_asin(asin_value, "变体 ASIN")
    parsed = parse_variant_string(raw, marketplace) if raw and ASIN_RE.search(raw) else {"asin": asin, "raw": raw, "attributes": [], "expected_title_terms": [], "unresolved_segments": []}
    if parsed["asin"] != asin:
        fail(f"变体对象的 asin 与 raw 中 ASIN 不一致：{asin} / {parsed['asin']}")
    attributes = [normalized_attribute(item, marketplace) for item in value.get("attributes", parsed["attributes"])]
    unresolved = split_text_values(value.get("unresolved_segments", parsed["unresolved_segments"]))
    expected = split_text_values(value.get("expected_title_terms"))
    if not expected:
        expected = [item["marketplace_value"] for item in attributes if item["type"] in {"color", "size", "model"} and item["marketplace_value"]]
    return {"asin": asin, "raw": raw, "attributes": attributes, "expected_title_terms": expected, "unresolved_segments": unresolved}


def source_record(path: Path, kind: str) -> dict[str, Any]:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        fail(f"输入文件不存在：{resolved}")
    return {"kind": kind, "path": str(resolved), "size": resolved.stat().st_size, "sha256": sha256(resolved)}


def normalize_job(raw: dict[str, Any], run_dir: Path) -> dict[str, Any]:
    if int(raw.get("schema_version", SCHEMA_VERSION)) != SCHEMA_VERSION:
        fail(f"不支持的 schema_version：{raw.get('schema_version')}")
    marketplace = normalize_marketplace(raw.get("marketplace"))
    primary = normalize_asin(raw.get("primary_asin"), "主 ASIN")
    keywords = split_text_values(raw.get("core_keywords"))
    if not keywords:
        fail("至少需要一个核心关键词。")
    chinese_name = str(raw.get("core_keyword_chinese_name") or "").strip()
    if not chinese_name and CHINESE_RE.search(keywords[0]):
        chinese_name = keywords[0]
    if not chinese_name or not CHINESE_RE.search(chinese_name):
        fail("core_keyword_chinese_name 必须提供首个核心关键词的准确中文产品语义翻译，且必须包含中文字符。")
    competitors = [normalize_asin(item, "竞品 ASIN") for item in split_text_values(raw.get("competitor_asins"))]
    if len(competitors) > 3:
        fail("竞品 ASIN 最多 3 个，请缩减后重试。")
    if len(set(competitors)) != len(competitors):
        fail("竞品 ASIN 存在重复。")
    if primary in competitors:
        fail("主 ASIN 不能同时作为竞品 ASIN。")
    selling_points = split_text_values(raw.get("selling_points"))
    if not selling_points:
        fail("产品核心卖点不能为空。")

    variants = [normalize_variant(item, marketplace) for item in (raw.get("variants") or [])]
    variant_asins = [item["asin"] for item in variants]
    if len(set(variant_asins)) != len(variant_asins):
        fail("变体 ASIN 存在重复。")
    if primary not in variant_asins:
        variants.insert(0, {"asin": primary, "raw": "", "attributes": [], "expected_title_terms": [], "unresolved_segments": []})
    for item in variants:
        item["is_primary"] = item["asin"] == primary

    source = raw.get("keyword_source") or {}
    mode = str(source.get("mode", "auto")).strip().casefold() or "auto"
    if mode not in {"auto", "uploaded", "sellersprite", "combined"}:
        fail("keyword_source.mode 必须是 auto、uploaded、sellersprite 或 combined。")
    keyword_files = [Path(item) for item in source.get("files", [])]
    if mode in {"uploaded", "combined"} and not keyword_files:
        fail("选择上传关键词文件时，至少需要一个 CSV 或 XLSX 文件。")
    if mode in {"auto", "sellersprite"} and keyword_files:
        fail(f"keyword_source.mode={mode} 时不应同时指定关键词文件。")
    for path in keyword_files:
        if path.suffix.lower() not in {".csv", ".xlsx"}:
            fail(f"不支持的关键词文件格式：{path}")

    workbook = Path(str(raw.get("category_workbook", ""))).expanduser()
    if workbook.suffix.lower() != ".xlsm":
        fail("分类商品报告必须是 .xlsm 文件。")
    sources = [source_record(path, "keyword") for path in keyword_files]
    sources.append(source_record(workbook, "category_workbook"))
    workspace = Path(str(raw.get("workspace") or Path.cwd())).expanduser().resolve()
    ppc_campaign = normalize_ppc_campaign(raw.get("ppc_campaign"), marketplace)
    model_review = any(item["unresolved_segments"] or any(not attr["marketplace_value"] for attr in item["attributes"] if attr["type"] != "audience_note") for item in variants)
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    return {
        "schema_version": SCHEMA_VERSION,
        "job_id": run_dir.name,
        "created_at": now,
        "primary_asin": primary,
        "marketplace": marketplace,
        "marketplace_language": MARKETPLACE_LANGUAGES[marketplace],
        "sellersprite_marketplace": SELLERSPRITE_MARKETPLACES[marketplace],
        "core_keywords": keywords,
        "core_keyword_chinese_name": chinese_name,
        "competitor_asins": competitors,
        "selling_points": selling_points,
        "variants": variants,
        "target_asins": [item["asin"] for item in variants],
        "keyword_source": {
            "mode": mode,
            "use_uploaded_files": mode in {"uploaded", "combined"},
            "use_sellersprite": mode in {"sellersprite", "combined"},
            "use_keyword_research_fallback": True,
            "minimum_usable_keywords": KEYWORD_COUNT_MINIMUM,
            "files": [str(path.expanduser().resolve()) for path in keyword_files],
            "relevance_min": 30,
            "monthly_search_volume_min": 1,
        },
        "ppc_campaign": ppc_campaign,
        "category_workbook": str(workbook.resolve()),
        "workspace": str(workspace),
        "tone": "Professional",
        "title_option": 1,
        "model_review_required": model_review,
        "source_files": sources,
        "raw_input": raw.get("raw_input", ""),
    }


def build_legacy_mode_b_request(
    job: dict[str, Any], keyword_pool: dict[str, Any], asin: str,
    invocation_id: str, created_at: str,
) -> dict[str, Any]:
    """Build the only request shape allowed for amazon-listing-optimization Mode B."""
    target_asin = normalize_asin(asin, "传统 Listing ASIN")
    eligible_keywords = [
        {"keyword": str(item.get("keyword", "")).strip(), "metrics": item.get("metrics", {})}
        for item in keyword_pool.get("keywords", [])
        if isinstance(item, dict) and item.get("listing_eligible") is True
    ]
    if not eligible_keywords:
        fail("传统 Listing 调用缺少 listing_eligible 关键词。")
    manifest = {
        "asin_target": {"provided": True, "status": "used", "note": "传入目标 ASIN 供 Mode B 自行获取 Listing。", "data": {"asin": target_asin}},
        "marketplace_language_tone": {"provided": True, "status": "used", "note": "传入站点、语言与语气。", "data": {"marketplace": job["marketplace"], "language": job["marketplace_language"], "tone": job["tone"]}},
        "core_keywords": {"provided": True, "status": "used", "note": "传入用户核心关键词。", "data": list(job["core_keywords"])},
        "listing_eligible_keywords": {"provided": True, "status": "used", "note": "传入已验证关键词及其真实指标。", "data": eligible_keywords},
    }
    return {
        "source_skill": "amazon-listing-optimization",
        "source_skill_path": "/Users/apple/.agents/skills/amazon-listing-optimization/SKILL.md",
        "mode": "B",
        "asin": target_asin,
        "marketplace": job["marketplace"],
        "language": job["marketplace_language"],
        "tone": job["tone"],
        "core_keywords": list(job["core_keywords"]),
        "listing_eligible_keywords": eligible_keywords,
        "invocation_id": invocation_id,
        "created_at": created_at,
        "input_manifest": manifest,
    }


def materialize_legacy_result(args: argparse.Namespace) -> dict[str, Any]:
    """Parse sealed Mode B Markdown through the KJXJ-owned adapter."""
    adapter = Path(__file__).with_name("external_adapters") / "listing_mode_b_adapter.py"
    completed = subprocess.run(
        [sys.executable, str(adapter), "materialize", "--request", args.request,
         "--raw-output", args.raw_output, "--receipt", args.receipt, "--output", args.output],
        text=True, capture_output=True, check=False,
    )
    if completed.returncode:
        fail("Listing Adapter 无法生成传统 Listing 交接结果：" + (completed.stderr.strip() or completed.stdout.strip()))
    return {"ok": True, "result": str(Path(args.output).resolve())}


def prepare(args: argparse.Namespace) -> dict[str, Any]:
    input_path = Path(args.input).expanduser().resolve()
    run_dir = Path(args.run_dir).expanduser().resolve()
    if run_dir.exists() and any(run_dir.iterdir()):
        fail(f"运行目录必须为空或不存在：{run_dir}")
    run_dir.mkdir(parents=True, exist_ok=True)
    job = normalize_job(load_json(input_path), run_dir)
    job_path = run_dir / "job.json"
    status_path = run_dir / "status.json"
    write_json(job_path, job)
    write_json(status_path, {"job_id": job["job_id"], "phase": "prepared", "updated_at": dt.datetime.now(dt.timezone.utc).isoformat(), "artifacts": {"job": str(job_path)}})
    return {
        "ok": True,
        "job": str(job_path),
        "status": str(status_path),
        "model_review_required": job["model_review_required"],
        "target_asins": job["target_asins"],
        "ppc_missing_user_inputs": job["ppc_campaign"]["missing_user_inputs"],
        "ppc_defaults_applied": job["ppc_campaign"]["defaults_applied"],
        "ppc_selling_price_requires_discovery": job["ppc_campaign"]["selling_price_requires_discovery"],
    }


def transition(args: argparse.Namespace) -> dict[str, Any]:
    status_path = Path(args.status).resolve()
    status = load_json(status_path)
    if args.phase not in PHASES:
        fail(f"不支持的运行阶段：{args.phase}")
    current_phase = status.get("phase")
    required_previous = {
        "keywords_ready": "prepared",
        "listing_ready": "keywords_ready",
        "title_ready": "listing_ready",
        "sync_planned": "title_ready",
        "ppc_ready": "sync_planned",
        "qa_ready": "ppc_ready",
        "awaiting_confirmation": "qa_ready",
        "confirmed": "awaiting_confirmation",
        "applied": "confirmed",
        "verified": "applied",
        "completed": "verified",
    }
    if args.phase in required_previous and current_phase != required_previous[args.phase]:
        fail(f"阶段 {args.phase} 必须从 {required_previous[args.phase]} 进入，当前为 {current_phase}。")
    artifacts = status.setdefault("artifacts", {})
    for item in args.artifact:
        if "=" not in item:
            fail(f"artifact 必须使用 key=/absolute/path：{item}")
        key, value = item.split("=", 1)
        artifacts[key] = str(Path(value).expanduser().resolve())
    validation_keys = {
        "keywords_ready": "keyword_validation",
        "listing_ready": "report_validation",
        "title_ready": "title_validation",
        "ppc_ready": "ppc_validation",
        "qa_ready": "rufus_qa_validation",
        "verified": "delivery_validation",
    }
    validation_key = validation_keys.get(args.phase)
    if validation_key:
        validation_path = Path(str(artifacts.get(validation_key, ""))).expanduser()
        if not validation_path.is_file():
            fail(f"阶段 {args.phase} 缺少验证文件 artifact：{validation_key}")
        validation = load_json(validation_path.resolve())
        if validation.get("ok") is not True or validation.get("errors"):
            fail(f"阶段 {args.phase} 的 {validation_key} 未通过。")
        if args.phase == "keywords_ready":
            job_path = Path(str(artifacts.get("job", ""))).expanduser().resolve()
            load_keyword_validation(load_json(job_path), validation_path.resolve())
    status["phase"] = args.phase
    status["updated_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    write_json(status_path, status)
    return status


def artifact_validation(path: Path, required: dict[str, Any] | None = None) -> tuple[bool, str]:
    """Return a user-actionable validation result without changing a run."""
    if not path.is_file():
        return False, f"文件不存在：{path}"
    try:
        value = load_json(path)
    except ValueError as exc:
        return False, str(exc)
    if value.get("ok") is not True or value.get("errors"):
        return False, "验证结果未通过或包含错误。"
    for key, expected in (required or {}).items():
        if value.get(key) != expected:
            return False, f"验证结果字段 {key} 不匹配。"
    return True, "已通过验证。"


def inspect_run(args: argparse.Namespace) -> dict[str, Any]:
    """Inspect recoverable stages. This command is deliberately read-only."""
    run_dir = Path(args.run_dir).expanduser().resolve()
    job_path, status_path = run_dir / "job.json", run_dir / "status.json"
    if not job_path.is_file() or not status_path.is_file():
        fail("运行目录必须包含 job.json 和 status.json。")
    job, status = load_json(job_path), load_json(status_path)
    artifacts = status.get("artifacts") if isinstance(status.get("artifacts"), dict) else {}
    stages = [("keywords_ready", "keyword_validation"), ("listing_ready", "report_validation"),
              ("title_ready", "report_validation"), ("ppc_ready", "ppc_validation"),
              ("qa_ready", "rufus_qa_validation")]
    completed, missing = [], []
    for phase, key in stages:
        candidate = Path(str(artifacts.get(key, run_dir / f"{key.replace('_validation', '')}-validation.json"))).expanduser()
        ok, reason = artifact_validation(candidate)
        if ok:
            completed.append({"phase": phase, "artifact": str(candidate.resolve()), "reason": reason})
        else:
            missing.append({"phase": phase, "artifact": str(candidate), "reason": reason})
    plan = run_dir / "listing-sync-plan.json"
    try:
        plan_ready = plan.is_file() and load_json(plan).get("status") == "ready"
    except ValueError:
        plan_ready = False
    if plan_ready:
        completed.append({"phase": "sync_planned", "artifact": str(plan), "reason": "同步计划为 ready。"})
    else:
        missing.append({"phase": "sync_planned", "artifact": str(plan), "reason": "缺少状态为 ready 的同步计划。"})
    sealed = run_dir / "confirmation-request.json"
    if sealed.is_file():
        completed.append({"phase": "awaiting_confirmation", "artifact": str(sealed), "reason": "存在已封存确认请求。"})
    next_phase = missing[0]["phase"] if missing else ("awaiting_confirmation" if not sealed.is_file() else "confirmed")
    return {"ok": True, "run_dir": str(run_dir), "recorded_phase": status.get("phase"),
            "completed_stages": completed, "missing_stages": missing, "next_phase": next_phase,
            "next_action": "运行 resume-plan --run-dir <运行目录>" if missing else "展示并确认封存审核包。"}


def resume_plan(args: argparse.Namespace) -> dict[str, Any]:
    """Repair a stale status from existing validation artifacts; never unseal a run."""
    inspection = inspect_run(args)
    run_dir = Path(inspection["run_dir"])
    status_path = run_dir / "status.json"
    status = load_json(status_path)
    if status.get("phase") in {"awaiting_confirmation", "confirmed", "applied", "verified", "completed"}:
        return {**inspection, "resumed": False, "message": "运行已封存或已进入确认后阶段，不会自动改写状态。"}
    phases = ["prepared", "keywords_ready", "listing_ready", "title_ready", "sync_planned", "ppc_ready", "qa_ready"]
    reached = "prepared"
    found = {item["phase"] for item in inspection["completed_stages"]}
    for phase in phases[1:]:
        if phase in found:
            reached = phase
        else:
            break
    status["phase"] = reached
    status["updated_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    for item in inspection["completed_stages"]:
        if item["phase"].endswith("_ready"):
            key = {"keywords_ready": "keyword_validation", "listing_ready": "report_validation",
                   "title_ready": "title_validation", "ppc_ready": "ppc_validation",
                   "qa_ready": "rufus_qa_validation"}.get(item["phase"])
            if key:
                status.setdefault("artifacts", {})[key] = item["artifact"]
    write_json(status_path, status)
    return {**inspection, "resumed": True, "phase": reached,
            "message": "已从通过验证的产物恢复状态；未完成阶段仍需按流程生成。"}


def emit_templates(args: argparse.Namespace) -> dict[str, Any]:
    """Create schema-aligned, non-executable drafting templates in a run directory."""
    run_dir = Path(args.run_dir).expanduser().resolve()
    job = load_json(run_dir / "job.json")
    output = run_dir / "templates"
    output.mkdir(exist_ok=True)
    templates = {
        "legacy-listing-request-template.json": {"source_skill": "amazon-listing-optimization", "source_skill_path": "/Users/apple/.agents/skills/amazon-listing-optimization/SKILL.md", "mode": "A|B", "asin": job["primary_asin"], "invocation_id": "", "created_at": "", "input_manifest": {}},
        "legacy-listing-result-template.json": {"source_skill": "amazon-listing-optimization", "source_skill_path": "/Users/apple/.agents/skills/amazon-listing-optimization/SKILL.md", "source_skill_sha256": "", "invocation_id": "", "started_at": "", "request_path": "", "request_sha256": "", "raw_output_path": "", "raw_output_sha256": "", "mode": "A|B", "status": "complete", "input_manifest": {}, "legacy_listing": {"title": "", "bullets": ["", "", "", "", ""], "description": "", "backend_search_terms": ""}, "keyword_priority": {}, "keyword_coverage": {}, "keyword_gaps": {}, "limitations": []},
        "ppc-campaign-plan-template.json": {"schema_version": 1, "mode": "build", "primary_asin": job["primary_asin"], "marketplace": job["marketplace"], "currency": job["ppc_campaign"]["currency"], "financial_framework": {}, "keyword_sources": [], "campaigns": [], "copy_blocks": {key: [] for key in PPC_COPY_BLOCKS}, "budget_summary": [], "launch_schedule": [], "optimization_plan_4_weeks": [], "risk_notes": []},
        "rufus-qa-template.json": {"status": "evidence_insufficient", "target_count": 12, "marketplace": job["marketplace"], "language": job["marketplace_language"], "items": [], "evidence": [], "limitations": ["补充直接商品事实后再生成 Q/A。"]},
    }
    paths = []
    for name, value in templates.items():
        path = output / name
        if not path.exists():
            write_json(path, value)
        paths.append(str(path))
    return {"ok": True, "templates": paths, "message": "模板仅用于补齐产物，不表示已通过验证。"}


def review_job(args: argparse.Namespace) -> dict[str, Any]:
    job_path = Path(args.job).expanduser().resolve()
    job = load_json(job_path)
    review = load_json(Path(args.review).expanduser().resolve())
    raw_variants = review.get("variants")
    if not isinstance(raw_variants, list) or not raw_variants:
        fail("复核文件必须包含非空 variants 数组。")
    reviewed = [normalize_variant(item, job["marketplace"]) for item in raw_variants]
    expected = set(job["target_asins"])
    received = {item["asin"] for item in reviewed}
    if len(received) != len(reviewed):
        fail("复核后的变体 ASIN 存在重复。")
    if received != expected:
        fail("复核前后目标 ASIN 集合不一致。")
    by_asin = {item["asin"]: item for item in reviewed}
    ordered = []
    for asin in job["target_asins"]:
        item = by_asin[asin]
        item["is_primary"] = asin == job["primary_asin"]
        if item["unresolved_segments"]:
            fail(f"{asin} 仍有未解决变体片段：{', '.join(item['unresolved_segments'])}")
        missing = [attr["source_value"] for attr in item["attributes"] if not attr["marketplace_value"]]
        if missing:
            fail(f"{asin} 仍有未本地化变体属性：{', '.join(missing)}")
        ordered.append(item)
    job["variants"] = ordered
    job["model_review_required"] = False
    job["reviewed_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    output = Path(args.output).expanduser().resolve() if args.output else job_path
    write_json(output, job)
    return {"ok": True, "job": str(output), "target_asins": job["target_asins"], "model_review_required": False}


def normalize_keyword(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(text.split()).casefold()


def keyword_terms_from_section(value: Any) -> list[str]:
    """Extract keyword values without treating report metadata as keywords."""
    if isinstance(value, list):
        result: list[str] = []
        for item in value:
            result.extend(keyword_terms_from_section(item))
        return result
    if isinstance(value, dict):
        for key in ("keyword", "term", "query"):
            if key in value and str(value[key]).strip():
                return [str(value[key]).strip()]
        result = []
        ignored = {"status", "priority", "coverage", "location", "volume", "source", "reason", "notes"}
        for key, item in value.items():
            if key not in ignored:
                result.extend(keyword_terms_from_section(item))
        return result
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def validate_keyword_pool(args: argparse.Namespace) -> dict[str, Any]:
    job = load_json(Path(args.job).expanduser().resolve())
    pool_path = Path(args.pool).expanduser().resolve()
    pool = load_json(pool_path)
    errors: list[str] = []

    if pool.get("schema_version") != SCHEMA_VERSION:
        errors.append("关键词池 schema_version 必须为 1。")
    if str(pool.get("marketplace", "")).upper() != job["marketplace"]:
        errors.append("关键词池站点与 job.json 不一致。")
    if pool.get("language") != job["marketplace_language"]:
        errors.append("关键词池语言与 job.json 不一致。")
    minimum = pool.get("minimum_usable_keywords")
    if isinstance(minimum, bool) or not isinstance(minimum, int) or minimum <= 0:
        errors.append("关键词池 minimum_usable_keywords 必须是正整数。")
        minimum_value = KEYWORD_COUNT_MINIMUM
    else:
        minimum_value = minimum
    if minimum != job["keyword_source"].get("minimum_usable_keywords", KEYWORD_COUNT_MINIMUM):
        errors.append("关键词池最低有效词数量与 job.json 不一致。")

    keywords = pool.get("keywords")
    if not isinstance(keywords, list) or not keywords:
        errors.append("关键词池 keywords 必须是非空数组。")
        keywords = []
    ids: set[str] = set()
    normalized_terms: set[str] = set()
    autocomplete_matched_positions: list[int] = []
    autocomplete_unmatched_positions: list[int] = []
    core_positions: list[int] = []
    non_core_positions: list[int] = []
    source_collected_terms: set[str] = set()
    post_fallback_terms: set[str] = set()
    accepted_autocomplete = 0
    for index, entry in enumerate(keywords):
        if not isinstance(entry, dict):
            errors.append(f"关键词池第 {index + 1} 项必须是对象。")
            continue
        keyword_id = str(entry.get("id", "")).strip()
        keyword = str(entry.get("keyword", "")).strip()
        normalized = normalize_keyword(keyword)
        if not KEYWORD_ID_RE.fullmatch(keyword_id):
            errors.append(f"关键词 ID 无效：{keyword_id or '空'}。")
        elif keyword_id != f"KW-{index + 1:03d}":
            errors.append(f"关键词 ID 必须按最终优先级连续编号：预期 KW-{index + 1:03d}，实际 {keyword_id}。")
        elif keyword_id in ids:
            errors.append(f"关键词 ID 重复：{keyword_id}。")
        ids.add(keyword_id)
        if not normalized:
            errors.append(f"{keyword_id or index + 1} 的 keyword 不能为空。")
        if entry.get("normalized_keyword") != normalized:
            errors.append(f"{keyword_id or index + 1} 的 normalized_keyword 不正确。")
        if normalized in normalized_terms:
            errors.append(f"关键词池包含规范化后的重复词：{keyword}。")
        normalized_terms.add(normalized)
        sources = entry.get("sources")
        if not isinstance(sources, list) or not sources:
            errors.append(f"{keyword_id or keyword} 缺少关键词来源。")
            sources = []
        invalid_sources = {str(item) for item in sources} - KEYWORD_SOURCES
        if invalid_sources:
            errors.append(f"{keyword_id or keyword} 包含无效来源：{', '.join(sorted(invalid_sources))}。")
        if not isinstance(entry.get("metrics"), dict):
            errors.append(f"{keyword_id or keyword} 的 metrics 必须是对象。")
        if entry.get("relevance") not in KEYWORD_RELEVANCE:
            errors.append(f"{keyword_id or keyword} 的 relevance 必须为 high 或 medium。")
        matches = entry.get("selling_point_matches")
        if not isinstance(matches, list) or any(not str(item).strip() for item in matches):
            errors.append(f"{keyword_id or keyword} 的 selling_point_matches 必须是字符串数组。")
            matches = []
        for eligibility in ("listing_eligible", "ppc_eligible", "qa_eligible"):
            if not isinstance(entry.get(eligibility), bool):
                errors.append(f"{keyword_id or keyword} 的 {eligibility} 必须是布尔值。")
        if {"uploaded_file", "sellersprite"} & set(sources):
            source_collected_terms.add(normalized)
        if {"uploaded_file", "sellersprite", "amazon_autocomplete"} & set(sources):
            post_fallback_terms.add(normalized)
        if "user_core" in sources:
            core_positions.append(index)
        else:
            non_core_positions.append(index)
        if "amazon_autocomplete" in sources and "user_core" not in sources:
            accepted_autocomplete += 1
            if matches:
                autocomplete_matched_positions.append(index)
            else:
                autocomplete_unmatched_positions.append(index)

    if autocomplete_matched_positions and autocomplete_unmatched_positions:
        if max(autocomplete_matched_positions) > min(autocomplete_unmatched_positions):
            errors.append("Amazon autocomplete 关键词必须把匹配核心卖点的词排在未匹配词之前。")
    if core_positions and non_core_positions and max(core_positions) > min(non_core_positions):
        errors.append("用户核心关键词必须排在统一关键词池的最前面。")

    pre_count = pool.get("pre_fallback_usable_count")
    if pre_count != len(source_collected_terms):
        errors.append("pre_fallback_usable_count 与上传文件/卖家精灵有效唯一词数量不一致。")
    post_count = pool.get("post_fallback_usable_count")
    if post_count != len(post_fallback_terms):
        errors.append("post_fallback_usable_count 与补词后的有效唯一采集词数量不一致。")
    should_fallback = len(source_collected_terms) < minimum_value
    fallback = pool.get("fallback")
    if not isinstance(fallback, dict):
        errors.append("关键词池缺少 fallback 对象。")
        fallback = {}
    triggered = fallback.get("triggered")
    if triggered is not should_fallback:
        errors.append("关键词研究补充触发状态与少于 10 个有效采集词的规则不一致。")
    status = fallback.get("status")
    if status not in KEYWORD_FALLBACK_STATUSES:
        errors.append("fallback.status 必须为 complete、insufficient、failed 或 not_needed。")
    if should_fallback:
        if fallback.get("skill") != "amazon-keyword-research":
            errors.append("触发补词时必须记录 amazon-keyword-research 技能来源。")
        seeds = fallback.get("seeds")
        if not isinstance(seeds, list) or not seeds:
            errors.append("触发补词时必须为每个核心关键词记录本地化种子。")
            seeds = []
        expected_core = {normalize_keyword(item) for item in job["core_keywords"]}
        received_core: set[str] = set()
        successful_seed_count = 0
        for seed in seeds:
            if not isinstance(seed, dict):
                errors.append("补词种子必须是对象。")
                continue
            original = normalize_keyword(seed.get("original"))
            localized = str(seed.get("localized", "")).strip()
            received_core.add(original)
            if not localized:
                errors.append("补词种子缺少目标站点语言的 localized 值。")
            if str(seed.get("marketplace", "")).upper() != job["marketplace"]:
                errors.append("补词种子的站点与任务不一致。")
            seed_status = seed.get("status")
            if seed_status == "complete":
                successful_seed_count += 1
                artifact = Path(str(seed.get("source_artifact", ""))).expanduser().resolve()
                if not artifact.is_file() or seed.get("source_sha256") != sha256(artifact):
                    errors.append(f"补词种子的 Amazon autocomplete 原始证据无效：{artifact}。")
                else:
                    raw_evidence = load_json(artifact)
                    suggestions = raw_evidence.get("suggestions")
                    if str(raw_evidence.get("marketplace", "")).upper() != job["marketplace"]:
                        errors.append(f"补词原始证据站点与任务不一致：{artifact}。")
                    if normalize_keyword(raw_evidence.get("keyword")) != normalize_keyword(localized):
                        errors.append(f"补词原始证据种子与 localized 不一致：{artifact}。")
                    if not isinstance(suggestions, list) or any(not str(item).strip() for item in suggestions):
                        errors.append(f"补词原始证据 suggestions 必须是字符串数组：{artifact}。")
            elif seed_status == "failed":
                if not str(seed.get("error", "")).strip():
                    errors.append("失败的补词种子必须记录 error。")
            else:
                errors.append("补词种子 status 必须为 complete 或 failed。")
        if received_core != expected_core:
            errors.append("补词种子没有完整覆盖 job.json 的全部核心关键词。")
        if status == "failed" and successful_seed_count:
            errors.append("存在成功种子时 fallback.status 不得为 failed。")
        if status in {"complete", "insufficient"} and not successful_seed_count:
            errors.append(f"fallback.status={status} 时至少需要一个成功种子。")
        if status == "complete" and len(post_fallback_terms) < minimum_value:
            errors.append("有效关键词仍少于阈值时 fallback.status 不得为 complete。")
        if status == "insufficient" and len(post_fallback_terms) >= minimum_value:
            errors.append("有效关键词达到阈值时 fallback.status 不得为 insufficient。")
    elif status != "not_needed":
        errors.append("未触发补词时 fallback.status 必须为 not_needed。")

    core_terms = {normalize_keyword(item) for item in job["core_keywords"]}
    represented_core = {
        normalize_keyword(entry.get("keyword")) for entry in keywords
        if isinstance(entry, dict) and "user_core" in (entry.get("sources") or [])
    }
    if not core_terms.issubset(represented_core):
        errors.append("统一关键词池必须保留全部用户核心关键词并标记 user_core 来源。")

    excluded = pool.get("excluded")
    if not isinstance(excluded, list):
        errors.append("关键词池 excluded 必须是数组。")
    else:
        for item in excluded:
            if (
                not isinstance(item, dict)
                or not str(item.get("keyword", "")).strip()
                or not CHINESE_RE.search(str(item.get("reason", "")))
            ):
                errors.append("每个排除关键词必须包含 keyword 和中文 reason。")
                break
    warnings = pool.get("warnings")
    if not isinstance(warnings, list):
        errors.append("关键词池 warnings 必须是数组。")
        warnings = []
    if should_fallback and len(post_fallback_terms) < minimum_value:
        if not any("关键词数据不足提醒" in str(item) for item in warnings):
            errors.append("补词后仍不足 10 个时必须记录中文“关键词数据不足提醒”。")

    validation = {
        "ok": not errors,
        "marketplace": job["marketplace"],
        "language": job["marketplace_language"],
        "pool": str(pool_path),
        "pool_sha256": sha256(pool_path),
        "minimum_usable_keywords": minimum_value,
        "pre_fallback_usable_count": len(source_collected_terms),
        "post_fallback_usable_count": len(post_fallback_terms),
        "usable_keyword_count": len(normalized_terms),
        "accepted_autocomplete_count": accepted_autocomplete,
        "fallback_triggered": should_fallback,
        "keyword_ids": sorted(ids),
        "errors": errors,
    }
    if args.output:
        write_json(Path(args.output).expanduser().resolve(), validation)
    if errors:
        fail("关键词池验证失败：" + "；".join(errors))
    return validation


def load_keyword_validation(job: dict[str, Any], validation_path: Path) -> tuple[dict[str, Any], dict[str, Any]]:
    if not validation_path.is_file():
        fail(f"关键词验证文件不存在：{validation_path}")
    validation = load_json(validation_path)
    if validation.get("ok") is not True or validation.get("errors"):
        fail("统一关键词池尚未通过验证。")
    if str(validation.get("marketplace", "")).upper() != job["marketplace"]:
        fail("关键词验证结果的站点与任务不一致。")
    if validation.get("language") != job["marketplace_language"]:
        fail("关键词验证结果的语言与任务不一致。")
    pool_path = Path(str(validation.get("pool", ""))).expanduser().resolve()
    if not pool_path.is_file() or validation.get("pool_sha256") != sha256(pool_path):
        fail("统一关键词池与验证时内容不一致。")
    return validation, load_json(pool_path)


def contains_term(text: str, term: str) -> bool:
    if not term:
        return True
    if re.fullmatch(r"[A-Za-z0-9]+", term):
        return bool(re.search(rf"(?<![A-Za-z0-9]){re.escape(term)}(?![A-Za-z0-9])", text, re.IGNORECASE))
    return term.casefold() in text.casefold()


def title_components(option: dict[str, Any], asin: str) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    value = option.get("title_components")
    if not isinstance(value, dict):
        return {}, [f"{asin} 标题方案 1 缺少 title_components。"]
    for key, label in (
        ("brand", "品牌名"),
        ("core_product_phrase", "产品核心关键词（品类词）"),
        ("differentiator", "产品核心差异化卖点"),
    ):
        if not str(value.get(key, "")).strip():
            errors.append(f"{asin} 标题方案 1 缺少{label}组件。")
    sku_attributes = value.get("sku_attributes")
    if not isinstance(sku_attributes, list):
        errors.append(f"{asin} title_components.sku_attributes 必须是数组。")
    placement = str(option.get("sku_attribute_placement", "")).strip()
    if placement not in {"title", "item_highlights", "not_applicable"}:
        errors.append(f"{asin} sku_attribute_placement 无效。")
    if placement == "item_highlights" and not str(option.get("sku_move_reason", "")).strip():
        errors.append(f"{asin} SKU 属性转移到商品亮点时必须提供中文 sku_move_reason。")
    elif placement == "item_highlights" and not CHINESE_RE.search(str(option.get("sku_move_reason", ""))):
        errors.append(f"{asin} sku_move_reason 必须使用中文说明。")
    return value, errors


def load_title_validator():
    path = Path(__file__).with_name("validate-title-highlights.py")
    spec = importlib.util.spec_from_file_location("title_highlights_validator", path)
    if not spec or not spec.loader:
        fail(f"无法加载标题验证器：{path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def load_legacy_handoff_validator():
    path = Path(__file__).with_name("validate-legacy-handoff.py")
    spec = importlib.util.spec_from_file_location("legacy_handoff_validator", path)
    if not spec or not spec.loader:
        fail(f"无法加载 Listing 交接验证器：{path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def combine_reference_title(title_field: Any, item_highlight_field: Any) -> str:
    return " ".join(
        value for value in (str(title_field or "").strip(), str(item_highlight_field or "").strip())
        if value
    )


def listing_values_equal(left: Any, right: Any) -> bool:
    if isinstance(left, list) or isinstance(right, list):
        if not isinstance(left, list) or not isinstance(right, list):
            return False
        return [str(item).strip() for item in left] == [str(item).strip() for item in right]
    return str(left or "").strip() == str(right or "").strip()


def validate_retrieval_ledger(
    retrieval: Any,
    asin: str,
    marketplace: str,
    target_role: str,
    original_listing: dict[str, Any] | None = None,
) -> list[str]:
    errors: list[str] = []
    if not isinstance(retrieval, dict):
        return [f"{asin} 缺少 source_retrieval，无法验证网页优先采集链。"]
    if retrieval.get("target_role") != target_role:
        errors.append(f"{asin} source_retrieval.target_role 必须为 {target_role}。")
    if str(retrieval.get("asin", "")).strip().upper() != asin:
        errors.append(f"{asin} source_retrieval.asin 不一致。")
    try:
        ledger_marketplace = normalize_marketplace(retrieval.get("marketplace"))
    except ValueError:
        ledger_marketplace = ""
        errors.append(f"{asin} source_retrieval.marketplace 无效。")
    if ledger_marketplace and ledger_marketplace != marketplace:
        errors.append(f"{asin} source_retrieval.marketplace 与报告站点不一致。")

    required_fields = retrieval.get("required_fields")
    if not isinstance(required_fields, list):
        errors.append(f"{asin} source_retrieval.required_fields 必须为数组。")
        required_fields = []
    required_field_set = {str(field).strip() for field in required_fields if str(field).strip()}
    missing_required = set(REQUIRED_ORIGINAL_LISTING_FIELDS) - required_field_set
    if missing_required:
        errors.append(f"{asin} 网页采集必需字段缺少：{', '.join(sorted(missing_required))}。")

    attempts = retrieval.get("attempts")
    if not isinstance(attempts, list) or not attempts:
        return errors + [f"{asin} source_retrieval.attempts 必须为非空数组。"]

    previous_rank = -1
    candidates: dict[str, list[tuple[int, str, Any]]] = {}
    verified_by_web: set[str] = set()
    remaining_fields = set(required_field_set)
    observed_sources: list[str] = []
    xlsm_index: int | None = None
    authentication_blocked = False
    for index, attempt in enumerate(attempts, start=1):
        if not isinstance(attempt, dict):
            errors.append(f"{asin} 第 {index} 个采集记录不是对象。")
            continue
        source = str(attempt.get("source", "")).strip()
        rank = LISTING_SOURCE_ORDER.get(source)
        if rank is None:
            errors.append(f"{asin} 第 {index} 个采集来源无效：{source or '空'}。")
            continue
        if rank <= previous_rank:
            errors.append(f"{asin} 采集来源必须按固定链逐项执行且不得重复。")
        previous_rank = rank
        observed_sources.append(source)
        if source == "xlsm" and xlsm_index is None:
            xlsm_index = index

        if attempt.get("target_role") != target_role:
            errors.append(f"{asin} 第 {index} 个采集记录 target_role 不一致。")
        if str(attempt.get("asin", "")).strip().upper() != asin:
            errors.append(f"{asin} 第 {index} 个采集记录 ASIN 不一致。")
        try:
            attempt_marketplace = normalize_marketplace(attempt.get("marketplace"))
        except ValueError:
            attempt_marketplace = ""
            errors.append(f"{asin} 第 {index} 个采集记录 marketplace 无效。")
        if attempt_marketplace and attempt_marketplace != marketplace:
            errors.append(f"{asin} 第 {index} 个采集记录 marketplace 不一致。")
        for field in ("executor", "url", "retrieved_at", "evidence_path", "evidence_sha256"):
            if not str(attempt.get(field, "")).strip():
                errors.append(f"{asin} 第 {index} 个采集记录缺少 {field}。")
        timestamp = str(attempt.get("retrieved_at", "")).strip()
        if timestamp:
            try:
                dt.datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
            except ValueError:
                errors.append(f"{asin} 第 {index} 个采集记录 retrieved_at 不是 ISO 8601。")
        evidence_hash = str(attempt.get("evidence_sha256", "")).strip()
        evidence_path = Path(str(attempt.get("evidence_path", ""))).expanduser()
        if evidence_hash and not SHA256_RE.fullmatch(evidence_hash):
            errors.append(f"{asin} 第 {index} 个采集记录 evidence_sha256 无效。")
        if str(attempt.get("evidence_path", "")).strip():
            if not evidence_path.is_file():
                errors.append(f"{asin} 第 {index} 个采集证据文件不存在：{evidence_path}。")
            elif evidence_hash and sha256(evidence_path) != evidence_hash:
                errors.append(f"{asin} 第 {index} 个采集证据哈希不一致。")
            else:
                try:
                    receipt = load_json(evidence_path)
                except (ValueError, OSError, json.JSONDecodeError):
                    errors.append(f"{asin} 第 {index} 个采集证据必须为 JSON 回执。")
                else:
                    receipt_fields = (
                        "source", "target_role", "asin", "marketplace", "executor", "url",
                        "retrieved_at", "status", "verified_fields", "missing_fields",
                        "field_values", "failure_reason",
                    )
                    mismatched = [
                        field for field in receipt_fields
                        if receipt.get(field) != attempt.get(field)
                    ]
                    if mismatched:
                        errors.append(
                            f"{asin} 第 {index} 个采集证据与尝试记录不一致："
                            + "、".join(mismatched)
                            + "。"
                        )

        status = str(attempt.get("status", "")).strip()
        if status not in LISTING_ATTEMPT_STATUSES:
            errors.append(f"{asin} 第 {index} 个采集记录 status 无效。")
        if status in {"failed", "authentication_blocked"} and not str(
            attempt.get("failure_reason", "")
        ).strip():
            errors.append(f"{asin} 第 {index} 个失败记录必须填写 failure_reason。")
        if source in {"browser", "chrome"} and status == "authentication_blocked":
            authentication_blocked = True

        verified_fields = attempt.get("verified_fields")
        missing_fields = attempt.get("missing_fields")
        field_values = attempt.get("field_values")
        if not isinstance(verified_fields, list):
            errors.append(f"{asin} 第 {index} 个采集记录 verified_fields 必须为数组。")
            verified_fields = []
        if not isinstance(missing_fields, list):
            errors.append(f"{asin} 第 {index} 个采集记录 missing_fields 必须为数组。")
        if not isinstance(field_values, dict):
            errors.append(f"{asin} 第 {index} 个采集记录 field_values 必须为对象。")
            field_values = {}
        verified_set = {str(field).strip() for field in verified_fields if str(field).strip()}
        missing_set = {
            str(field).strip() for field in missing_fields if str(field).strip()
        } if isinstance(missing_fields, list) else set()
        already_resolved = verified_set.intersection(required_field_set) - remaining_fields
        if already_resolved:
            errors.append(
                f"{asin} 第 {index} 个来源重复验证已由高优先级来源取得的字段："
                + "、".join(sorted(already_resolved))
                + "。"
            )
        expected_missing = remaining_fields - verified_set
        if missing_set != expected_missing:
            errors.append(f"{asin} 第 {index} 个来源 missing_fields 未准确反映当时剩余字段。")
        if set(field_values) != verified_set:
            errors.append(f"{asin} 第 {index} 个来源 field_values 必须且只能包含 verified_fields。")
        if status == "complete" and (not verified_set or expected_missing):
            errors.append(f"{asin} 第 {index} 个 complete 来源必须取得全部当时剩余字段。")
        if status == "partial" and (not verified_set or not expected_missing):
            errors.append(f"{asin} 第 {index} 个 partial 来源必须同时包含已验证和仍缺失字段。")
        if status in {"failed", "authentication_blocked"} and verified_set:
            errors.append(f"{asin} 第 {index} 个失败或认证阻塞来源不得声明已验证字段。")
        for field in verified_fields:
            field_name = str(field).strip()
            value = field_values.get(field_name)
            if value in (None, "", []):
                errors.append(
                    f"{asin} 第 {index} 个来源声明已验证 {field_name}，但缺少 field_values.{field_name}。"
                )
                continue
            candidates.setdefault(field_name, []).append((index, source, value))
            if source != "xlsm":
                verified_by_web.add(field_name)
            elif field_name in verified_by_web:
                errors.append(f"{asin} XLSM 不得覆盖网页已验证字段：{field_name}。")
        remaining_fields -= verified_set.intersection(required_field_set)

    if authentication_blocked:
        errors.append(f"{asin} Browser 或 Chrome 存在登录/CAPTCHA/OTP 阻塞，禁止继续回退。")

    resolved_fields = retrieval.get("resolved_fields")
    if not isinstance(resolved_fields, dict):
        errors.append(f"{asin} 缺少 source_retrieval.resolved_fields。")
        resolved_fields = {}
    for field, field_candidates in candidates.items():
        resolved = resolved_fields.get(field)
        if not isinstance(resolved, dict):
            errors.append(f"{asin} 缺少 resolved_fields.{field}。")
            continue
        expected_index, expected_source, expected_value = field_candidates[0]
        if (
            resolved.get("attempt_index") != expected_index
            or resolved.get("source") != expected_source
            or not retrieval_values_equal(resolved.get("value"), expected_value)
        ):
            errors.append(f"{asin} 字段 {field} 未使用最高优先级的已验证来源 {expected_source}。")
        if original_listing is not None and field in required_field_set:
            if field not in original_listing or original_listing.get(field) in (None, "", []):
                errors.append(f"{asin} original_listing 缺少已解析字段 {field}。")
            elif not retrieval_values_equal(original_listing.get(field), expected_value):
                errors.append(f"{asin} original_listing.{field} 与已解析来源值不一致。")

    unresolved_fields = required_field_set - set(candidates)
    raw_missing_fields = retrieval.get("missing_fields")
    if not isinstance(raw_missing_fields, list):
        errors.append(f"{asin} source_retrieval.missing_fields 必须为数组。")
        raw_missing_fields = []
    ledger_missing = {str(field).strip() for field in raw_missing_fields if str(field).strip()}
    if unresolved_fields != ledger_missing:
        errors.append(f"{asin} source_retrieval.missing_fields 与实际未解析字段不一致。")
    if unresolved_fields:
        limitations = retrieval.get("limitations")
        if not isinstance(limitations, list) or not any(
            "数据获取提醒" in str(item) for item in limitations
        ):
            errors.append(f"{asin} 全部来源后仍缺字段时必须包含中文数据获取提醒。")

    all_required_resolved_before_end = bool(required_field_set) and required_field_set.issubset(
        verified_by_web
    )
    if xlsm_index is not None:
        if target_role != "optimized_asin":
            errors.append(f"{asin} 竞品证据禁止使用 XLSM。")
        expected_prefix = list(WEB_LISTING_SOURCE_CHAIN)
        actual_prefix = observed_sources[: xlsm_index - 1]
        if actual_prefix != expected_prefix:
            errors.append(
                f"{asin} 使用 XLSM 前必须完整执行网页来源链："
                + " → ".join(WEB_LISTING_SOURCE_CHAIN)
                + "。"
            )
        if all_required_resolved_before_end:
            errors.append(f"{asin} 网页已取得全部必需 Listing 字段，不得再使用 XLSM。")
    elif unresolved_fields and observed_sources != list(WEB_LISTING_SOURCE_CHAIN):
        errors.append(
            f"{asin} Listing 字段仍缺失时必须完成全部网页来源链，才能判定无数据。"
        )
    return errors


def validate_original_listing_binding(report: dict[str, Any], asin: str) -> list[str]:
    errors: list[str] = []
    original = report.get("original_listing")
    if not isinstance(original, dict):
        return [f"{asin} 缺少 original_listing，无法确认优化前 Listing。"]
    original_title = str(original.get("title", "")).strip()
    if not original_title:
        errors.append(f"{asin} 缺少 original_listing.title。")

    retrieval = report.get("source_retrieval")
    try:
        marketplace = normalize_marketplace(report.get("marketplace"))
    except ValueError:
        marketplace = ""
        errors.append(f"{asin} 报告站点无效，无法验证网页采集来源。")
    errors.extend(validate_retrieval_ledger(
        retrieval, asin, marketplace, "optimized_asin", original,
    ))
    if not isinstance(retrieval, dict):
        return errors
    resolved_title = (retrieval.get("resolved_fields") or {}).get("title")
    if not isinstance(resolved_title, dict):
        return errors + [f"{asin} 缺少 source_retrieval.resolved_fields.title。"]
    resolved_value = str(resolved_title.get("value", "")).strip()
    resolved_source = str(resolved_title.get("source", "")).strip()
    if resolved_value != original_title:
        errors.append(f"{asin} original_listing.title 与已解析标题值不一致。")
    if resolved_source not in LISTING_SOURCE_ORDER:
        errors.append(f"{asin} 优化前标题来源无效：{resolved_source or '空'}。")

    if resolved_source == "xlsm":
        xlsm = retrieval.get("xlsm_fallback")
        if not isinstance(xlsm, dict):
            errors.append(f"{asin} 使用 XLSM 标题时必须提供 source_retrieval.xlsm_fallback。")
        else:
            combined = combine_reference_title(xlsm.get("title_field"), xlsm.get("item_highlight_field"))
            if str(xlsm.get("reference_title", "")).strip() != combined:
                errors.append(f"{asin} XLSM reference_title 未按标题字段和亮点字段组合。")
            if resolved_value != combined:
                errors.append(f"{asin} 优化前标题与 XLSM reference_title 不一致。")

    before_after = report.get("before_after") or []
    title_rows = [
        row for row in before_after
        if isinstance(row, dict) and str(row.get("section", "")).strip().casefold() in {"title", "标题", "商品标题"}
    ]
    if not title_rows:
        errors.append(f"{asin} before_after 缺少标题对比行。")
    elif any(str(row.get("before", "")).strip() != original_title for row in title_rows):
        errors.append(f"{asin} before_after 的优化前标题与 original_listing.title 不一致。")
    return errors


def validate_reports(args: argparse.Namespace) -> dict[str, Any]:
    job = load_json(Path(args.job).resolve())
    keyword_validation_path = Path(args.keyword_validation).expanduser().resolve()
    keyword_validation, keyword_pool = load_keyword_validation(job, keyword_validation_path)
    keyword_pool_sha = keyword_validation["pool_sha256"]
    listing_keywords = {
        normalize_keyword(item.get("keyword")) for item in keyword_pool.get("keywords", [])
        if isinstance(item, dict) and item.get("listing_eligible") is True
    }
    expected = set(job["target_asins"])
    report_paths = [Path(item).expanduser().resolve() for item in args.reports]
    reports: dict[str, tuple[Path, dict[str, Any]]] = {}
    errors: list[str] = []
    for path in report_paths:
        report = load_json(path)
        try:
            asin = normalize_asin(report.get("asin"), f"报告 {path.name} 的 ASIN")
        except ValueError as exc:
            errors.append(str(exc))
            continue
        if asin in reports:
            errors.append(f"ASIN {asin} 存在重复报告。")
        reports[asin] = (path, report)
    missing = sorted(expected - set(reports))
    extra = sorted(set(reports) - expected)
    if missing:
        errors.append("缺少报告：" + ", ".join(missing))
    if extra:
        errors.append("存在非目标报告：" + ", ".join(extra))

    validator = load_title_validator()
    handoff_validator = load_legacy_handoff_validator()
    variants = {item["asin"]: item for item in job["variants"]}
    color_terms = {
        item["asin"]: [attr["marketplace_value"] for attr in item["attributes"] if attr["type"] == "color" and attr["marketplace_value"]]
        for item in job["variants"]
    }
    checked: list[dict[str, Any]] = []
    for asin in sorted(expected & set(reports)):
        path, report = reports[asin]
        legacy = report.get("legacy_generation") if isinstance(report.get("legacy_generation"), dict) else {}
        legacy_artifact_path = Path(str(legacy.get("artifact_path", ""))).expanduser()
        if not legacy_artifact_path.is_file():
            errors.append(f"{asin} 缺少传统 Listing 交接文件：{legacy_artifact_path}")
            legacy_artifact = None
        else:
            legacy_artifact_path = legacy_artifact_path.resolve()
            legacy_artifact = load_json(legacy_artifact_path)
        handoff_result = handoff_validator.validate_report(report, legacy_artifact)
        errors.extend(f"{asin} Listing 交接：{item}" for item in handoff_result.get("errors", []))
        if normalize_marketplace(report.get("marketplace")) != job["marketplace"]:
            errors.append(f"{asin} 报告站点与任务站点不一致。")
        if report.get("keyword_pool_sha256") != keyword_pool_sha:
            errors.append(f"{asin} 报告未绑定当前统一关键词池。")
        missing_audit = [section for section in REQUIRED_AUDIT_SECTIONS if report.get(section) in (None, "", [], {})]
        audit = report.get("audit") if isinstance(report.get("audit"), dict) else {}
        if not audit.get("dimensions"):
            missing_audit.append("audit.dimensions")
        if missing_audit:
            errors.append(f"{asin} 缺少确认审核包必需字段：{', '.join(missing_audit)}。")
        for section in ("keyword_priority", "keyword_coverage", "keyword_gaps"):
            if section not in report:
                errors.append(f"{asin} 报告缺少 {section}。")
                continue
            terms = keyword_terms_from_section(report.get(section))
            if section != "keyword_gaps" and not terms:
                errors.append(f"{asin} 报告的 {section} 没有引用统一关键词池。")
            unknown = sorted({term for term in terms if normalize_keyword(term) not in listing_keywords})
            if unknown:
                errors.append(f"{asin} 报告的 {section} 包含统一关键词池外的词：{', '.join(unknown)}。")
        listing = report.get("listing") or {}
        for field in ("title", "description", "backend_search_terms"):
            if not str(listing.get(field, "")).strip():
                errors.append(f"{asin} 缺少 listing.{field}。")
        bullets = listing.get("bullets") or []
        if len(bullets) != 5 or any(not str(item).strip() for item in bullets):
            errors.append(f"{asin} 必须包含 5 条非空五点描述。")
        options = report.get("title_options_2026") or []
        option = next((item for item in options if item.get("option") == 1), None)
        if not option:
            errors.append(f"{asin} 缺少 2026 标题方案 1。")
            continue
        title = str(option.get("title", ""))
        highlights = str(option.get("item_highlights", ""))
        components, component_errors = title_components(option, asin)
        errors.extend(component_errors)
        sku_attributes = components.get("sku_attributes") if isinstance(components.get("sku_attributes"), list) else []
        placement = str(option.get("sku_attribute_placement", ""))
        validation = validator.validate(
            title,
            highlights,
            job["marketplace"],
            brand=str(components.get("brand", "")),
            core_product_phrase=str(components.get("core_product_phrase", "")),
            differentiator=str(components.get("differentiator", "")),
            sku_attributes=[str(item) for item in sku_attributes],
            sku_attribute_placement=placement,
        )
        if not validation["valid"]:
            errors.append(f"{asin} 标题方案 1 未通过校验：{'；'.join(validation['errors'])}")
        for term in variants[asin].get("expected_title_terms", []):
            if not any(contains_term(str(item), term) for item in sku_attributes):
                errors.append(f"{asin} 标题方案 1 的 SKU 组件缺少变体属性：{term}")
            target_field = highlights if placement == "item_highlights" else title
            if not contains_term(target_field, term):
                field_name = "商品亮点" if placement == "item_highlights" else "标题"
                errors.append(f"{asin} 标题方案 1 的{field_name}缺少变体属性：{term}")
        if variants[asin].get("expected_title_terms"):
            if placement not in {"title", "item_highlights"}:
                errors.append(f"{asin} 子体必须声明 SKU 属性位于标题或商品亮点。")
        elif placement != "not_applicable":
            errors.append(f"{asin} 无 SKU 核心属性时必须使用 sku_attribute_placement: not_applicable。")
        foreign_colors = {
            term for other_asin, terms in color_terms.items() if other_asin != asin for term in terms
        }
        own_colors = set(color_terms.get(asin, []))
        for term in sorted(foreign_colors - own_colors):
            if contains_term(title, term) or contains_term(highlights, term):
                errors.append(f"{asin} 标题方案 1 残留其他变体颜色：{term}")
        checked.append({
            "asin": asin,
            "report": str(path),
            "report_sha256": sha256(path),
            "legacy_artifact": str(legacy_artifact_path) if legacy_artifact is not None else "",
            "legacy_artifact_sha256": sha256(legacy_artifact_path) if legacy_artifact is not None else "",
            "legacy_request": str(Path(str(legacy.get("request_path", ""))).expanduser().resolve()),
            "legacy_request_sha256": str(legacy.get("request_sha256", "")),
            "legacy_raw_output": str(Path(str(legacy.get("raw_output_path", ""))).expanduser().resolve()),
            "legacy_raw_output_sha256": str(legacy.get("raw_output_sha256", "")),
            "source_skill_sha256": str(legacy.get("source_skill_sha256", "")),
            "title": title,
            "item_highlights": highlights,
            "title_validation": validation,
        })
    result = {
        "ok": not errors,
        "marketplace": job["marketplace"],
        "keyword_validation": str(keyword_validation_path),
        "keyword_pool_sha256": keyword_pool_sha,
        "expected_asins": sorted(expected),
        "checked": checked,
        "errors": errors,
    }
    if args.output:
        write_json(Path(args.output).resolve(), result)
    if errors:
        fail("报告验证失败：" + "；".join(errors))
    return result


def verify_report_validation(
    job: dict[str, Any], validation_path: Path,
    report_paths: list[Path], keyword_pool_sha: str,
) -> dict[str, Any]:
    validation = load_json(validation_path)
    if validation.get("ok") is not True or validation.get("errors"):
        fail("Listing 与标题报告尚未通过验证。")
    if str(validation.get("marketplace", "")).upper() != job["marketplace"]:
        fail("报告验证结果的站点与任务不一致。")
    if validation.get("keyword_pool_sha256") != keyword_pool_sha:
        fail("报告验证结果未绑定当前统一关键词池。")
    if set(validation.get("expected_asins") or []) != set(job["target_asins"]):
        fail("报告验证结果的 ASIN 集合与任务不一致。")
    checked = {
        str(item.get("asin", "")).upper(): item
        for item in validation.get("checked", []) if isinstance(item, dict)
    }
    expected_paths = {str(path.resolve()): path.resolve() for path in report_paths}
    if set(checked) != set(job["target_asins"]):
        fail("报告验证结果未覆盖全部目标 ASIN。")
    for asin, item in checked.items():
        report_path = Path(str(item.get("report", ""))).expanduser().resolve()
        if str(report_path) not in expected_paths:
            fail(f"报告验证结果引用了错误文件：{report_path}")
        if not report_path.is_file() or item.get("report_sha256") != sha256(report_path):
            fail(f"已验证的 Listing 报告发生变化：{asin}")
        report = load_json(report_path)
        report_legacy = report.get("legacy_generation") or {}
        source_skill_path = Path(str(report_legacy.get("source_skill_path", ""))).expanduser().resolve()
        if (
            not source_skill_path.is_file()
            or item.get("source_skill_sha256") != sha256(source_skill_path)
            or item.get("source_skill_sha256") != report_legacy.get("source_skill_sha256")
        ):
            fail(f"传统 Listing 来源技能已变化或未绑定：{asin}")
        legacy_path = Path(str(item.get("legacy_artifact", ""))).expanduser().resolve()
        report_legacy_path = Path(
            str(report_legacy.get("artifact_path", ""))
        ).expanduser().resolve()
        if legacy_path != report_legacy_path:
            fail(f"报告验证结果引用的传统 Listing 交接文件不一致：{asin}")
        if not legacy_path.is_file() or item.get("legacy_artifact_sha256") != sha256(legacy_path):
            fail(f"已验证的传统 Listing 交接文件发生变化：{asin}")
        for label, checked_path_key, checked_hash_key, report_path_key, report_hash_key in (
            ("传统 Listing 请求", "legacy_request", "legacy_request_sha256", "request_path", "request_sha256"),
            ("传统 Listing 原始输出", "legacy_raw_output", "legacy_raw_output_sha256", "raw_output_path", "raw_output_sha256"),
        ):
            artifact_path = Path(str(item.get(checked_path_key, ""))).expanduser().resolve()
            expected_path = Path(str(report_legacy.get(report_path_key, ""))).expanduser().resolve()
            if artifact_path != expected_path:
                fail(f"报告验证结果引用的{label}不一致：{asin}")
            if not artifact_path.is_file() or item.get(checked_hash_key) != sha256(artifact_path):
                fail(f"已验证的{label}发生变化：{asin}")
            if item.get(checked_hash_key) != report_legacy.get(report_hash_key):
                fail(f"{label}哈希未绑定 Listing 报告：{asin}")
    return validation


def markdown_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, list):
        return "；".join(markdown_value(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).replace("\n", "<br>").replace("|", "\\|")


def markdown_table(records: list[dict[str, Any]], columns: list[str] | None = None) -> list[str]:
    if not records:
        return ["无"]
    if columns is None:
        columns = []
        for record in records:
            for key in record:
                if key not in columns:
                    columns.append(key)
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for record in records:
        lines.append("| " + " | ".join(markdown_value(record.get(column, "")) for column in columns) + " |")
    return lines


def render_review_section(title: str, value: Any, level: int = 4) -> list[str]:
    lines = [f"{'#' * level} {title}", ""]
    if value in (None, "", [], {}):
        return lines + ["无", ""]
    if isinstance(value, list):
        if all(isinstance(item, dict) for item in value):
            lines.extend(markdown_table(value))
        else:
            lines.extend(f"{index}. {markdown_value(item)}" for index, item in enumerate(value, 1))
        return lines + [""]
    if isinstance(value, dict):
        scalars = {key: item for key, item in value.items() if not isinstance(item, (list, dict))}
        if scalars:
            lines.extend(markdown_table([{"项目": key, "内容": item} for key, item in scalars.items()], ["项目", "内容"]))
            lines.append("")
        for key, item in value.items():
            if isinstance(item, (list, dict)):
                lines.extend(render_review_section(key, item, level + 1))
        return lines
    return lines + [markdown_value(value), ""]


def build_listing_plan_markdown(packet: dict[str, Any]) -> str:
    lines = [
        "# Listing 优化方案", "",
        "> 以下为确认执行前的完整 Listing 优化结果与审核依据。", "",
    ]
    for item in packet["listings"]:
        listing = item["legacy_listing"]
        audit_report = item["audit_report"]
        legacy_generation = audit_report.get("legacy_generation") or {}
        title_options = audit_report.get("title_options_2026") or []
        option_one = next(
            (option for option in title_options if isinstance(option, dict) and option.get("option") == 1),
            {},
        )
        lines.extend([
            f"## ASIN {item['asin']}", "",
            "### amazon-listing-optimization 输出（旧标题格式）", "",
            f"- 生成来源：`{legacy_generation.get('source_skill', '')}`",
            f"- 来源技能：`{legacy_generation.get('source_skill_path', '')}`",
            f"- 执行模式：`{legacy_generation.get('mode', '')}`",
            f"- 原始输出：`{legacy_generation.get('raw_output_path', '')}`", "",
        ])
        lines.extend(["#### 优化后的旧格式标题", "", "```text", listing["title"], "```", ""])
        lines.extend([f"字符数：{len(str(listing['title']))}/200", ""])
        for index, bullet in enumerate(listing["bullets"], 1):
            lines.extend([f"#### Bullet Point {index}", "", "```text", bullet, "```", ""])
        lines.extend(["#### Product Description", "", "```text", listing["description"], "```", ""])
        lines.extend(["#### Backend Search Terms", "", "```text", listing["backend_search_terms"], "```", ""])
        lines.extend(["### 标题来源链", ""])
        lines.extend(markdown_table([
            {"阶段": "传统 Listing 优化", "标题": listing.get("title", ""), "来源": "amazon-listing-optimization Mode B 原始输出"},
            {"阶段": "2026 推荐方案 1", "标题": option_one.get("title", ""), "来源": "本技能内置 2026 标题逻辑"},
        ], ["阶段", "标题", "来源"]))
        lines.append("")
        lines.extend(["### Listing 优化审核报告", ""])
        for section in AUDIT_SECTIONS:
            lines.extend(render_review_section(AUDIT_LABELS.get(section, section), audit_report.get(section)))
    return "\n".join(lines)


def build_short_title_markdown(packet: dict[str, Any]) -> str:
    lines = [
        "# 2026 短标题与商品亮点方案", "",
        "> 以下方案基于 amazon-listing-optimization 生成的精确优化标题，方案 1 为默认写入方案。", "",
    ]
    for item in packet["listings"]:
        report = item["audit_report"]
        generation = report.get("title_options_generation") or {}
        legacy = report.get("legacy_generation") or {}
        reference_title = generation.get("reference_title") or (legacy.get("legacy_listing") or {}).get("title", "")
        lines.extend([f"## ASIN {item['asin']}", ""])
        # This is an exact-title audit field, not a table cell. Preserve its
        # literal content so preview verification can compare it verbatim.
        lines.extend(["### 旧版优化参考标题", "", str(reference_title), ""])
        lines.extend(render_review_section("核心流量词选择", generation.get("core_keyword_selection"), 3))
        lines.extend(render_review_section("标题组合事实词组池", generation.get("title_pair_fact_pool"), 3))
        options = sorted(
            (option for option in report.get("title_options_2026", []) if isinstance(option, dict)),
            key=lambda option: int(option.get("rank", option.get("option", 999))),
        )
        for option in options:
            option_number = option.get("option", "")
            lines.extend([
                f"### 方案 {option_number}", "",
                f"- 排名：{markdown_value(option.get('rank', option_number))}",
                f"- 状态：{markdown_value(option.get('status', ''))}",
                f"- 技能内部选优分：{markdown_value((option.get('quality_assessment') or {}).get('total', option.get('score', '')))}/100（不是 Amazon 官方算法分）", "",
                "#### 商品标题", "", "```text", str(option.get("title", "")), "```", "",
                f"字符数：{markdown_value(option.get('title_characters', ''))}/75", "",
                "#### 商品亮点", "", "```text", str(option.get("item_highlights", "")), "```", "",
                f"字符数：{markdown_value(option.get('item_highlights_characters', ''))}/125", "",
            ])
            lines.extend(render_review_section("分项评分", option.get("score_breakdown"), 4))
            lines.extend(render_review_section("标题组件", option.get("title_components"), 4))
            lines.extend(render_review_section("SKU 属性放置", {
                "位置": option.get("sku_attribute_placement", ""),
                "移动原因": option.get("sku_move_reason", ""),
                "适配计算": option.get("sku_fit_evaluation", {}),
            }, 4))
            lines.extend(render_review_section("跨字段互补与去重说明", option.get("deduplication_notes"), 4))
            lines.extend(render_review_section("标题组合分配表", option.get("title_pair_coverage_audit"), 4))
            lines.extend(render_review_section("核心策略", option.get("core_strategy"), 4))
        lines.extend(render_review_section("参考标题词组分析", generation.get("reference_phrase_analysis"), 3))
    return "\n".join(lines)


def build_xlsm_replacement_markdown(packet: dict[str, Any]) -> str:
    lines = [
        "# XLSM 文件替换方案", "",
        "> 以下内容来自只读同步计划，代表确认后将写入分类商品报告的实际位置和值。", "",
        "## 执行计划", "",
    ]
    plan_rows = []
    for item in packet["listings"]:
        destination = item["destination"]
        plan_rows.append({
            "站点": packet["marketplace"], "ASIN": item["asin"],
            "分类商品报告": destination["workbook"], "工作表": destination["sheet"],
            "行": ", ".join(map(str, destination["rows"])), "标题模式": destination["title_mode"],
            "站点证据": ", ".join(destination["station_evidence"]),
        })
    lines.extend(markdown_table(plan_rows))
    lines.extend(["", "## 结构化变体", ""])
    variant_rows = []
    for variant in packet["variants"]:
        attributes = [
            f"{attr.get('type', '')}: {attr.get('source_value', '')} -> {attr.get('marketplace_value', '')}"
            for attr in variant.get("attributes", [])
        ]
        variant_rows.append({
            "ASIN": variant["asin"], "主ASIN": variant.get("is_primary", False),
            "原始输入": variant.get("raw", ""), "解析属性": "；".join(attributes) or "无",
        })
    lines.extend(markdown_table(variant_rows))
    lines.append("")
    for item in packet["listings"]:
        listing = item["pending_listing"]
        destination = item["destination"]
        lines.extend([f"## ASIN {item['asin']}", "", "### 即将替换的完整字段值", ""])
        values = [
            ("Title", listing["title"]), ("Item Name", listing["item_name"]),
            ("Item Highlight", listing["item_highlight"]),
        ]
        values.extend((f"Bullet Point {index}", bullet) for index, bullet in enumerate(listing["bullets"], 1))
        values.extend([
            ("Product Description", listing["description"]),
            ("Backend Search Terms", listing["backend_search_terms"]),
        ])
        for label, value in values:
            lines.extend([f"#### {label}", "", "```text", value, "```", ""])
        lines.extend(["### 字段写入位置", ""])
        lines.extend(markdown_table(destination["planned_updates"], ["field", "letter", "hidden", "source_state", "action"]))
        lines.append("")
    return "\n".join(lines)


def require_text_file(path: Path, label: str) -> str:
    if not path.is_file():
        fail(f"{label}不存在：{path}")
    content = path.read_text(encoding="utf-8")
    if not content.strip():
        fail(f"{label}不能为空：{path}")
    return content


def verify_listing_preview(packet: dict[str, Any], path: Path) -> str:
    content = require_text_file(path, "Listing 优化方案文本文件")
    required_snippets = ["Listing", "审核"]
    required_section_labels = {
        "audit": ("审核报告", "评分"),
        "keyword_coverage": ("关键词覆盖",),
        "before_after": ("修改对比", "优化前后"),
        "issues_fixed": ("已修复问题",),
        "recommendations": ("建议",),
        "working_well": ("原Listing中表现较好的部分", "原 Listing 中表现较好的部分"),
    }
    for item in packet["listings"]:
        listing = item["legacy_listing"]
        legacy_generation = (item.get("audit_report") or {}).get("legacy_generation") or {}
        required_snippets.extend([
            item["asin"], "amazon-listing-optimization 输出（旧标题格式）",
            legacy_generation.get("source_skill", ""), listing["title"],
            *listing["bullets"], listing["description"], listing["backend_search_terms"],
        ])
    missing = [
        snippet for snippet in required_snippets
        if str(snippet).strip() and str(snippet).strip() not in content
    ]
    if missing:
        fail(
            "Listing 优化方案未完整展示待写入字段或审核内容："
            + "；".join(dict.fromkeys(map(str, missing)))
        )
    missing_sections = [
        section for section, labels in required_section_labels.items()
        if not any(label in content for label in labels)
    ]
    if missing_sections:
        fail("Listing 优化方案缺少完整中文审核章节：" + "、".join(missing_sections))
    return content


def verify_short_title_preview(packet: dict[str, Any], path: Path) -> str:
    content = require_text_file(path, "2026 短标题与商品亮点方案文本文件")
    required: list[str] = ["2026 短标题与商品亮点方案", "方案 1", "方案 2", "方案 3"]
    for item in packet["listings"]:
        report = item["audit_report"]
        reference_title = str((report.get("title_options_generation") or {}).get("reference_title", ""))
        legacy_title = str(item["legacy_listing"].get("title", ""))
        if reference_title != legacy_title:
            fail(f"{item['asin']} 的 2026 参考标题与传统旧格式标题不一致。")
        required.append(reference_title)
        options = [option for option in report.get("title_options_2026", []) if isinstance(option, dict)]
        if len(options) != 3:
            fail(f"{item['asin']} 必须包含 3 组 2026 标题与商品亮点方案。")
        option_one = next((option for option in options if option.get("option") == 1), None)
        if not option_one or option_one.get("status") != "recommended_for_current_upload":
            fail(f"{item['asin']} 的方案 1 必须标记为推荐写入方案。")
        required.append(item["asin"])
        for option in options:
            required.extend([str(option.get("title", "")), str(option.get("item_highlights", ""))])
    missing = [value for value in required if value.strip() and value.strip() not in content]
    if missing:
        fail("2026 短标题文档未完整展示方案：" + "；".join(dict.fromkeys(missing)))
    return content


def verify_rufus_preview(
    report: dict[str, Any], job: dict[str, Any], keyword_pool: dict[str, Any],
    keyword_pool_sha: str, path: Path,
) -> str:
    validator = load_rufus_qa_module()
    qa = validator.extract_qa(report)
    errors = list(validator.validate_rufus_qa(qa, job["marketplace"]).get("errors") or [])
    if qa.get("keyword_pool_sha256") != keyword_pool_sha:
        errors.append("Rufus Q/A 方案未绑定当前统一关键词池。")
    eligible_keywords = {
        str(item.get("id")): normalize_keyword(item.get("keyword"))
        for item in keyword_pool.get("keywords", [])
        if isinstance(item, dict) and item.get("qa_eligible") is True
    }
    for item in qa.get("items", []):
        if not isinstance(item, dict):
            continue
        qa_id = str(item.get("id", "") or "未命名 Q/A")
        refs = item.get("keyword_refs")
        if not isinstance(refs, list) or not refs:
            errors.append(f"{qa_id} 缺少 keyword_refs。")
            continue
        invalid_refs = [str(ref) for ref in refs if str(ref) not in eligible_keywords]
        if invalid_refs:
            errors.append(f"{qa_id} 引用了非 qa_eligible 关键词：{', '.join(invalid_refs)}。")
            continue
        qa_text = normalize_keyword(f"{item.get('question', '')} {item.get('answer', '')}")
        if not any(eligible_keywords[str(ref)] in qa_text for ref in refs):
            errors.append(f"{qa_id} 的问题或答案未自然包含 keyword_refs 对应关键词。")
    if errors:
        fail("Rufus Q/A 方案验证失败：" + "；".join(errors))
    content = require_text_file(path, "Rufus Q/A 方案文本文件")
    expected = validator.render_markdown(qa).rstrip()
    if content.rstrip() != expected:
        fail("Rufus Q/A 方案 Markdown 与主 ASIN Listing 报告中的已验证 Q/A 不一致。")
    return content


def validate_rufus_plan(args: argparse.Namespace) -> dict[str, Any]:
    job = load_json(Path(args.job).expanduser().resolve())
    keyword_validation_path = Path(args.keyword_validation).expanduser().resolve()
    keyword_validation, keyword_pool = load_keyword_validation(job, keyword_validation_path)
    keyword_pool_sha = keyword_validation["pool_sha256"]
    ppc_validation_path = Path(args.ppc_validation).expanduser().resolve()
    verify_ppc_validation(job, ppc_validation_path, keyword_pool_sha)
    report_path = Path(args.report).expanduser().resolve()
    report = load_json(report_path)
    errors: list[str] = []
    if str(report.get("asin", "")).upper() != job["primary_asin"]:
        errors.append("Rufus Q/A 只能使用主 ASIN Listing 报告。")
    if str(report.get("marketplace", "")).upper() != job["marketplace"]:
        errors.append("Rufus Q/A Listing 报告站点与任务不一致。")
    if report.get("keyword_pool_sha256") != keyword_pool_sha:
        errors.append("Rufus Q/A Listing 报告未绑定当前统一关键词池。")
    validator = load_rufus_qa_module()
    qa = validator.extract_qa(report)
    errors.extend(validator.validate_rufus_qa(qa, job["marketplace"]).get("errors") or [])
    if qa.get("keyword_pool_sha256") != keyword_pool_sha:
        errors.append("Rufus Q/A 未绑定当前统一关键词池。")
    eligible_keywords = {
        str(item.get("id")): normalize_keyword(item.get("keyword"))
        for item in keyword_pool.get("keywords", [])
        if isinstance(item, dict) and item.get("qa_eligible") is True
    }
    for item in qa.get("items", []):
        if not isinstance(item, dict):
            continue
        qa_id = str(item.get("id", "") or "未命名 Q/A")
        refs = item.get("keyword_refs")
        if not isinstance(refs, list) or not refs:
            errors.append(f"{qa_id} 缺少 keyword_refs。")
            continue
        invalid_refs = [str(ref) for ref in refs if str(ref) not in eligible_keywords]
        if invalid_refs:
            errors.append(f"{qa_id} 引用了非 qa_eligible 关键词：{', '.join(invalid_refs)}。")
            continue
        qa_text = normalize_keyword(f"{item.get('question', '')} {item.get('answer', '')}")
        if not any(eligible_keywords[str(ref)] in qa_text for ref in refs):
            errors.append(f"{qa_id} 的问题或答案未自然包含 keyword_refs 对应关键词。")
    markdown_path = Path(args.markdown).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    result = {
        "ok": not errors,
        "primary_asin": job["primary_asin"],
        "marketplace": job["marketplace"],
        "keyword_pool_sha256": keyword_pool_sha,
        "ppc_validation": str(ppc_validation_path),
        "report": str(report_path),
        "report_sha256": sha256(report_path),
        "markdown": str(markdown_path),
        "errors": errors,
    }
    if errors:
        write_json(output_path, result)
        fail("Rufus Q/A 方案验证失败：" + "；".join(errors))
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(validator.render_markdown(qa), encoding="utf-8")
    result["markdown_sha256"] = sha256(markdown_path)
    write_json(output_path, result)
    return result


def verify_rufus_validation(
    job: dict[str, Any], validation_path: Path, report_path: Path,
    markdown_path: Path, keyword_pool_sha: str,
) -> dict[str, Any]:
    validation = load_json(validation_path)
    if validation.get("ok") is not True or validation.get("errors"):
        fail("Rufus Q/A 方案尚未通过验证。")
    if str(validation.get("primary_asin", "")).upper() != job["primary_asin"]:
        fail("Rufus Q/A 验证结果的主 ASIN 与任务不一致。")
    if str(validation.get("marketplace", "")).upper() != job["marketplace"]:
        fail("Rufus Q/A 验证结果的站点与任务不一致。")
    if validation.get("keyword_pool_sha256") != keyword_pool_sha:
        fail("Rufus Q/A 验证结果未绑定当前统一关键词池。")
    if Path(str(validation.get("report", ""))).expanduser().resolve() != report_path:
        fail("Rufus Q/A 验证结果引用的 Listing 报告不一致。")
    if validation.get("report_sha256") != sha256(report_path):
        fail("Rufus Q/A 验证后的 Listing 报告发生变化。")
    if Path(str(validation.get("markdown", ""))).expanduser().resolve() != markdown_path:
        fail("Rufus Q/A 验证结果引用的 Markdown 不一致。")
    if validation.get("markdown_sha256") != sha256(markdown_path):
        fail("Rufus Q/A 验证后的 Markdown 发生变化。")
    return validation


def verify_ppc_validation(
    job: dict[str, Any], validation_path: Path, keyword_pool_sha: str | None = None,
) -> dict[str, Any]:
    if not validation_path.is_file():
        fail(f"PPC 验证文件不存在：{validation_path}")
    validation = load_json(validation_path)
    if validation.get("ok") is not True or validation.get("errors"):
        fail("PPC 方案尚未通过验证，不能请求确认执行。")
    if str(validation.get("primary_asin", "")).upper() != job["primary_asin"]:
        fail("PPC 验证结果的主 ASIN 与任务不一致。")
    if str(validation.get("marketplace", "")).upper() != job["marketplace"]:
        fail("PPC 验证结果的站点与任务不一致。")
    if keyword_pool_sha and validation.get("keyword_pool_sha256") != keyword_pool_sha:
        fail("PPC 验证结果未绑定当前统一关键词池。")
    plan_path = Path(str(validation.get("plan", ""))).expanduser().resolve()
    markdown_path = Path(str(validation.get("markdown", ""))).expanduser().resolve()
    if not plan_path.is_file() or validation.get("plan_sha256") != sha256(plan_path):
        fail("PPC JSON 与验证时内容不一致。")
    require_text_file(markdown_path, "PPC 广告方案文本文件")
    if validation.get("markdown_sha256") != sha256(markdown_path):
        fail("PPC Markdown 与验证时内容不一致。")
    return validation


def build_review_markdown(packet: dict[str, Any], contents: dict[str, str]) -> str:
    deliverables = packet["pre_confirmation_deliverables"]
    lines = [
        "# 确认执行前审核包", "",
        "> 必须完整展示以下五个文件的绝对路径和正文后，才能请求用户回复“确认执行”。", "",
        "## 预确认文件", "",
    ]
    rows = [
        {"文件": item["label"], "绝对路径": item["path"], "SHA-256": item["sha256"]}
        for item in deliverables
    ]
    lines.extend(markdown_table(rows, ["文件", "绝对路径", "SHA-256"]))
    for item in deliverables:
        lines.extend(["", f"## {item['label']}", "", f"绝对路径：`{item['path']}`", "", contents[item["kind"]].rstrip(), ""])
    lines.extend(["## 确认提示", "", "请检查以上五个文件的全部内容。确认无误后回复：`确认执行`。", ""])
    return "\n".join(lines)


def build_review_packet(args: argparse.Namespace) -> dict[str, Any]:
    job_path = Path(args.job).expanduser().resolve()
    plan_path = Path(args.plan).expanduser().resolve()
    job = load_json(job_path)
    plan = load_json(plan_path)
    if plan.get("status") != "ready":
        fail("同步计划不是 ready，不能生成确认审核包。")
    keyword_validation_path = Path(args.keyword_validation).expanduser().resolve()
    keyword_validation, keyword_pool = load_keyword_validation(job, keyword_validation_path)
    keyword_pool_sha = keyword_validation["pool_sha256"]
    ppc_validation_path = Path(args.ppc_validation).expanduser().resolve()
    ppc_validation = verify_ppc_validation(job, ppc_validation_path, keyword_pool_sha)
    report_paths = [Path(item).expanduser().resolve() for item in args.reports]
    report_validation_path = Path(args.report_validation).expanduser().resolve()
    verify_report_validation(job, report_validation_path, report_paths, keyword_pool_sha)
    reports: dict[str, tuple[Path, dict[str, Any]]] = {}
    for path in report_paths:
        report = load_json(path)
        asin = normalize_asin(report.get("asin"), f"报告 {path.name} 的 ASIN")
        if asin in reports:
            fail(f"ASIN {asin} 存在重复报告。")
        if report.get("keyword_pool_sha256") != keyword_pool_sha:
            fail(f"{asin} 优化报告未绑定当前统一关键词池。")
        reports[asin] = (path, report)
    actions = {item.get("asin"): item for item in plan.get("actions", []) if item.get("action") == "update"}
    expected = set(job["target_asins"])
    if set(reports) != expected:
        fail("优化报告 ASIN 集合与任务不一致。")
    if set(actions) != expected:
        fail("同步计划 ASIN 集合与任务不一致。")

    listings = []
    for asin in job["target_asins"]:
        report_path, report = reports[asin]
        action = actions[asin]
        missing_audit = [section for section in REQUIRED_AUDIT_SECTIONS if report.get(section) in (None, "", [], {})]
        audit = report.get("audit") or {}
        if not audit.get("dimensions"):
            missing_audit.append("audit.dimensions")
        if missing_audit:
            fail(f"{asin} 缺少完整审核报告字段：{', '.join(missing_audit)}")
        fields = action.get("fields") or {}
        bullets = list(fields.get("bullets") or [])[:5]
        if len(bullets) != 5:
            fail(f"{asin} 的同步计划未包含完整 5 条五点。")
        legacy_generation = report.get("legacy_generation") or {}
        legacy_listing = legacy_generation.get("legacy_listing") or report.get("listing") or {}
        if report.get("listing") != legacy_listing:
            fail(f"{asin} 顶层传统 Listing 与 amazon-listing-optimization 交接内容不一致。")
        legacy_bullets = legacy_listing.get("bullets") or []
        if len(legacy_bullets) != 5:
            fail(f"{asin} 的传统 Listing 未包含完整 5 条五点。")
        option_one = next(
            (
                option for option in (report.get("title_options_2026") or [])
                if isinstance(option, dict) and option.get("option") == 1
            ),
            None,
        )
        if not option_one:
            fail(f"{asin} 缺少 2026 标题方案 1，无法核对同步计划。")
        sync_bindings = (
            ("title", fields.get("title"), option_one.get("title")),
            ("item_highlight", fields.get("item_highlight"), option_one.get("item_highlights")),
            ("bullets", bullets, legacy_bullets),
            ("description", fields.get("description"), legacy_listing.get("description")),
            ("search_terms", fields.get("search_terms"), legacy_listing.get("backend_search_terms")),
        )
        mismatched_sync_fields = [
            field for field, actual, expected_value in sync_bindings
            if not listing_values_equal(actual, expected_value)
        ]
        if mismatched_sync_fields:
            fail(
                f"{asin} 同步计划字段未绑定传统 Listing 或 2026 方案 1："
                + "、".join(mismatched_sync_fields)
            )
        listings.append({
            "asin": asin,
            "report": str(report_path),
            "legacy_listing": {
                "title": str(legacy_listing.get("title", "")),
                "bullets": list(legacy_bullets),
                "description": str(legacy_listing.get("description", "")),
                "backend_search_terms": str(legacy_listing.get("backend_search_terms", "")),
            },
            "pending_listing": {
                "title": fields.get("title", ""),
                "item_name": fields.get("title", ""),
                "item_highlight": fields.get("item_highlight", ""),
                "bullets": bullets,
                "description": fields.get("description", ""),
                "backend_search_terms": fields.get("search_terms", ""),
            },
            "destination": {
                "workbook": action.get("workbook", ""), "sheet": action.get("sheet", ""),
                "rows": action.get("rows", []), "title_mode": fields.get("title_mode", ""),
                "station_evidence": action.get("station_evidence", []),
                "planned_updates": action.get("planned_updates", []),
            },
            "audit_report": {
                **{section: report.get(section) for section in AUDIT_SECTIONS},
                "legacy_generation": legacy_generation,
                "title_options_generation": report.get("title_options_generation"),
                "title_options_2026": report.get("title_options_2026"),
            },
        })
    packet = {
        "schema_version": 1,
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "status": "ready_for_user_review",
        "job": str(job_path), "plan": str(plan_path), "marketplace": job["marketplace"],
        "keyword_validation": str(keyword_validation_path),
        "keyword_pool": keyword_validation["pool"],
        "keyword_pool_sha256": keyword_pool_sha,
        "ppc_plan": ppc_validation["plan"],
        "ppc_validation": str(ppc_validation_path),
        "report_validation": str(report_validation_path),
        "variants": job["variants"], "listings": listings,
    }
    output = Path(args.output).expanduser().resolve()
    requested_markdown = Path(args.markdown).expanduser().resolve()
    markdown = requested_markdown.parent / confirmation_review_filename(job)
    # Keep programmatic callers aligned with the generated, audit-safe filename.
    args.markdown = str(markdown)
    listing_markdown = Path(args.listing_markdown).expanduser().resolve()
    short_title_markdown = Path(args.short_title_markdown).expanduser().resolve()
    xlsm_markdown = Path(args.xlsm_markdown).expanduser().resolve()
    rufus_markdown = Path(args.rufus_markdown).expanduser().resolve()
    rufus_validation_path = Path(args.rufus_validation).expanduser().resolve()
    ppc_markdown = Path(ppc_validation["markdown"]).expanduser().resolve()
    expected_names = {
        listing_markdown: f"listing-optimization-plan-{job['primary_asin']}.md",
        short_title_markdown: f"short-title-highlights-plan-{job['primary_asin']}.md",
        ppc_markdown: f"ppc-campaign-plan-{job['primary_asin']}.md",
        xlsm_markdown: "xlsm-replacement-plan.md",
        rufus_markdown: f"rufus-qa-plan-{job['primary_asin']}.md",
    }
    if len(expected_names) != 5:
        fail("五个预确认文本文件必须使用不同路径。")
    for path, expected_name in expected_names.items():
        if path.name != expected_name:
            fail(f"预确认文本文件名必须为 {expected_name}：{path}")
    xlsm_markdown.parent.mkdir(parents=True, exist_ok=True)
    xlsm_markdown.write_text(build_xlsm_replacement_markdown(packet), encoding="utf-8")
    listing_markdown.parent.mkdir(parents=True, exist_ok=True)
    listing_markdown.write_text(build_listing_plan_markdown(packet), encoding="utf-8")
    short_title_markdown.parent.mkdir(parents=True, exist_ok=True)
    short_title_markdown.write_text(build_short_title_markdown(packet), encoding="utf-8")
    primary_report = reports[job["primary_asin"]][1]
    primary_report_path = reports[job["primary_asin"]][0]
    verify_rufus_validation(
        job, rufus_validation_path, primary_report_path,
        rufus_markdown, keyword_pool_sha,
    )
    contents = {
        "listing_optimization": verify_listing_preview(packet, listing_markdown),
        "short_title_highlights": verify_short_title_preview(packet, short_title_markdown),
        "ppc_campaign": require_text_file(ppc_markdown, "PPC 广告方案文本文件"),
        "xlsm_replacement": require_text_file(xlsm_markdown, "XLSM 文件替换方案文本文件"),
        "rufus_qa": verify_rufus_preview(
            primary_report, job, keyword_pool, keyword_pool_sha, rufus_markdown
        ),
    }
    packet["pre_confirmation_deliverables"] = [
        {
            "kind": "listing_optimization", "label": "Listing 优化方案文本文件",
            "path": str(listing_markdown), "sha256": sha256(listing_markdown),
        },
        {
            "kind": "short_title_highlights", "label": "2026 短标题与商品亮点方案文本文件",
            "path": str(short_title_markdown), "sha256": sha256(short_title_markdown),
        },
        {
            "kind": "ppc_campaign", "label": "PPC 广告方案文本文件",
            "path": str(ppc_markdown), "sha256": sha256(ppc_markdown),
        },
        {
            "kind": "xlsm_replacement", "label": "XLSM 文件替换方案文本文件",
            "path": str(xlsm_markdown), "sha256": sha256(xlsm_markdown),
        },
        {
            "kind": "rufus_qa", "label": "Rufus Q/A 方案文本文件",
            "path": str(rufus_markdown), "sha256": sha256(rufus_markdown),
        },
    ]
    packet["rufus_validation"] = str(rufus_validation_path)
    write_json(output, packet)
    markdown.parent.mkdir(parents=True, exist_ok=True)
    markdown.write_text(build_review_markdown(packet, contents), encoding="utf-8")
    return {
        "ok": True,
        "review_packet": str(output),
        "markdown": str(markdown),
        "pre_confirmation_files": [item["path"] for item in packet["pre_confirmation_deliverables"]],
        "asins": job["target_asins"],
    }


def verify_sources(job: dict[str, Any]) -> None:
    for item in job.get("source_files", []):
        path = Path(item["path"])
        if not path.is_file() or sha256(path) != item["sha256"]:
            fail(f"输入文件在预检后发生变化：{path}")


def seal_plan(args: argparse.Namespace) -> dict[str, Any]:
    job_path = Path(args.job).expanduser().resolve()
    plan_path = Path(args.plan).expanduser().resolve()
    job = load_json(job_path)
    plan = load_json(plan_path)
    verify_sources(job)
    if plan.get("status") != "ready" or not plan.get("actions"):
        fail("同步计划不是 ready，不能请求确认。")
    plan_asins = {item.get("asin") for item in plan["actions"] if item.get("action") == "update"}
    if plan_asins != set(job["target_asins"]):
        fail("同步计划 ASIN 集合与任务不一致。")
    review_path = Path(args.review).expanduser().resolve()
    review_markdown_path = Path(args.review_markdown).expanduser().resolve()
    review = load_json(review_path)
    if review.get("status") != "ready_for_user_review":
        fail("确认审核包状态无效，不能请求确认。")
    if Path(str(review.get("job", ""))).expanduser().resolve() != job_path:
        fail("确认审核包与当前 job.json 不一致。")
    if Path(str(review.get("plan", ""))).expanduser().resolve() != plan_path:
        fail("确认审核包与当前同步计划不一致。")
    keyword_validation_path = Path(str(review.get("keyword_validation", ""))).expanduser().resolve()
    keyword_validation, _keyword_pool = load_keyword_validation(job, keyword_validation_path)
    keyword_pool_path = Path(keyword_validation["pool"]).expanduser().resolve()
    if review.get("keyword_pool_sha256") != keyword_validation["pool_sha256"]:
        fail("确认审核包未绑定当前统一关键词池。")
    if Path(str(review.get("keyword_pool", ""))).expanduser().resolve() != keyword_pool_path:
        fail("确认审核包引用的统一关键词池路径不一致。")
    review_listing_items = [item for item in review.get("listings", []) if isinstance(item, dict)]
    review_report_paths = [Path(str(item.get("report", ""))).expanduser().resolve() for item in review_listing_items]
    report_validation_path = Path(str(review.get("report_validation", ""))).expanduser().resolve()
    verify_report_validation(
        job, report_validation_path, review_report_paths,
        keyword_validation["pool_sha256"],
    )
    ppc_validation_path = Path(str(review.get("ppc_validation", ""))).expanduser().resolve()
    ppc_validation = verify_ppc_validation(job, ppc_validation_path, keyword_validation["pool_sha256"])
    ppc_plan_path = Path(ppc_validation["plan"]).expanduser().resolve()
    if Path(str(review.get("ppc_plan", ""))).expanduser().resolve() != ppc_plan_path:
        fail("确认审核包引用的 PPC JSON 与验证结果不一致。")
    deliverables = review.get("pre_confirmation_deliverables")
    if not isinstance(deliverables, list):
        fail("确认审核包缺少五个预确认文本文件。")
    by_kind = {item.get("kind"): item for item in deliverables if isinstance(item, dict)}
    required_kinds = {
        "listing_optimization", "short_title_highlights", "ppc_campaign",
        "xlsm_replacement", "rufus_qa",
    }
    if set(by_kind) != required_kinds:
        fail("确认审核包必须且只能包含 Listing、短标题、PPC、XLSM 和 Rufus Q/A 五个预确认文本文件。")
    expected_names = {
        "listing_optimization": f"listing-optimization-plan-{job['primary_asin']}.md",
        "short_title_highlights": f"short-title-highlights-plan-{job['primary_asin']}.md",
        "ppc_campaign": f"ppc-campaign-plan-{job['primary_asin']}.md",
        "xlsm_replacement": "xlsm-replacement-plan.md",
        "rufus_qa": f"rufus-qa-plan-{job['primary_asin']}.md",
    }
    validated_ppc_markdown = Path(ppc_validation["markdown"]).expanduser().resolve()
    rufus_validation_path = Path(str(review.get("rufus_validation", ""))).expanduser().resolve()
    primary_review = next(
        (item for item in review.get("listings", []) if isinstance(item, dict) and item.get("asin") == job["primary_asin"]),
        None,
    )
    if not primary_review:
        fail("确认审核包缺少主 ASIN Listing 报告。")
    primary_report_path = Path(str(primary_review.get("report", ""))).expanduser().resolve()
    rufus_markdown_path = Path(str(by_kind["rufus_qa"].get("path", ""))).expanduser().resolve()
    verify_rufus_validation(
        job, rufus_validation_path, primary_report_path,
        rufus_markdown_path, keyword_validation["pool_sha256"],
    )
    review_markdown = require_text_file(review_markdown_path, "确认执行前综合审核包")
    sealed_artifacts: dict[str, dict[str, str]] = {}
    for kind in sorted(required_kinds):
        item = by_kind[kind]
        artifact_path = Path(str(item.get("path", ""))).expanduser().resolve()
        if artifact_path.name != expected_names[kind]:
            fail(f"预确认文本文件名必须为 {expected_names[kind]}：{artifact_path}")
        if kind == "ppc_campaign" and artifact_path != validated_ppc_markdown:
            fail("确认审核包中的 PPC 文本文件不是已验证的 PPC Markdown。")
        content = require_text_file(artifact_path, str(item.get("label") or kind))
        digest = sha256(artifact_path)
        if item.get("sha256") != digest:
            fail(f"预确认文本文件在审核包生成后发生变化：{artifact_path}")
        if str(artifact_path) not in review_markdown or content.rstrip() not in review_markdown:
            fail(f"综合审核包未完整展示预确认文件路径和正文：{artifact_path}")
        sealed_artifacts[kind] = {"path": str(artifact_path), "sha256": digest}
    sealed_artifacts.update({
        "job": {"path": str(job_path), "sha256": sha256(job_path)},
        "sync_plan": {"path": str(plan_path), "sha256": sha256(plan_path)},
        "review_packet": {"path": str(review_path), "sha256": sha256(review_path)},
        "review_markdown": {"path": str(review_markdown_path), "sha256": sha256(review_markdown_path)},
        "ppc_plan_json": {"path": str(ppc_plan_path), "sha256": sha256(ppc_plan_path)},
        "ppc_validation": {"path": str(ppc_validation_path), "sha256": sha256(ppc_validation_path)},
        "report_validation": {"path": str(report_validation_path), "sha256": sha256(report_validation_path)},
        "rufus_validation": {"path": str(rufus_validation_path), "sha256": sha256(rufus_validation_path)},
        "keyword_pool": {"path": str(keyword_pool_path), "sha256": sha256(keyword_pool_path)},
        "keyword_validation": {"path": str(keyword_validation_path), "sha256": sha256(keyword_validation_path)},
    })
    review_listings = review.get("listings")
    if not isinstance(review_listings, list) or not review_listings:
        fail("确认审核包缺少全部目标 ASIN 的 Listing 报告引用。")
    review_asins = {
        normalize_asin(item.get("asin"), "确认审核包 Listing ASIN")
        for item in review_listings if isinstance(item, dict)
    }
    if review_asins != set(job["target_asins"]):
        fail("确认审核包中的 Listing 报告 ASIN 集合与任务不一致。")
    for item in review_listings:
        if not isinstance(item, dict):
            fail("确认审核包中的 Listing 报告引用无效。")
        asin = normalize_asin(item.get("asin"), "确认审核包 Listing ASIN")
        report_path = Path(str(item.get("report", ""))).expanduser().resolve()
        if not report_path.is_file():
            fail(f"确认审核包引用的 Listing 报告不存在：{report_path}")
        sealed_artifacts[f"listing_report_{asin}"] = {
            "path": str(report_path), "sha256": sha256(report_path),
        }
        report = load_json(report_path)
        legacy_path = Path(
            str((report.get("legacy_generation") or {}).get("artifact_path", ""))
        ).expanduser().resolve()
        if not legacy_path.is_file():
            fail(f"确认审核包引用的传统 Listing 交接文件不存在：{legacy_path}")
        sealed_artifacts[f"legacy_listing_{asin}"] = {
            "path": str(legacy_path), "sha256": sha256(legacy_path),
        }
        legacy = report.get("legacy_generation") or {}
        for suffix, path_key, hash_key in (
            ("request", "request_path", "request_sha256"),
            ("raw_output", "raw_output_path", "raw_output_sha256"),
        ):
            evidence_path = Path(str(legacy.get(path_key, ""))).expanduser().resolve()
            if not evidence_path.is_file() or sha256(evidence_path) != legacy.get(hash_key):
                fail(f"确认审核包引用的传统 Listing {suffix} 证据无效：{asin}")
            sealed_artifacts[f"legacy_{suffix}_{asin}"] = {
                "path": str(evidence_path), "sha256": sha256(evidence_path),
            }
    request = {
        "schema_version": 1,
        "status": "awaiting_confirmation",
        "created_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "job": str(job_path),
        "job_sha256": sha256(job_path),
        "plan": str(plan_path),
        "plan_sha256": sha256(plan_path),
        "review_packet": str(review_path),
        "review_markdown": str(review_markdown_path),
        "sealed_artifacts": sealed_artifacts,
    }
    output = Path(args.output).expanduser().resolve()
    write_json(output, request)
    return {"ok": True, "confirmation_request": str(output), "plan_sha256": request["plan_sha256"]}


def current_hashes(request: dict[str, Any]) -> None:
    job_path = Path(request["job"])
    plan_path = Path(request["plan"])
    if sha256(job_path) != request["job_sha256"]:
        fail("job.json 在确认前发生变化，原计划已失效。")
    if sha256(plan_path) != request["plan_sha256"]:
        fail("同步计划在确认前发生变化，原计划已失效。")
    sealed_artifacts = request.get("sealed_artifacts")
    if not isinstance(sealed_artifacts, dict) or not sealed_artifacts:
        fail("确认请求缺少预确认文件封存信息，原计划已失效。")
    for name, item in sealed_artifacts.items():
        if not isinstance(item, dict):
            fail(f"确认请求中的封存信息无效：{name}")
        path = Path(str(item.get("path", ""))).expanduser().resolve()
        if not path.is_file() or sha256(path) != item.get("sha256"):
            fail(f"预确认封存文件发生变化，原确认已失效：{path}")
    verify_sources(load_json(job_path))


def confirm(args: argparse.Namespace) -> dict[str, Any]:
    request_path = Path(args.request).expanduser().resolve()
    request = load_json(request_path)
    if args.phrase.strip() not in CONFIRM_PHRASES:
        fail("未收到有效的明确执行确认。")
    current_hashes(request)
    confirmation = {
        **request,
        "status": "confirmed",
        "confirmed_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "confirmation_phrase": args.phrase.strip(),
        "request": str(request_path),
        "request_sha256": sha256(request_path),
    }
    output = Path(args.output).expanduser().resolve()
    write_json(output, confirmation)
    return {"ok": True, "confirmation": str(output), "plan": request["plan"]}


def verify_confirmation_file(path: Path) -> dict[str, Any]:
    confirmation = load_json(path)
    if confirmation.get("status") != "confirmed":
        fail("确认文件状态不是 confirmed。")
    request_path = Path(confirmation["request"])
    if sha256(request_path) != confirmation["request_sha256"]:
        fail("确认请求文件发生变化。")
    request = load_json(request_path)
    current_hashes(request)
    if confirmation.get("plan_sha256") != request.get("plan_sha256"):
        fail("确认文件与当前计划不一致。")
    return confirmation


def verify_confirmation(args: argparse.Namespace) -> dict[str, Any]:
    confirmation = verify_confirmation_file(Path(args.confirmation).expanduser().resolve())
    return {"ok": True, "plan": confirmation["plan"], "plan_sha256": confirmation["plan_sha256"]}


def load_sync_module():
    path = Path("/Users/apple/.codex/skills/kjxj-sync-product-listing/scripts/listing_sync.py")
    spec = importlib.util.spec_from_file_location("listing_sync", path)
    if not spec or not spec.loader:
        fail(f"无法加载同步脚本：{path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def apply_sync(args: argparse.Namespace) -> dict[str, Any]:
    confirmation = verify_confirmation_file(Path(args.confirmation).expanduser().resolve())
    plan = load_json(Path(confirmation["plan"]))
    sync = load_sync_module()
    results = sync.apply_plan(plan, Path(args.output_dir).expanduser().resolve())
    payload = {
        "ok": True,
        "applied_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "plan": confirmation["plan"],
        "plan_sha256": confirmation["plan_sha256"],
        "confirmation": str(Path(args.confirmation).expanduser().resolve()),
        "results": results,
    }
    output = Path(args.result).expanduser().resolve()
    write_json(output, payload)
    return {"ok": True, "result": str(output), "outputs": [item["output"] for item in results]}


def green(cell) -> bool:
    color = cell.fill.fgColor
    return color.type == "rgb" and str(color.rgb).upper().endswith("C6EFCE")


def verify_output(args: argparse.Namespace) -> dict[str, Any]:
    confirmation = verify_confirmation_file(Path(args.confirmation).expanduser().resolve())
    plan = load_json(Path(confirmation["plan"]))
    raw_result = load_json(Path(args.result).expanduser().resolve())
    results = raw_result.get("results", raw_result if isinstance(raw_result, list) else None)
    if not isinstance(results, list) or not results:
        fail("apply 结果中没有可验证的 results。")
    sync = load_sync_module()
    errors: list[str] = []
    checked: list[dict[str, Any]] = []
    result_by_source = {str(Path(item["source"]).resolve()): item for item in results}
    actions_by_source: dict[str, list[dict[str, Any]]] = {}
    for action in plan["actions"]:
        actions_by_source.setdefault(str(Path(action["workbook"]).resolve()), []).append(action)
    for source_name, actions in actions_by_source.items():
        result = result_by_source.get(source_name)
        if not result:
            errors.append(f"缺少源文件的 apply 结果：{source_name}")
            continue
        source = Path(source_name)
        output = Path(result["output"]).resolve()
        if output.suffix.lower() != ".xlsm" or not output.is_file():
            errors.append(f"输出不是有效的 XLSM 文件：{output}")
            continue
        try:
            with zipfile.ZipFile(output) as archive:
                if archive.testzip() is not None:
                    errors.append(f"OOXML 压缩包损坏：{output}")
                output_parts = set(archive.namelist())
            with zipfile.ZipFile(source) as archive:
                source_parts = set(archive.namelist())
                source_vba = archive.read("xl/vbaProject.bin") if "xl/vbaProject.bin" in source_parts else None
            if output_parts != source_parts:
                errors.append(f"输出 OOXML 部件列表发生非预期变化：{output}")
            if source_vba is not None:
                with zipfile.ZipFile(output) as archive:
                    if archive.read("xl/vbaProject.bin") != source_vba:
                        errors.append(f"宏部件发生变化：{output}")
            workbook = load_workbook(output, data_only=False, keep_vba=True)
            for action in actions:
                ws = workbook[action["sheet"]]
                fields = sync.header_map(ws)
                expected_values = {
                    fields["title"][0]: action["fields"]["title"],
                    fields["item_name"][0]: action["fields"]["title"],
                    fields["description"][0]: "".join(f"<p>{p}</p>" for p in action["fields"]["description"].split("\n\n")),
                    fields["search_terms"][0]: action["fields"]["search_terms"],
                }
                if action["fields"].get("item_highlight") and fields.get("item_highlight"):
                    expected_values[fields["item_highlight"][0]] = action["fields"]["item_highlight"]
                expected_values.update({column: bullet for column, bullet in zip(fields["bullet"][:5], action["fields"]["bullets"][:5])})
                skipped = set(action.get("skip_fields", []))
                skipped_columns = {column for field, columns in {"title": fields["title"][:1], "item_name": fields["item_name"][:1], "item_highlight": fields.get("item_highlight", [])[:1], "description": fields["description"][:1], "bullet": fields["bullet"][:5], "search_terms": fields["search_terms"][:1]}.items() if field in skipped for column in columns}
                # apply_plan compacts retained rows from row 7 onward. Resolve the
                # target ASIN again in the output instead of using source row IDs.
                product_id_columns = fields.get("product_id", [])[:1]
                output_rows = [
                    row for row in range(7, ws.max_row + 1)
                    if any(str(ws.cell(row, column).value or "").strip().upper() == action["asin"].upper()
                           for column in product_id_columns)
                ]
                if len(output_rows) != len(action["rows"]):
                    errors.append(
                        f"{action['asin']} 在输出中匹配到 {len(output_rows)} 行，预期 {len(action['rows'])} 行。"
                    )
                for row in output_rows:
                    for column, expected_value in expected_values.items():
                        if column in skipped_columns:
                            continue
                        cell = ws.cell(row, column)
                        if cell.value != expected_value:
                            errors.append(f"{action['asin']} 的 {cell.coordinate} 内容不一致。")
                        if not green(cell):
                            errors.append(f"{action['asin']} 的 {cell.coordinate} 未标浅绿色。")
            checked.append({"source": str(source), "output": str(output), "asins": result.get("asins", []), "rows": result.get("rows", 0), "written_cells": result.get("written_cells", 0), "retained_parent_rows": result.get("retained_parent_rows", 0), "deleted_data_rows": result.get("deleted_data_rows", 0)})
        except Exception as exc:
            errors.append(f"验证输出失败 {output}：{exc}")
    validation = {"ok": not errors, "checked": checked, "errors": errors}
    if args.output:
        write_json(Path(args.output).expanduser().resolve(), validation)
    if errors:
        fail("交付验证失败：" + "；".join(errors))
    return validation


def ppc_financial_source(ppc: dict[str, Any], field: str) -> dict[str, Any]:
    defaults = ppc.get("defaults_applied") or {}
    default = defaults.get(field) if isinstance(defaults.get(field), dict) else {}
    if default.get("used") is True:
        return {
            "value": ppc.get(field), "source": "default",
            "reason": str(default.get("reason") or "按编排流程使用默认财务参数"),
        }
    if ppc.get(field) is None:
        return {
            "value": None, "source": "not_provided",
            "reason": "用户未提供该财务参数，流程未用假设值补写。",
        }
    return {
        "value": ppc.get(field), "source": "user_provided",
        "reason": "用户在任务输入中提供的财务参数",
    }


def ppc_campaign_label(recommended: dict[str, Any]) -> str:
    """Build a non-strategic campaign identity from verified title components."""
    components = recommended.get("title_components")
    if not isinstance(components, dict):
        fail("PPC 数据包需要推荐 2026 标题的 title_components 作为商品身份来源。")
    brand = str(components.get("brand", "")).strip()
    core_product = str(components.get("core_product_phrase", "")).strip()
    title = str(recommended.get("title", "")).strip()
    if not brand or not core_product or not title:
        fail("PPC 数据包需要完整的品牌、核心品类词和推荐 2026 标题。")
    label = " ".join((brand, core_product))
    if not contains_term(title, brand) or not contains_term(title, core_product):
        fail("PPC Campaign 身份必须由推荐 2026 标题中可验证的品牌和核心品类词组成。")
    return label


def build_ppc_source_packet(args: argparse.Namespace) -> dict[str, Any]:
    """Create a fact-only Mode A handoff; never make PPC strategy decisions here."""
    job_path = Path(args.job).expanduser().resolve()
    job = load_json(job_path)
    keyword_validation_path = Path(args.keyword_validation).expanduser().resolve()
    keyword_validation, keyword_pool = load_keyword_validation(job, keyword_validation_path)
    report_path = Path(args.report).expanduser().resolve()
    report = load_json(report_path)
    if str(report.get("asin", "")).upper() != job["primary_asin"]:
        fail("PPC 数据包只能使用主 ASIN 的 Listing 报告。")
    if str(report.get("marketplace", "")).upper() != job["marketplace"]:
        fail("PPC 数据包的 Listing 报告站点与任务不一致。")
    legacy = report.get("legacy_generation") if isinstance(report.get("legacy_generation"), dict) else {}
    listing = legacy.get("legacy_listing") if isinstance(legacy.get("legacy_listing"), dict) else {}
    if not listing:
        fail("PPC 数据包需要 amazon-listing-optimization Mode B 的传统 Listing 输出。")
    title_options = report.get("title_options_2026") if isinstance(report.get("title_options_2026"), list) else []
    recommended = next(
        (item for item in title_options if isinstance(item, dict) and item.get("status") == "recommended_for_current_upload"),
        None,
    )
    if not recommended:
        fail("PPC 数据包需要 2026 推荐标题方案。")
    campaign_label = ppc_campaign_label(recommended)
    ppc = job.get("ppc_campaign") if isinstance(job.get("ppc_campaign"), dict) else {}
    keywords = []
    for item in keyword_pool.get("keywords", []):
        if not isinstance(item, dict) or item.get("ppc_eligible") is not True:
            continue
        keywords.append({
            "id": item.get("id"), "keyword": item.get("keyword"),
            "sources": item.get("sources", []), "metrics": item.get("metrics", {}),
            "selling_point_matches": item.get("selling_point_matches", []),
            "ppc_eligible": True,
        })
    if not keywords:
        fail("PPC 数据包没有 ppc_eligible 关键词。")
    financial = {
        key: ppc_financial_source(ppc, key)
        for key in (
            "selling_price", "landed_cost", "amazon_fees", "break_even_acos",
            "monthly_ad_budget", "conversion_rate", "product_stage",
        )
    }
    default_used = any(item["source"] == "default" for item in financial.values())
    packet = {
        "schema_version": SCHEMA_VERSION,
        "mode": "A",
        "primary_asin": job["primary_asin"],
        "marketplace": job["marketplace"],
        "currency": ppc.get("currency"),
        "data_quality": "assumption_limited" if default_used else "fact_supported",
        "financial": financial,
        # This is an immutable provenance boundary for the PPC executor.
        "financial_defaults_applied": ppc.get("defaults_applied", {}),
        "product_context": {
            "campaign_label": campaign_label,
            "campaign_label_sources": [
                "title_options_2026.recommended.title_components.brand",
                "title_options_2026.recommended.title_components.core_product_phrase",
            ],
            "legacy_listing": listing,
            "recommended_2026_title": recommended.get("title", ""),
            "recommended_2026_highlights": recommended.get("item_highlights", ""),
            "source_report": str(report_path), "source_report_sha256": sha256(report_path),
            "legacy_raw_output": legacy.get("raw_output_path", ""),
            "legacy_raw_output_sha256": legacy.get("raw_output_sha256", ""),
        },
        "keywords": keywords,
        "keyword_pool": {
            "path": str(Path(keyword_validation["pool"]).resolve()),
            "sha256": keyword_validation["pool_sha256"],
        },
        "competitor_asins": list(job.get("competitor_asins", [])),
        "historical_ad_data": {"included": False, "reason": "Mode A 建站流程不接收历史广告表现或搜索词报表"},
        "constraints": {
            "orchestrator_must_not_create_strategy": True,
            "no_unverified_conversion_rate": True,
            "no_unverified_suggested_bid": True,
        },
    }
    output = Path(args.output).expanduser().resolve()
    write_json(output, packet)
    return {"ok": True, "packet": str(output), "packet_sha256": sha256(output), "keyword_count": len(keywords)}


def ppc_packet_has_strategy(value: Any) -> bool:
    if isinstance(value, dict):
        return any(key in PPC_STRATEGY_KEYS or ppc_packet_has_strategy(item) for key, item in value.items())
    if isinstance(value, list):
        return any(ppc_packet_has_strategy(item) for item in value)
    return False


def validate_ppc_provenance(plan: dict[str, Any], args: argparse.Namespace, errors: list[str]) -> dict[str, Any] | None:
    """Verify that semantic PPC output was emitted by amazon-ppc-campaign, not this workflow."""
    packet_arg = getattr(args, "source_packet", None)
    request_arg = getattr(args, "request", None)
    raw_arg = getattr(args, "raw_output", None)
    if not any((packet_arg, request_arg, raw_arg)):
        return  # Direct function callers retain backwards-compatible fixture support.
    if not all((packet_arg, request_arg, raw_arg)):
        errors.append("PPC 来源校验必须同时提供 source_packet、request 和 raw_output。")
        return
    packet_path = Path(str(packet_arg)).expanduser().resolve()
    request_path = Path(str(request_arg)).expanduser().resolve()
    raw_path = Path(str(raw_arg)).expanduser().resolve()
    for label, path in (("PPC 数据包", packet_path), ("PPC 调用请求", request_path), ("PPC 原始输出", raw_path)):
        if not path.is_file():
            errors.append(f"{label}不存在：{path}")
    if errors:
        return
    try:
        packet = load_json(packet_path)
        request = load_json(request_path)
        raw = raw_path.read_text(encoding="utf-8")
    except (ValueError, OSError) as exc:
        errors.append(f"PPC 来源文件不可读取：{exc}")
        return
    if ppc_packet_has_strategy(packet):
        errors.append("PPC 数据包不得包含 Campaign、出价、否定词或关键词分组策略。")
    if packet.get("mode") != "A" or packet.get("historical_ad_data", {}).get("included") is not False:
        errors.append("PPC 数据包必须为不含历史广告数据的 Mode A 输入。")
    packet_defaults = packet.get("financial_defaults_applied")
    if not isinstance(packet_defaults, dict):
        errors.append("PPC 数据包 financial_defaults_applied 必须是对象。")
    else:
        for field in ("monthly_ad_budget", "break_even_acos", "selling_price"):
            entry = packet_defaults.get(field)
            if not isinstance(entry, dict) or not isinstance(entry.get("used"), bool):
                errors.append(f"PPC 数据包 financial_defaults_applied.{field} 必须保留 used 标记。")
    context = packet.get("product_context") if isinstance(packet.get("product_context"), dict) else {}
    campaign_label = str(context.get("campaign_label", "")).strip()
    if not campaign_label:
        errors.append("PPC 数据包缺少 product_context.campaign_label。")
    elif not contains_term(str(context.get("recommended_2026_title", "")), campaign_label):
        errors.append("PPC 数据包 campaign_label 必须完整出现在推荐 2026 标题中。")
    provenance = plan.get("execution_provenance") if isinstance(plan.get("execution_provenance"), dict) else {}
    adapter_path = Path(__file__).with_name("external_adapters") / "ppc_mode_a_adapter.py"
    if provenance.get("adapter") != "kjxj-optimize-sync-listing/scripts/external_adapters/ppc_mode_a_adapter.py":
        errors.append("PPC execution_provenance.adapter 不正确。")
    if not adapter_path.is_file() or provenance.get("adapter_sha256") != sha256(adapter_path):
        errors.append("PPC 适配器文件或 SHA-256 不一致。")
    upstream = provenance.get("upstream_skill") if isinstance(provenance.get("upstream_skill"), dict) else {}
    if upstream.get("name") != PPC_SOURCE_SKILL or upstream.get("path") != PPC_SOURCE_SKILL_PATH:
        errors.append("PPC execution_provenance.upstream_skill 不正确。")
    invocation_id = str(provenance.get("adapter_invocation_id", "")).strip()
    if not PPC_INVOCATION_ID_RE.fullmatch(invocation_id):
        errors.append("PPC adapter_invocation_id 必须是稳定的实际调用 ID。")
    for label, expected, actual in (
        ("数据包路径", str(packet_path), provenance.get("source_packet_path")),
        ("数据包 SHA-256", sha256(packet_path), provenance.get("source_packet_sha256")),
        ("原始输出路径", str(raw_path), (provenance.get("external_raw_output") or {}).get("path")),
        ("原始输出 SHA-256", sha256(raw_path), (provenance.get("external_raw_output") or {}).get("sha256")),
    ):
        if actual != expected:
            errors.append(f"PPC 生成来源的{label}不一致。")
    if request.get("adapter") != provenance.get("adapter"):
        errors.append("PPC 调用请求未绑定 KJXJ 适配器。")
    if request.get("mode") != "A" or request.get("adapter_invocation_id") != invocation_id:
        errors.append("PPC 调用请求模式或 invocation_id 不一致。")
    if request.get("source_packet_path") != str(packet_path) or request.get("source_packet_sha256") != sha256(packet_path):
        errors.append("PPC 调用请求未正确绑定数据包。")
    if f'"primary_asin":"{packet.get("primary_asin", "")}"' not in raw:
        errors.append("PPC 原始输出未包含与事实数据包一致的规范 JSON。")
    return packet


def ppc_copy_values(copy_blocks: dict[str, Any], key: str) -> list[str]:
    value = copy_blocks.get(key)
    if not isinstance(value, list):
        fail(f"PPC copy_blocks.{key} 必须是数组。")
    result: list[str] = []
    seen: set[str] = set()
    for item in value:
        text = str(item).strip()
        if not text:
            fail(f"PPC copy_blocks.{key} 包含空值。")
        if re.search(r"[\r\n|]", text) or re.match(r"^(?:[-*•]|\d+[.)])\s*", text):
            fail(f"PPC copy_blocks.{key} 必须每项只包含可复制的纯值：{text}")
        normalized = normalize_keyword(text)
        if normalized in seen:
            fail(f"PPC copy_blocks.{key} 包含重复值：{text}")
        seen.add(normalized)
        result.append(text)
    return result


PPC_NON_EXPLANATORY_KEYS = {
    "type", "name", "campaign", "campaign_name", "match_type", "currency", "mode",
    "financial_confidence",
    "asin", "keyword", "keywords", "target", "targets", "source_url", "url",
}


def validate_chinese_explanation(
    value: Any,
    field: str,
    errors: list[str],
    key: str = "",
) -> None:
    if key in PPC_NON_EXPLANATORY_KEYS or value is None:
        return
    if isinstance(value, dict):
        for child_key, child_value in value.items():
            validate_chinese_explanation(child_value, f"{field}.{child_key}", errors, child_key)
        return
    if isinstance(value, list):
        for index, item in enumerate(value, start=1):
            validate_chinese_explanation(item, f"{field}[{index}]", errors, key)
        return
    if isinstance(value, str):
        text = value.strip()
        if text and not text.startswith(("http://", "https://")) and not CHINESE_RE.search(text):
            errors.append(f"PPC 说明字段 {field} 必须使用中文：{text}")


def validate_ppc_markdown_language(markdown: str) -> list[str]:
    errors: list[str] = []
    in_code_block = False
    for line_number, raw_line in enumerate(markdown.splitlines(), start=1):
        stripped = raw_line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block or not stripped:
            continue
        content = re.sub(r"^(?:#{1,6}|>|[-*+] |\d+[.)] )\s*", "", stripped).strip()
        if content in PPC_COPY_BLOCK_LABELS:
            continue
        if stripped.startswith("|") or re.fullmatch(r"[-:| ]+", stripped):
            continue
        if content.startswith(("http://", "https://")):
            continue
        if re.search(r"[A-Za-z]", content) and not CHINESE_RE.search(content):
            errors.append(f"PPC Markdown 第 {line_number} 行的说明文字必须使用中文：{content}")
    return errors


def numeric_bid_paths(value: Any, path: str = "") -> list[str]:
    findings: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            child_path = f"{path}.{key}" if path else str(key)
            if key in PPC_AMOUNT_BID_KEYS and isinstance(item, (int, float)) and not isinstance(item, bool):
                findings.append(child_path)
            findings.extend(numeric_bid_paths(item, child_path))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            findings.extend(numeric_bid_paths(item, f"{path}[{index}]"))
    return findings


def validate_ppc_plan(args: argparse.Namespace) -> dict[str, Any]:
    job = load_json(Path(args.job).expanduser().resolve())
    keyword_validation_path = Path(args.keyword_validation).expanduser().resolve()
    keyword_validation, keyword_pool = load_keyword_validation(job, keyword_validation_path)
    keyword_pool_sha = keyword_validation["pool_sha256"]
    ppc_keywords = {
        normalize_keyword(item.get("keyword")) for item in keyword_pool.get("keywords", [])
        if isinstance(item, dict) and item.get("ppc_eligible") is True
    }
    plan_path = Path(args.plan).expanduser().resolve()
    markdown_path = Path(args.markdown).expanduser().resolve()
    plan = load_json(plan_path)
    errors: list[str] = []

    if plan.get("schema_version") != SCHEMA_VERSION:
        errors.append("schema_version 必须为 1。")
    if plan.get("mode") != "build":
        errors.append("PPC 方案必须使用 Mode A build。")
    if str(plan.get("primary_asin", "")).upper() != job["primary_asin"]:
        errors.append("PPC 方案 ASIN 必须是 job.json 的主 ASIN。")
    if str(plan.get("marketplace", "")).upper() != job["marketplace"]:
        errors.append("PPC 方案站点与 job.json 不一致。")
    if plan.get("keyword_pool_sha256") != keyword_pool_sha:
        errors.append("PPC 方案未绑定当前统一关键词池。")
    ppc_job = job.get("ppc_campaign") or {}
    if ppc_job.get("missing_user_inputs"):
        errors.append("job.json 仍缺少用户专属 PPC 参数：" + ", ".join(ppc_job["missing_user_inputs"]))
    if plan.get("currency") != ppc_job.get("currency"):
        errors.append("PPC 方案币种与 job.json 不一致。")
    for key in PPC_REQUIRED_PLAN_FIELDS:
        if not plan.get(key):
            errors.append(f"PPC 方案缺少完整字段：{key}")

    source_packet = validate_ppc_provenance(plan, args, errors)

    financial = plan.get("financial_framework") if isinstance(plan.get("financial_framework"), dict) else {}
    try:
        selling_price = optional_positive_number(financial.get("selling_price"), "PPC 售价")
        monthly_budget = optional_positive_number(financial.get("monthly_ad_budget"), "PPC 月度预算")
        profit_before_ads = optional_positive_number(financial.get("profit_before_ads"), "PPC 广告前利润")
        break_even_acos = optional_rate(financial.get("break_even_acos"), "PPC 盈亏平衡 ACoS")
        target_launch = optional_rate(financial.get("target_acos_launch"), "PPC 新品目标 ACoS")
        target_mature = optional_rate(financial.get("target_acos_mature"), "PPC 成熟期目标 ACoS")
        conversion_rate = optional_rate(financial.get("conversion_rate"), "PPC 转化率")
        max_cpc = optional_positive_number(financial.get("max_cpc"), "PPC Max CPC")
    except ValueError as exc:
        errors.append(str(exc))
        selling_price = monthly_budget = profit_before_ads = break_even_acos = None
        target_launch = target_mature = conversion_rate = max_cpc = None

    if selling_price is None:
        errors.append("PPC 方案缺少已验证的主 ASIN 售价。")
    if monthly_budget is None:
        errors.append("PPC 方案缺少月度广告预算。")
    elif ppc_job.get("monthly_ad_budget") and abs(monthly_budget - ppc_job["monthly_ad_budget"]) > 0.01:
        errors.append("PPC 方案月度预算与 job.json 不一致。")
    if ppc_job.get("selling_price") and selling_price and abs(selling_price - ppc_job["selling_price"]) > 0.01:
        errors.append("PPC 方案售价与 job.json 不一致。")
    if break_even_acos is None:
        errors.append("PPC 方案缺少盈亏平衡 ACoS。")
    if target_launch is None or target_mature is None:
        errors.append("PPC 方案必须同时给出新品期和成熟期目标 ACoS。")
    elif break_even_acos and (target_launch > break_even_acos or target_mature > break_even_acos):
        errors.append("目标 ACoS 不得高于盈亏平衡 ACoS。")
    if ppc_job.get("break_even_acos") and break_even_acos and abs(break_even_acos - ppc_job["break_even_acos"]) > 0.001:
        errors.append("PPC 方案盈亏平衡 ACoS 与 job.json 不一致。")
    if selling_price and ppc_job.get("landed_cost") and ppc_job.get("amazon_fees"):
        expected_profit = selling_price - ppc_job["landed_cost"] - ppc_job["amazon_fees"]
        if expected_profit <= 0:
            errors.append("用户成本数据导致广告前利润不为正，不能创建盈利型 PPC 财务框架。")
        else:
            if profit_before_ads is None or abs(profit_before_ads - expected_profit) > 0.02:
                errors.append("PPC 广告前利润计算与售价、到岸成本和 Amazon 费用不一致。")
            expected_break_even = expected_profit / selling_price
            if break_even_acos is None or abs(break_even_acos - expected_break_even) > 0.001:
                errors.append("PPC 盈亏平衡 ACoS 计算不正确。")
    if ppc_job.get("conversion_rate") is not None:
        if conversion_rate is None or abs(conversion_rate - ppc_job["conversion_rate"]) > 0.001:
            errors.append("PPC 转化率与用户输入不一致。")
    elif conversion_rate is not None and not financial.get("conversion_rate_source"):
        errors.append("非用户提供的转化率必须注明可核验的数据来源。")
    if conversion_rate is not None and max_cpc is None:
        errors.append("已有转化率时必须计算 Max CPC。")
    if not isinstance(financial.get("data_sources"), list) or not financial.get("data_sources"):
        errors.append("PPC 财务框架必须列出数据来源。")
    if not isinstance(financial.get("assumptions"), list):
        errors.append("PPC 财务框架 assumptions 必须是数组。")
    defaults_applied = ppc_job.get("defaults_applied") or {}
    default_used = any(
        isinstance(defaults_applied.get(key), dict) and defaults_applied[key].get("used") is True
        for key in ("monthly_ad_budget", "break_even_acos", "selling_price")
    )
    if default_used:
        if financial.get("defaults_applied") != defaults_applied:
            errors.append("PPC 财务框架必须原样记录 job.json 的 defaults_applied。")
        assumptions = financial.get("assumptions")
        if not isinstance(assumptions, list) or not any(str(item).strip() for item in assumptions):
            errors.append("PPC 使用默认预算或 ACoS 时必须在 assumptions 中说明。")
    if isinstance(source_packet, dict) and "financial_defaults_applied" in source_packet:
        if source_packet.get("financial_defaults_applied") != defaults_applied:
            errors.append("PPC 数据包的 financial_defaults_applied 与 job.json 不一致。")
        if financial.get("defaults_applied") != source_packet.get("financial_defaults_applied"):
            errors.append("PPC 财务框架必须原样记录 PPC 数据包的 financial_defaults_applied。")
        if financial.get("financial_inputs") != source_packet.get("financial"):
            errors.append("PPC 财务框架必须原样记录 PPC 数据包的财务字段来源。")

    data_quality = plan.get("data_quality") if isinstance(plan.get("data_quality"), dict) else {}
    expected_quality = "assumption_limited" if default_used else "fact_supported"
    if data_quality.get("financial_confidence") != expected_quality:
        errors.append(f"PPC data_quality.financial_confidence 必须为 {expected_quality}。")
    for key in ("input_data_quality", "financial_limitations", "keyword_boundary"):
        value = str(data_quality.get(key, "")).strip()
        if not value or not CHINESE_RE.search(value):
            errors.append(f"PPC data_quality.{key} 必须是中文说明。")

    bid_guidance = plan.get("bid_guidance") if isinstance(plan.get("bid_guidance"), dict) else {}
    if conversion_rate is None:
        if max_cpc is not None:
            errors.append("未提供可核验转化率时不得计算或填写 Max CPC。")
        if bid_guidance.get("mode") != "formula_and_coefficients":
            errors.append("未提供转化率时 bid_guidance.mode 必须为 formula_and_coefficients。")
        if not isinstance(bid_guidance.get("formula"), str) or not CHINESE_RE.search(str(bid_guidance.get("formula", ""))):
            errors.append("未提供转化率时 bid_guidance.formula 必须说明 Max CPC 计算条件。")
        if not isinstance(bid_guidance.get("coefficients"), dict) or not bid_guidance.get("coefficients"):
            errors.append("未提供转化率时 bid_guidance.coefficients 必须提供匹配类型系数。")
        numeric_bids = numeric_bid_paths(plan)
        if numeric_bids:
            errors.append("未提供转化率时不得输出金额 CPC/出价：" + ", ".join(numeric_bids))
    elif bid_guidance.get("mode") != "capped_amounts":
        errors.append("提供可核验转化率时 bid_guidance.mode 必须为 capped_amounts。")

    for field, value in (
        ("financial_framework.data_sources", financial.get("data_sources")),
        ("financial_framework.assumptions", financial.get("assumptions")),
        ("data_quality", data_quality),
        ("bid_guidance", bid_guidance),
        ("keyword_sources", plan.get("keyword_sources")),
        ("campaigns", plan.get("campaigns")),
        ("budget_summary", plan.get("budget_summary")),
        ("launch_schedule", plan.get("launch_schedule")),
        ("optimization_plan_4_weeks", plan.get("optimization_plan_4_weeks")),
        ("risk_notes", plan.get("risk_notes")),
    ):
        validate_chinese_explanation(value, field, errors)

    campaigns = plan.get("campaigns") if isinstance(plan.get("campaigns"), list) else []
    campaign_types = {item.get("type") for item in campaigns if isinstance(item, dict)}
    for required in {"auto", "manual_exact", "manual_broad"}:
        if required not in campaign_types:
            errors.append(f"PPC 方案缺少 {required} Campaign。")
    if isinstance(source_packet, dict):
        context = source_packet.get("product_context") if isinstance(source_packet.get("product_context"), dict) else {}
        campaign_label = str(context.get("campaign_label", "")).strip()
        if not campaign_label:
            errors.append("PPC 数据包缺少 product_context.campaign_label，不能校验 Campaign 身份一致性。")
        else:
            identity_text = " ".join([
                campaign_label,
                str(context.get("recommended_2026_title", "")),
                str((context.get("legacy_listing") or {}).get("title", "")),
                " ".join(str(item.get("keyword", "")) for item in source_packet.get("keywords", []) if isinstance(item, dict)),
                " ".join(str(item) for item in source_packet.get("competitor_asins", [])),
            ])
            allowed_tokens = set(re.findall(r"[^\W\d_]+|\d+", identity_text.casefold(), re.UNICODE))
            allowed_tokens.update({"auto", "manual", "exact", "broad", "product", "targeting", "asin"})
            for index, campaign in enumerate(campaigns, start=1):
                if not isinstance(campaign, dict):
                    continue
                name = str(campaign.get("name", "")).strip()
                if not contains_term(name, campaign_label):
                    errors.append(f"PPC Campaign {index} 名称必须包含数据包商品身份：{campaign_label}。")
                    continue
                unknown_tokens = sorted({
                    token for token in re.findall(r"[^\W\d_]+|\d+", name.casefold(), re.UNICODE)
                    if token not in allowed_tokens
                })
                if unknown_tokens:
                    errors.append(
                        f"PPC Campaign {index} 名称包含数据包商品身份外的词：{', '.join(unknown_tokens)}。"
                    )

    copy_blocks = plan.get("copy_blocks") if isinstance(plan.get("copy_blocks"), dict) else {}
    copy_values: dict[str, list[str]] = {}
    for key in PPC_COPY_BLOCKS:
        try:
            copy_values[key] = ppc_copy_values(copy_blocks, key)
        except ValueError as exc:
            errors.append(str(exc))
            copy_values[key] = []
    exact = {normalize_keyword(item) for item in copy_values["manual_exact_keywords"]}
    broad = {normalize_keyword(item) for item in copy_values["manual_broad_keywords"]}
    auto_negatives = {normalize_keyword(item) for item in copy_values["auto_negative_exact_keywords"]}
    broad_negatives = {normalize_keyword(item) for item in copy_values["broad_negative_exact_keywords"]}
    phrase_negatives = {normalize_keyword(item) for item in copy_values["negative_phrase_keywords"]}
    if not exact:
        errors.append("Manual Exact Keywords 不能为空。")
    if not broad:
        errors.append("Manual Broad Keywords 不能为空。")
    outside_pool = sorted((exact | broad) - ppc_keywords)
    if outside_pool:
        errors.append("PPC Exact/Broad 包含统一关键词池 ppc_eligible 集合外的词：" + ", ".join(outside_pool))
    autocomplete_used = any(
        isinstance(item, dict) and "amazon_autocomplete" in (item.get("sources") or [])
        for item in keyword_pool.get("keywords", [])
    )
    keyword_sources = plan.get("keyword_sources") if isinstance(plan.get("keyword_sources"), list) else []
    if autocomplete_used and not any("Amazon autocomplete" in str(item) for item in keyword_sources):
        errors.append("统一关键词池使用自动补词时，PPC keyword_sources 必须记录 Amazon autocomplete。")
    if not exact.issubset(auto_negatives):
        errors.append("所有 Exact 关键词都必须加入 Auto Negative Exact。")
    if not exact.issubset(broad_negatives):
        errors.append("所有 Exact 关键词都必须加入 Broad Negative Exact。")
    if (exact | broad) & phrase_negatives:
        errors.append("Negative Phrase Keywords 与投放关键词存在直接冲突。")

    competitors = set(job.get("competitor_asins", []))
    try:
        product_targets = {normalize_asin(item, "PPC Product Target") for item in copy_values["product_targeting_asins"]}
    except ValueError as exc:
        errors.append(str(exc))
        product_targets = set()
    if not product_targets.issubset(competitors):
        errors.append("Product Targeting ASIN 必须来自本任务竞品集合。")
    if competitors and ("product_targeting" not in campaign_types or not product_targets):
        errors.append("任务含竞品时必须生成 Product Targeting Campaign 和复制区。")
    if not competitors and product_targets:
        errors.append("任务无竞品时不得生成 Product Targeting ASIN。")

    serialized = json.dumps(plan, ensure_ascii=False)
    if PPC_PLACEHOLDER_RE.search(serialized):
        errors.append("PPC JSON 包含占位符。")
    if not markdown_path.is_file():
        errors.append(f"PPC Markdown 文件不存在：{markdown_path}")
        markdown = ""
    else:
        markdown = markdown_path.read_text(encoding="utf-8")
        if PPC_PLACEHOLDER_RE.search(markdown):
            errors.append("PPC Markdown 包含占位符。")
        errors.extend(validate_ppc_markdown_language(markdown))
        for section in PPC_REQUIRED_MARKDOWN_SECTIONS:
            if section not in markdown:
                errors.append(f"PPC Markdown 缺少质量说明章节：{section}")
        for label in PPC_COPY_BLOCK_LABELS:
            if label not in markdown:
                errors.append(f"PPC Markdown 缺少复制区：{label}")
        if markdown.count("```") < 12:
            errors.append("PPC Markdown 必须为六组复制区分别提供代码块。")
        for key, values in copy_values.items():
            for value in values:
                if f"\n{value}\n" not in markdown:
                    errors.append(f"PPC Markdown 未逐行展示 copy_blocks.{key}：{value}")

    validation = {
        "ok": not errors,
        "primary_asin": job["primary_asin"],
        "marketplace": job["marketplace"],
        "keyword_validation": str(keyword_validation_path),
        "keyword_pool_sha256": keyword_pool_sha,
        "plan": str(plan_path),
        "plan_sha256": sha256(plan_path),
        "markdown": str(markdown_path),
        "markdown_sha256": sha256(markdown_path) if markdown_path.is_file() else None,
        "campaign_types": sorted(item for item in campaign_types if item),
        "copy_block_counts": {key: len(values) for key, values in copy_values.items()},
        "errors": errors,
    }
    if args.output:
        write_json(Path(args.output).expanduser().resolve(), validation)
    if errors:
        fail("PPC 方案验证失败：" + "；".join(errors))
    return validation


def load_rufus_qa_module():
    path = Path(__file__).with_name("validate-rufus-qa.py")
    spec = importlib.util.spec_from_file_location("rufus_qa_validator", path)
    if not spec or not spec.loader:
        fail(f"无法加载 Rufus Q/A 验证器：{path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prepare_parser = sub.add_parser("prepare")
    prepare_parser.add_argument("--input", required=True)
    prepare_parser.add_argument("--run-dir", required=True)
    transition_parser = sub.add_parser("transition")
    transition_parser.add_argument("--status", required=True)
    transition_parser.add_argument("--phase", required=True)
    transition_parser.add_argument("--artifact", action="append", default=[])
    inspect_parser = sub.add_parser("inspect-run", help="只读检查可恢复阶段与缺失产物")
    inspect_parser.add_argument("--run-dir", required=True)
    resume_parser = sub.add_parser("resume-plan", help="从已验证产物恢复滞后状态")
    resume_parser.add_argument("--run-dir", required=True)
    template_parser = sub.add_parser("emit-templates", help="生成传统 Listing、PPC 与 Q/A 草稿模板")
    template_parser.add_argument("--run-dir", required=True)
    review_parser = sub.add_parser("review-job")
    review_parser.add_argument("--job", required=True)
    review_parser.add_argument("--review", required=True)
    review_parser.add_argument("--output")
    keyword_parser = sub.add_parser("validate-keyword-pool")
    keyword_parser.add_argument("--job", required=True)
    keyword_parser.add_argument("--pool", required=True)
    keyword_parser.add_argument("--output")
    reports_parser = sub.add_parser("validate-reports")
    reports_parser.add_argument("--job", required=True)
    reports_parser.add_argument("--keyword-validation", required=True)
    reports_parser.add_argument("--reports", nargs="+", required=True)
    reports_parser.add_argument("--output")
    ppc_packet_parser = sub.add_parser("prepare-ppc-source-packet")
    ppc_packet_parser.add_argument("--job", required=True)
    ppc_packet_parser.add_argument("--keyword-validation", required=True)
    ppc_packet_parser.add_argument("--report", required=True)
    ppc_packet_parser.add_argument("--output", required=True)
    materialize_parser = sub.add_parser("materialize-legacy-result")
    materialize_parser.add_argument("--request", required=True)
    materialize_parser.add_argument("--raw-output", required=True)
    materialize_parser.add_argument("--receipt", required=True)
    materialize_parser.add_argument("--output", required=True)
    packet_parser = sub.add_parser("build-review-packet")
    packet_parser.add_argument("--job", required=True)
    packet_parser.add_argument("--keyword-validation", required=True)
    packet_parser.add_argument("--plan", required=True)
    packet_parser.add_argument("--reports", nargs="+", required=True)
    packet_parser.add_argument("--report-validation", required=True)
    packet_parser.add_argument("--output", required=True)
    packet_parser.add_argument("--markdown", required=True)
    packet_parser.add_argument("--listing-markdown", required=True)
    packet_parser.add_argument("--short-title-markdown", required=True)
    packet_parser.add_argument("--xlsm-markdown", required=True)
    packet_parser.add_argument("--rufus-markdown", required=True)
    packet_parser.add_argument("--rufus-validation", required=True)
    packet_parser.add_argument("--ppc-validation", required=True)
    seal_parser = sub.add_parser("seal-plan")
    seal_parser.add_argument("--job", required=True)
    seal_parser.add_argument("--plan", required=True)
    seal_parser.add_argument("--review", required=True)
    seal_parser.add_argument("--review-markdown", required=True)
    seal_parser.add_argument("--output", required=True)
    confirm_parser = sub.add_parser("confirm")
    confirm_parser.add_argument("--request", required=True)
    confirm_parser.add_argument("--phrase", required=True)
    confirm_parser.add_argument("--output", required=True)
    verify_parser = sub.add_parser("verify-confirmation")
    verify_parser.add_argument("--confirmation", required=True)
    apply_parser = sub.add_parser("apply-sync")
    apply_parser.add_argument("--confirmation", required=True)
    apply_parser.add_argument("--output-dir", required=True)
    apply_parser.add_argument("--result", required=True)
    output_parser = sub.add_parser("verify-output")
    output_parser.add_argument("--confirmation", required=True)
    output_parser.add_argument("--result", required=True)
    output_parser.add_argument("--output")
    ppc_parser = sub.add_parser("validate-ppc-plan")
    ppc_parser.add_argument("--job", required=True)
    ppc_parser.add_argument("--keyword-validation", required=True)
    ppc_parser.add_argument("--plan", required=True)
    ppc_parser.add_argument("--markdown", required=True)
    ppc_parser.add_argument("--source-packet", required=True)
    ppc_parser.add_argument("--request", required=True)
    ppc_parser.add_argument("--raw-output", required=True)
    ppc_parser.add_argument("--output")
    rufus_parser = sub.add_parser("validate-rufus-plan")
    rufus_parser.add_argument("--job", required=True)
    rufus_parser.add_argument("--keyword-validation", required=True)
    rufus_parser.add_argument("--ppc-validation", required=True)
    rufus_parser.add_argument("--report", required=True)
    rufus_parser.add_argument("--markdown", required=True)
    rufus_parser.add_argument("--output", required=True)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    commands = {
        "prepare": prepare, "transition": transition, "review-job": review_job,
        "inspect-run": inspect_run, "resume-plan": resume_plan,
        "emit-templates": emit_templates,
        "validate-keyword-pool": validate_keyword_pool,
        "prepare-ppc-source-packet": build_ppc_source_packet,
        "materialize-legacy-result": materialize_legacy_result,
        "validate-reports": validate_reports, "build-review-packet": build_review_packet,
        "seal-plan": seal_plan, "confirm": confirm,
        "verify-confirmation": verify_confirmation, "apply-sync": apply_sync,
        "verify-output": verify_output, "validate-ppc-plan": validate_ppc_plan,
        "validate-rufus-plan": validate_rufus_plan,
    }
    try:
        result = commands[args.command](args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ValueError, KeyError, OSError, zipfile.BadZipFile) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
