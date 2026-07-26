#!/usr/bin/env python3
"""Plan and apply multi-marketplace Listing updates to Amazon category reports."""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"x": MAIN_NS, "r": REL_NS, "pr": PKG_REL_NS}
GREEN_FILL = "FFC6EFCE"
ET.register_namespace("", MAIN_NS)
ET.register_namespace("r", REL_NS)

MARKETPLACE_IDS = {
    "US": "ATVPDKIKX0DER", "CA": "A2EUQ1WTGCTBG2", "MX": "A1AM78C64UM0Y8",
    "BR": "A2Q3Y263D00KWC", "UK": "A1F83G8C2ARO7P", "DE": "A1PA6795UKMFR9",
    "FR": "A13V1IB3VIYZZH", "IT": "APJ6JRA9NG5V4", "ES": "A1RKKUPIHCS9HS",
    "NL": "A1805IZSGTT6HS", "SE": "A2NODRKZP88ZB9", "PL": "A1C3SOZRARQ6R3",
    "BE": "AMEN7PMS3EDWL", "JP": "A1VC38T7YXB528", "AU": "A39IBJ37TRP1C6",
    "IN": "A21TJRUUN4KGV", "AE": "A2VIGQ35RCS4UG", "SA": "A17E79C6D8DWNP",
    "SG": "A19VAU5U5O7RUS", "TR": "A33AVAJ2PDY3EV", "IE": "A28R8C7NBKEWEA",
}

FILENAME_MARKERS = {
    "US": ["美国", "us", "usa", "united states"],
    "UK": ["英国", "uk", "united kingdom"],
    "DE": ["德国", "de", "germany", "deutschland"],
    "FR": ["法国", "fr", "france"],
    "IT": ["意大利", "it", "italy", "italia"],
    "ES": ["西班牙", "es", "spain", "espana"],
    "NL": ["荷兰", "nl", "netherlands"],
    "SE": ["瑞典", "se", "sweden"],
    "JP": ["日本", "jp", "japan"],
    "CA": ["加拿大", "ca", "canada"],
    "AU": ["澳大利亚", "au", "australia"],
    "MX": ["墨西哥", "mx", "mexico"],
    "BR": ["巴西", "br", "brazil"],
    "BE": ["比利时", "be", "belgium", "belgie", "belgique"],
    "PL": ["波兰", "pl", "poland", "polska"],
    "SE": ["瑞典", "se", "sweden", "sverige"],
    "IN": ["印度", "in", "india"],
    "AE": ["阿联酋", "ae", "uae", "united arab emirates"],
    "SA": ["沙特", "sa", "saudi arabia"],
    "SG": ["新加坡", "sg", "singapore"],
    "TR": ["土耳其", "tr", "turkey", "türkiye"],
    "IE": ["爱尔兰", "ie", "ireland"],
}

MARKETPLACE_ALIASES = {
    "US": ["US", "USA", "Amazon US", "United States", "美国"],
    "CA": ["CA", "Amazon CA", "Canada", "加拿大"],
    "MX": ["MX", "Amazon MX", "Mexico", "México", "墨西哥"],
    "BR": ["BR", "Amazon BR", "Brazil", "Brasil", "巴西"],
    "UK": ["UK", "Amazon UK", "United Kingdom", "Great Britain", "England", "英国"],
    "DE": ["DE", "Amazon DE", "Germany", "Deutschland", "德国"],
    "FR": ["FR", "Amazon FR", "France", "法国"],
    "IT": ["IT", "Amazon IT", "Italy", "Italia", "意大利"],
    "ES": ["ES", "Amazon ES", "Spain", "España", "西班牙"],
    "NL": ["NL", "Amazon NL", "Netherlands", "Holland", "Nederland", "荷兰"],
    "SE": ["SE", "Amazon SE", "Sweden", "Sverige", "瑞典"],
    "PL": ["PL", "Amazon PL", "Poland", "Polska", "波兰"],
    "BE": ["BE", "Amazon BE", "Belgium", "Belgique", "België", "比利时"],
    "JP": ["JP", "Amazon JP", "Japan", "日本"],
    "AU": ["AU", "Amazon AU", "Australia", "澳大利亚"],
    "IN": ["IN", "Amazon IN", "India", "印度"],
    "AE": ["AE", "Amazon AE", "United Arab Emirates", "UAE", "阿联酋"],
    "SA": ["SA", "Amazon SA", "Saudi Arabia", "Saudi", "沙特"],
    "SG": ["SG", "Amazon SG", "Singapore", "新加坡"],
    "TR": ["TR", "Amazon TR", "Turkey", "Türkiye", "Turkiye", "土耳其"],
    "IE": ["IE", "Amazon IE", "Ireland", "爱尔兰"],
}

# Exact aliases are preferred because Amazon templates reuse the same translated
# labels across category files. Duplicate headers (for example five bullets)
# intentionally map to a list of columns.
ALIASES = {
    "sku": ["SKU", "Seller SKU", "Merchant SKU", "Venditore SKU", "SKU venditore", "卖家 SKU", "商家 SKU"],
    "parent_sku": ["Parent SKU", "Parentage SKU", "SKU genitore", "SKU parent", "父 SKU", "父商品 SKU"],
    "parent_product_id": ["Parent ASIN", "Parent Product ID", "ASIN genitore", "父 ASIN", "父商品 ASIN"],
    "product_id": ["Produkt-ID", "Product ID", "Product-id", "Identifiant du produit", "ID prodotto", "ID de producto", "ID do produto", "Produkt ID", "Identyfikator produktu", "商品ID", "商品 ID", "Ürün Kimliği", "معرف المنتج"],
    "product_id_type": ["Art der Produkt-ID", "Product ID Type", "Type d'identifiant produit", "Type d'identifiant de produit", "Tipo di ID prodotto", "Tipo ID di prodotto", "Tipo de ID de producto", "Tipo de identificador de producto", "Tipo de ID do produto", "Typ produkt-ID", "Typ identyfikatora produktu", "商品IDのタイプ", "Ürün Kimliği Türü", "نوع معرف المنتج"],
    "title": ["Title", "Titel", "Titre", "Titolo", "Título", "Título do produto", "Titel van het product", "Produktnamn", "Tytuł", "商品名", "Ürün Başlığı", "عنوان المنتج"],
    "item_name": ["Artikelname", "Item Name", "Nom de l'article", "Nom de l’article", "Nome articolo", "Nome dell'articolo", "Nome dell’articolo", "Nombre del producto", "Nome do item", "Artikelnaam", "Produktnamn", "Nazwa produktu", "商品名", "Ürün adı", "اسم المنتج"],
    "item_highlight": ["Artikel-Highlight", "Item Highlight", "Point fort de l'article", "Point fort de l’article", "Punto saliente dell'articolo", "Caratteristiche principali articolo", "Aspecto destacado del artículo", "Destaque do item", "Artikelhoogtepunt", "Produktfördel", "Najważniejsza cecha", "商品のハイライト", "Ürün öne çıkan özelliği", "أبرز ميزات المنتج"],
    "description": ["Beschreibung des Produkts", "Product Description", "Description du produit", "Descrizione del prodotto", "Descripción del producto", "Descrição do produto", "Productbeschrijving", "Produktbeskrivning", "Opis produktu", "商品の説明", "Ürün açıklaması", "وصف المنتج"],
    "bullet": ["Aufzählungspunkt", "Bullet Point", "Bullet", "Puces", "Puce", "Punto elenco", "Viñeta", "Marcador", "Opsommingsteken", "Punkt", "Punkt wypunktowania", "箇条書き", "Madde işareti", "نقطة تعداد"],
    "search_terms": ["Suchbegriffe", "Search Terms", "Termes de recherche", "Mots-clés de recherche", "Termini di ricerca", "Chiavi di ricerca", "Términos de búsqueda", "Termos de pesquisa", "Zoektermen", "Söktermer", "Wyszukiwane hasła", "検索キーワード", "Arama terimleri", "مصطلحات البحث"],
}

ATTRIBUTE_PATTERNS = {
    "sku": (r"(?:^|[._])(?:item|seller|merchant)_sku(?:\[|#|$)",),
    "parent_sku": (r"(?:^|[._])parent_sku(?:\[|#|$)",),
    "parent_product_id": (r"(?:^|[._])parent_(?:asin|product_id)(?:_value)?(?:\[|#|$)",),
    "product_id": (r"(?:^|[._])product_id(?:_value)?(?:\[|#|$)",),
    "product_id_type": (r"product_id_type(?:\[|#|$)",),
    "title": (r"(?:^|[._])title(?:\[|#|$)",),
    "item_name": (r"item_name(?:\[|#|$)",),
    "item_highlight": (r"(?:title_differentiation|item_highlight)(?:\[|#|$)",),
    "description": (r"product_description(?:\[|#|$)",),
    "bullet": (r"bullet_point(?:\[|#|$)",),
    "search_terms": (r"generic_keyword(?:\[|#|$)",),
}

LANGUAGE_MARKETPLACES = {
    "de": {"DE", "AT"}, "en": {"US", "CA", "UK", "AU", "IN", "AE", "SA", "SG", "IE"},
    "fr": {"FR", "BE", "CA"}, "it": {"IT"}, "es": {"ES", "MX"},
    "pt": {"BR"}, "nl": {"NL", "BE"}, "sv": {"SE"}, "pl": {"PL"},
    "ja": {"JP"}, "tr": {"TR"}, "ar": {"AE", "SA"},
}

LANGUAGE_SIGNATURES = {
    "de": {"artikelname", "produkt-id", "suchbegriffe"},
    "en": {"item name", "product id", "search terms"},
    "fr": {"nom de l'article", "identifiant du produit", "mots-clés de recherche"},
    "it": {"nome articolo", "id prodotto", "termini di ricerca"},
    "es": {"nombre del producto", "id de producto", "términos de búsqueda"},
    "pt": {"nome do item", "id do produto", "termos de pesquisa"},
    "nl": {"artikelnaam", "product-id", "zoektermen"},
    "sv": {"produktnamn", "produkt-id", "söktermer"},
    "pl": {"nazwa produktu", "identyfikator produktu", "wyszukiwane hasła"},
    "ja": {"商品名", "商品id", "検索キーワード"},
    "tr": {"ürün adı", "ürün kimliği", "arama terimleri"},
    "ar": {"اسم المنتج", "معرف المنتج", "مصطلحات البحث"},
}


def normalize(value: object) -> str:
    """Normalize user-facing labels across punctuation, accents, and spacing."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("’", "'").replace("‘", "'").replace("–", "-").replace("—", "-")
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return re.sub(r"\s+", " ", text.strip()).casefold()


def normalize_marketplace(value: object) -> str:
    normalized = normalize(value)
    for marketplace, aliases in MARKETPLACE_ALIASES.items():
        if normalized in {normalize(alias) for alias in aliases}:
            return marketplace
    return str(value or "").strip().upper()


def collect_paths(paths: list[str], directories: list[str], patterns: tuple[str, ...]) -> list[Path]:
    found = {Path(value).expanduser().resolve() for value in paths}
    for directory in directories:
        root = Path(directory).expanduser().resolve()
        if not root.is_dir():
            raise ValueError(f"指定目录不存在：{root}")
        for pattern in patterns:
            found.update(path.resolve() for path in root.rglob(pattern) if not path.name.startswith("~$"))
    return sorted(path for path in found if path.is_file())


def collect_report_paths(paths: list[str], directories: list[str]) -> list[Path]:
    explicit = {Path(value).expanduser().resolve() for value in paths}
    invalid = [path for path in explicit if path.suffix.lower() != ".json"]
    if invalid:
        raise ValueError(f"显式优化报告必须是 JSON 文件：{', '.join(map(str, invalid))}")
    discovered = set(explicit)
    for directory in directories:
        root = Path(directory).expanduser().resolve()
        if not root.is_dir():
            raise ValueError(f"指定目录不存在：{root}")
        discovered.update(root.rglob("listing-optimization-report-*.json"))
    return sorted(path.resolve() for path in discovered if path.is_file())


def parse_title_options(raw: dict, selections: dict[str, int], short_title_policy: str = "first") -> tuple[dict[str, str] | None, str | None]:
    options = raw.get("title_options_2026") or []
    asin = raw["asin"].upper()
    if not options:
        return {"title": raw["listing"]["title"], "item_highlight": "", "title_mode": "标准长标题"}, None
    selected = selections.get(asin)
    if selected is None and short_title_policy == "first":
        selected = next((item.get("option") for item in options if item.get("option") == 1), options[0].get("option"))
    if selected is None and len(options) == 1:
        selected = options[0].get("option")
    if selected is None and short_title_policy == "long":
        return {"title": raw["listing"]["title"], "item_highlight": "", "title_mode": "标准长标题（策略回退）"}, None
    if selected is None:
        return None, "存在多个短标题方案，需用户选择标题方案。"
    option = next((item for item in options if item.get("option") == selected), None)
    if option is None:
        return None, f"标题方案 {selected} 不存在。"
    return {"title": option["title"], "item_highlight": option.get("item_highlights", ""), "title_mode": f"短标题方案 {selected}"}, None


def load_reports(paths: list[Path], selections: dict[str, int], short_title_policy: str = "first") -> list[dict]:
    records = []
    for path in paths:
        try:
            raw = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            records.append({"asin": "", "marketplace": "", "source": str(path), "status": "blocked", "blockers": [f"无法读取优化报告 JSON：{exc}"], "fields": {}})
            continue
        asin = str(raw.get("asin", "")).strip().upper()
        marketplace = normalize_marketplace(raw.get("marketplace", ""))
        listing = raw.get("listing") or {}
        missing = [field for field in ["title", "bullets", "description", "backend_search_terms"] if not listing.get(field)]
        title, title_error = parse_title_options(raw, selections, short_title_policy) if asin and listing else (None, "缺少 ASIN 或 Listing 数据。")
        status = "ready" if asin and marketplace and not missing and not title_error else "blocked"
        records.append({
            "asin": asin, "marketplace": marketplace, "source": str(path), "status": status,
            "blockers": ([f"缺少字段：{', '.join(missing)}"] if missing else []) + ([title_error] if title_error else []),
            "fields": {
                "title": title["title"] if title else "",
                "item_highlight": title["item_highlight"] if title else "",
                "title_mode": title["title_mode"] if title else "",
                "description": listing.get("description", ""),
                "bullets": listing.get("bullets", []),
                "search_terms": listing.get("backend_search_terms", ""),
            },
        })
    duplicate_keys = defaultdict(list)
    for record in records:
        duplicate_keys[(record["marketplace"], record["asin"])].append(record)
    for key, members in duplicate_keys.items():
        if len(members) > 1:
            for member in members:
                member["status"] = "blocked"
                member["blockers"].append(f"存在重复优化报告：{key[0]} / {key[1]}。")
    return records


def header_map(ws) -> dict[str, list[int]]:
    fields, _evidence = scan_template_fields(ws)
    return fields


def scan_template_fields(ws) -> tuple[dict[str, list[int]], dict[str, list[dict]]]:
    """Prefer stable Amazon attribute IDs before localized display aliases."""
    result: dict[str, list[int]] = defaultdict(list)
    evidence: dict[str, list[dict]] = defaultdict(list)
    alias_map = {key: {normalize(alias) for alias in aliases} for key, aliases in ALIASES.items()}
    for column in range(1, ws.max_column + 1):
        attribute = str(ws.cell(5, column).value or "")
        display_values = [(row, normalize(ws.cell(row, column).value)) for row in (3, 4) if row <= ws.max_row]
        matches = [
            key for key, patterns in ATTRIBUTE_PATTERNS.items()
            if any(re.search(pattern, attribute, flags=re.IGNORECASE) for pattern in patterns)
        ]
        if "parent_product_id" in matches and "product_id" in matches:
            matches.remove("product_id")
        source = "attribute_id" if matches else ""
        if not matches:
            for key, aliases in alias_map.items():
                if any(value in aliases for _row, value in display_values if value):
                    matches.append(key)
            source = "display_alias" if matches else ""
        for key in matches:
            if column not in result[key]:
                result[key].append(column)
                evidence[key].append({
                    "column": column,
                    "letter": ws.cell(4, column).column_letter,
                    "source": source,
                    "attribute_id": attribute,
                    "display_name": next((str(ws.cell(row, column).value or "") for row in (4, 3) if ws.cell(row, column).value), ""),
                })
    return dict(result), dict(evidence)


def header_language_evidence(ws) -> set[str]:
    headers = {normalize(ws.cell(row, column).value) for row in (3, 4) if row <= ws.max_row for column in range(1, ws.max_column + 1)}
    evidence = set()
    for language, signatures in LANGUAGE_SIGNATURES.items():
        if len(headers & {normalize(value) for value in signatures}) >= 2:
            evidence.add(language)
    return evidence


def marker_matches(filename: str, marker: str) -> bool:
    if marker.isascii() and marker.isalpha() and len(marker) <= 3:
        return re.search(rf"(?<![a-z0-9]){re.escape(marker)}(?![a-z0-9])", filename) is not None
    return marker in filename


def field_metadata(ws, fields: dict[str, list[int]], evidence: dict[str, list[dict]] | None = None) -> dict[str, list[dict]]:
    metadata = defaultdict(list)
    for key, columns in fields.items():
        for column in columns:
            letter = ws.cell(4, column).column_letter
            item = {"column": column, "letter": letter, "hidden": bool(ws.column_dimensions[letter].hidden)}
            if evidence:
                item.update(next((value for value in evidence.get(key, []) if value["column"] == column), {}))
            metadata[key].append(item)
    return dict(metadata)


def marketplace_evidence(path: Path, workbook, worksheet) -> dict[str, set[str]]:
    evidence: dict[str, set[str]] = defaultdict(set)
    filename = path.name.casefold()
    probe = " ".join(
        str(ws.cell(row, column).value or "")
        for ws in workbook.worksheets for row in range(3, min(ws.max_row, 5) + 1)
        for column in range(1, ws.max_column + 1)
    ).upper()
    for marketplace, marketplace_id in MARKETPLACE_IDS.items():
        if marketplace_id in probe:
            evidence[marketplace].add("marketplace_id")
    for marketplace, markers in FILENAME_MARKERS.items():
        if any(marker_matches(filename, marker) for marker in markers):
            evidence[marketplace].add("filename")
    for language in header_language_evidence(worksheet):
        for marketplace in LANGUAGE_MARKETPLACES[language]:
            if marketplace in MARKETPLACE_IDS:
                evidence[marketplace].add(f"template_language:{language}")
    return evidence


def template_candidate(ws) -> dict | None:
    fields, evidence = scan_template_fields(ws)
    if not {"product_id", "product_id_type"}.issubset(fields):
        return None
    required_listing_fields = {"title", "item_name", "description", "bullet", "search_terms"}
    return {
        "sheet": ws.title,
        "fields": fields,
        "field_evidence": evidence,
        "field_metadata": field_metadata(ws, fields, evidence),
        "listing_field_count": len(required_listing_fields & set(fields)),
        "missing_listing_fields": sorted(required_listing_fields - set(fields)),
    }


def classify_workbook(path: Path) -> dict:
    extension = path.suffix.lower()
    if extension == ".xls":
        return {"path": str(path), "status": "blocked", "marketplace": "", "reason": "旧版 .xls 无法保证无损 OOXML 定点更新；需用户接受兼容转换输出。"}
    if extension not in {".xlsx", ".xlsm"}:
        return {"path": str(path), "status": "blocked", "marketplace": "", "reason": "不是支持的分类商品报告格式。"}
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=extension == ".xlsm")
    candidates = [candidate for ws in wb.worksheets if (candidate := template_candidate(ws))]
    if not candidates:
        return {"path": str(path), "status": "blocked", "marketplace": "", "reason": "未找到含 ASIN 标识字段的分类商品模板工作表。"}
    worksheet = wb[candidates[0]["sheet"]]
    evidence = marketplace_evidence(path, wb, worksheet)
    winners = [marketplace for marketplace, sources in evidence.items() if len(sources) >= 2]
    if len(winners) != 1:
        detail = {marketplace: sorted(sources) for marketplace, sources in evidence.items()}
        return {"path": str(path), "status": "blocked", "marketplace": "", "reason": "无法以至少两类证据唯一确认分类商品报告所属站点。", "station_evidence": detail}
    marketplace = winners[0]
    status = "ready" if len(candidates) == 1 else "ambiguous"
    result = {"path": str(path), "status": status, "marketplace": marketplace, "candidates": candidates, "station_evidence": sorted(evidence[marketplace])}
    if len(candidates) == 1:
        result.update({"sheet": candidates[0]["sheet"], "fields": candidates[0]["fields"], "field_metadata": candidates[0]["field_metadata"], "rows": wb[worksheet.title].max_row, "columns": wb[worksheet.title].max_column})
    else:
        result["reason"] = "存在多个可识别模板工作表；将在目标 ASIN 全量命中后消歧。"
    return result


def product_rows(workbook: dict, asin: str) -> list[int]:
    path = Path(workbook["path"])
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb[workbook["sheet"]]
    fields = workbook["fields"]
    product_id = fields["product_id"][0]
    product_id_type = fields["product_id_type"][0]
    matches = []
    for row in range(7, ws.max_row + 1):
        if normalize(ws.cell(row, product_id).value).upper() != asin:
            continue
        if normalize(ws.cell(row, product_id_type).value) != "asin":
            continue
        # Empty Listing fields are legitimate targets for optimization data.
        matches.append(row)
    return matches


def resolve_workbook_candidate(workbook: dict, asins: list[str]) -> tuple[dict | None, str]:
    """Resolve multi-template workbooks only when one sheet covers every target ASIN."""
    candidates = workbook.get("candidates") or [{
        "sheet": workbook["sheet"], "fields": workbook["fields"],
        "field_metadata": workbook.get("field_metadata", {}),
        "listing_field_count": len(workbook.get("fields", {})),
        "missing_listing_fields": [],
    }]
    path = Path(workbook["path"])
    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    matches = []
    for candidate in candidates:
        fields = candidate["fields"]
        if candidate["listing_field_count"] < 5:
            continue
        sheet_view = {"path": workbook["path"], "sheet": candidate["sheet"], "fields": fields}
        rows_by_asin = {asin: product_rows(sheet_view, asin) for asin in asins}
        if all(rows_by_asin.values()):
            matches.append((candidate, rows_by_asin))
    if not matches:
        return None, "没有单个模板工作表包含全部目标 ASIN。"
    highest = max(candidate["listing_field_count"] for candidate, _rows in matches)
    winners = [(candidate, rows) for candidate, rows in matches if candidate["listing_field_count"] == highest]
    if len(winners) != 1:
        return None, "多个模板工作表以相同字段完整度包含全部目标 ASIN。"
    candidate, rows_by_asin = winners[0]
    return {
        **workbook,
        "status": "ready",
        "sheet": candidate["sheet"],
        "fields": candidate["fields"],
        "field_metadata": candidate["field_metadata"],
        "rows_by_asin": rows_by_asin,
        "sheet_selection_reason": f"唯一工作表命中全部 {len(asins)} 个目标 ASIN，Listing 字段完整度 {candidate['listing_field_count']}/5。",
    }, ""


def planned_updates(workbook: dict, rows: list[int], fields: dict, skip_fields: set[str], has_highlight: bool) -> list[dict]:
    path = Path(workbook["path"])
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    ws = wb[workbook["sheet"]]
    field_columns = {
        "title": fields["title"][:1], "item_name": fields["item_name"][:1],
        "item_highlight": fields.get("item_highlight", [])[:1], "description": fields["description"][:1],
        "bullet": fields.get("bullet", [])[:5], "search_terms": fields["search_terms"][:1],
    }
    result = []
    for key, columns in field_columns.items():
        for column in columns:
            values = [ws.cell(row, column).value for row in rows]
            if key == "item_highlight" and not has_highlight:
                action = "保留原值"
            elif key in skip_fields:
                action = "跳过"
            else:
                action = "写入"
            result.append({
                "field": key, "column": column, "letter": ws.cell(4, column).column_letter,
                "hidden": bool(ws.column_dimensions[ws.cell(4, column).column_letter].hidden),
                "source_state": "空白，将填充" if any(value in (None, "") for value in values) else "已有内容，将覆盖",
                "action": action,
            })
    return result


def build_plan(reports: list[dict], workbooks: list[dict], skips: dict[str, set[str]]) -> dict:
    target_asins = defaultdict(list)
    for report in reports:
        if report.get("status") == "ready":
            target_asins[report["marketplace"]].append(report["asin"])
    resolved_workbooks = []
    for workbook in workbooks:
        if workbook.get("status") not in {"ready", "ambiguous"}:
            resolved_workbooks.append(workbook)
            continue
        resolved, reason = resolve_workbook_candidate(workbook, sorted(set(target_asins.get(workbook["marketplace"], []))))
        resolved_workbooks.append(resolved if resolved else {**workbook, "status": "blocked", "reason": reason})
    actions = []
    for report in reports:
        if report["status"] != "ready":
            actions.append({**report, "action": "blocked"})
            continue
        marketplace_books = [book for book in resolved_workbooks if book.get("marketplace") == report["marketplace"]]
        candidates = [book for book in marketplace_books if book["status"] == "ready"]
        if len(candidates) != 1:
            reasons = [book.get("reason", "") for book in marketplace_books if book.get("status") == "blocked" and book.get("reason")]
            actions.append({**report, "action": "blocked", "blockers": report["blockers"] + reasons[:1] + (["未找到唯一匹配的分类商品报告文件。"] if not reasons else [])})
            continue
        rows = candidates[0].get("rows_by_asin", {}).get(report["asin"], product_rows(candidates[0], report["asin"]))
        if not rows:
            actions.append({**report, "action": "blocked", "blockers": report["blockers"] + ["目标报告中未找到完整商品资料行。"]})
            continue
        skip_fields = skips.get(report["asin"], set())
        actions.append({
            **report, "action": "update", "workbook": candidates[0]["path"], "sheet": candidates[0]["sheet"],
            "rows": rows, "skip_fields": sorted(skip_fields),
            "station_evidence": candidates[0]["station_evidence"],
            "field_metadata": candidates[0]["field_metadata"],
            "sheet_selection_reason": candidates[0].get("sheet_selection_reason", "模板工作表唯一匹配。"),
            "planned_updates": planned_updates(candidates[0], rows, candidates[0]["fields"], skip_fields, bool(report["fields"]["item_highlight"])),
        })
    return {"generated_at": dt.datetime.now().isoformat(timespec="seconds"), "status": "ready" if all(item["action"] == "update" for item in actions) else "review_required", "actions": actions, "workbooks": resolved_workbooks}


def xml_paths(archive: zipfile.ZipFile, sheet_name: str) -> str:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {rel.get("Id"): rel.get("Target") for rel in rels.findall("pr:Relationship", NS)}
    for sheet in workbook.findall("x:sheets/x:sheet", NS):
        if sheet.get("name") == sheet_name:
            target = targets[sheet.get(f"{{{REL_NS}}}id")].lstrip("/")
            return target if target.startswith("xl/") else "xl/" + target
    raise ValueError(f"工作表不存在：{sheet_name}")


def shared_strings(archive: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    return ["".join(node.itertext()) for node in root.findall("x:si", NS)]


def column_letter(index: int) -> str:
    chars = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        chars = chr(65 + remainder) + chars
    return chars


def cell_column_index(cell: ET.Element) -> int:
    letters = re.sub(r"[^A-Z]", "", cell.get("r", "").upper())
    value = 0
    for letter in letters:
        value = value * 26 + ord(letter) - 64
    return value


def add_green_styles(styles_xml: bytes, old_styles: set[int]) -> tuple[bytes, dict[int, int]]:
    root = ET.fromstring(styles_xml)
    fills = root.find("x:fills", NS)
    cell_xfs = root.find("x:cellXfs", NS)
    fill_id = len(fills)
    fill = ET.SubElement(fills, f"{{{MAIN_NS}}}fill")
    pattern = ET.SubElement(fill, f"{{{MAIN_NS}}}patternFill", {"patternType": "solid"})
    ET.SubElement(pattern, f"{{{MAIN_NS}}}fgColor", {"rgb": GREEN_FILL})
    ET.SubElement(pattern, f"{{{MAIN_NS}}}bgColor", {"indexed": "64"})
    fills.set("count", str(fill_id + 1))
    original = list(cell_xfs)
    # openpyxl can expose a virtual style for an empty cell that inherits a
    # column style, while that id is not present in the OOXML cellXfs table.
    # Use the default cell style for those new cells before applying green fill.
    old_styles = {style if 0 <= style < len(original) else 0 for style in old_styles}
    mapping = {}
    for old in sorted(old_styles):
        clone = copy.deepcopy(original[old])
        clone.set("fillId", str(fill_id))
        clone.set("applyFill", "1")
        mapping[old] = len(cell_xfs)
        cell_xfs.append(clone)
    cell_xfs.set("count", str(len(cell_xfs)))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), mapping


def patch_sheet(xml: bytes, fields: dict[str, list[int]], rows: list[int], record: dict, style_map: dict[int, int]) -> tuple[bytes, int]:
    root = ET.fromstring(xml)
    row_nodes = {int(node.get("r")): node for node in root.findall("x:sheetData/x:row", NS)}
    updates_by_column = {
        fields["title"][0]: record["fields"]["title"],
        fields["item_name"][0]: record["fields"]["title"],
        fields["description"][0]: "".join(f"<p>{p}</p>" for p in record["fields"]["description"].split("\n\n")),
        fields["search_terms"][0]: record["fields"]["search_terms"],
    }
    if fields.get("item_highlight") and record["fields"]["item_highlight"]:
        updates_by_column[fields["item_highlight"][0]] = record["fields"]["item_highlight"]
    for column, bullet in zip(fields["bullet"][:5], record["fields"]["bullets"][:5]):
        updates_by_column[column] = bullet
    skipped = set(record.get("skip_fields", []))
    field_columns = {
        "title": fields["title"][:1], "item_name": fields["item_name"][:1], "item_highlight": fields.get("item_highlight", [])[:1],
        "description": fields["description"][:1], "bullet": fields.get("bullet", [])[:5], "search_terms": fields["search_terms"][:1],
    }
    skipped_columns = {column for key, columns in field_columns.items() if key in skipped for column in columns}
    written = 0
    for row_number in rows:
        if row_number not in row_nodes:
            raise ValueError(f"模板缺少目标行：{row_number}")
        cells = {node.get("r"): node for node in row_nodes[row_number].findall("x:c", NS)}
        for column, value in updates_by_column.items():
            if column in skipped_columns:
                continue
            ref = f"{column_letter(column)}{row_number}"
            cell = cells.get(ref)
            if cell is None:
                original_style = record["field_style_ids"][column]
                cell = ET.Element(f"{{{MAIN_NS}}}c", {"r": ref, "s": str(original_style)})
                row_nodes[row_number].append(cell)
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "inlineStr")
            style_id = int(cell.get("s", "0"))
            fallback_style = style_map.get(0, next(iter(style_map.values())))
            cell.set("s", str(style_map.get(style_id, fallback_style)))
            inline = ET.SubElement(cell, f"{{{MAIN_NS}}}is")
            text = ET.SubElement(inline, f"{{{MAIN_NS}}}t")
            text.text = value
            written += 1
        ordered = sorted(row_nodes[row_number].findall("x:c", NS), key=cell_column_index)
        for node in row_nodes[row_number].findall("x:c", NS):
            row_nodes[row_number].remove(node)
        for node in ordered:
            row_nodes[row_number].append(node)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), written


def related_parent_rows(ws, fields: dict[str, list[int]], modified_rows: set[int]) -> set[int]:
    """Resolve direct parent rows from Parent SKU or Parent ASIN references."""
    product_id_columns = fields.get("product_id", [])[:1]
    sku_columns = fields.get("sku", [])[:1]
    parent_reference_columns = fields.get("parent_sku", []) + fields.get("parent_product_id", [])
    if not parent_reference_columns:
        return set()

    rows_by_reference: dict[str, set[int]] = defaultdict(set)
    for row in range(7, ws.max_row + 1):
        for column in product_id_columns + sku_columns:
            value = normalize(ws.cell(row, column).value).upper()
            if value:
                rows_by_reference[value].add(row)

    parents = set()
    for row in modified_rows:
        for column in parent_reference_columns:
            reference = normalize(ws.cell(row, column).value).upper()
            if reference:
                parents.update(rows_by_reference.get(reference, set()))
    return parents - modified_rows


def prune_sheet_data(xml: bytes, keep_rows: set[int]) -> tuple[bytes, int]:
    """Delete unrelated data rows and compact retained rows from row 7."""
    root = ET.fromstring(xml)
    sheet_data = root.find("x:sheetData", NS)
    deleted = 0
    retained_data_rows = []
    for row in list(sheet_data.findall("x:row", NS)):
        row_number = int(row.get("r"))
        if row_number >= 7 and row_number not in keep_rows:
            sheet_data.remove(row)
            deleted += 1
        elif row_number >= 7:
            retained_data_rows.append(row)

    retained_data_rows.sort(key=lambda row: int(row.get("r")))
    for new_row_number, row in enumerate(retained_data_rows, 7):
        row.set("r", str(new_row_number))
        for cell in row.findall("x:c", NS):
            reference = cell.get("r", "")
            cell.set("r", re.sub(r"\d+$", str(new_row_number), reference))

    last_row = 6 + len(retained_data_rows)
    dimension = root.find("x:dimension", NS)
    if dimension is not None and dimension.get("ref"):
        dimension.set("ref", re.sub(r"([A-Z]+)\d+$", rf"\g<1>{last_row}", dimension.get("ref")))
    auto_filter = root.find("x:autoFilter", NS)
    if auto_filter is not None and auto_filter.get("ref"):
        auto_filter.set("ref", re.sub(r"([A-Z]+)\d+$", rf"\g<1>{last_row}", auto_filter.get("ref")))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True), deleted


def safe_output_dir(output_dir: Path, sources: list[Path]) -> Path:
    resolved = output_dir.expanduser().resolve()
    for source in sources:
        source_dir = source.parent.resolve()
        if resolved == source_dir or resolved.is_relative_to(source_dir):
            raise ValueError(f"输出目录不能位于源报告目录或其子目录：{resolved}")
    return resolved


def next_output_path(output_dir: Path, marketplace: str, asin_count: int, extension: str) -> Path:
    prefix = f"{dt.date.today().isoformat()}_Amazon{marketplace}_分类商品报告_Listing同步完成_{asin_count}ASIN_"
    version = 1
    while True:
        candidate = output_dir / f"{prefix}v{version}{extension}"
        if not candidate.exists():
            return candidate
        version += 1


def apply_plan(plan: dict, output_dir: Path) -> list[dict]:
    if plan.get("status") != "ready":
        raise ValueError("计划存在阻塞项，必须先修正并重新生成审核计划。")
    grouped = defaultdict(list)
    for action in plan["actions"]:
        grouped[action["workbook"]].append(action)
    results = []
    output_dir = safe_output_dir(output_dir, [Path(name) for name in grouped])
    output_dir.mkdir(parents=True, exist_ok=True)
    for source_name, actions in grouped.items():
        source = Path(source_name)
        marketplace = actions[0]["marketplace"]
        output = next_output_path(output_dir, marketplace, len(actions), source.suffix)
        with zipfile.ZipFile(source) as archive:
            sheet_paths = {action["sheet"]: xml_paths(archive, action["sheet"]) for action in actions}
            styles = archive.read("xl/styles.xml")
            style_ids = set()
            wb = load_workbook(source, read_only=False, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
            for action in actions:
                ws = wb[action["sheet"]]
                fields = header_map(ws)
                action["field_columns"] = fields
                action["field_style_ids"] = {}
                columns_to_style = [fields["title"][0], fields["item_name"][0], fields["description"][0], fields["search_terms"][0], *fields["bullet"][:5]]
                if action["fields"]["item_highlight"]:
                    columns_to_style.extend(fields.get("item_highlight", [])[:1])
                for row in action["rows"]:
                    for column in columns_to_style:
                        style_id = ws.cell(row, column).style_id
                        style_ids.add(style_id)
                        action["field_style_ids"][column] = style_id
            updated_styles, style_map = add_green_styles(styles, style_ids)
            modified_sheets = {}
            written = 0
            retained_parent_rows = 0
            deleted_data_rows = 0
            for action in actions:
                path = sheet_paths[action["sheet"]]
                old = modified_sheets.get(path, archive.read(path))
                modified_sheets[path], count = patch_sheet(old, action["field_columns"], action["rows"], action, style_map)
                written += count
            for sheet_name in {action["sheet"] for action in actions}:
                path = sheet_paths[sheet_name]
                sheet_actions = [action for action in actions if action["sheet"] == sheet_name]
                modified_rows = {row for action in sheet_actions for row in action["rows"]}
                parents = related_parent_rows(wb[sheet_name], header_map(wb[sheet_name]), modified_rows)
                retained_parent_rows += len(parents)
                modified_sheets[path], deleted = prune_sheet_data(modified_sheets[path], modified_rows | parents)
                deleted_data_rows += deleted
            with zipfile.ZipFile(output, "w") as target:
                for info in archive.infolist():
                    if info.filename == "xl/styles.xml":
                        payload = updated_styles
                    elif info.filename in modified_sheets:
                        payload = modified_sheets[info.filename]
                    else:
                        payload = archive.read(info.filename)
                    target.writestr(info, payload)
        results.append({
            "source": str(source), "output": str(output), "asins": [a["asin"] for a in actions],
            "rows": sum(len(a["rows"]) for a in actions), "written_cells": written,
            "retained_parent_rows": retained_parent_rows, "deleted_data_rows": deleted_data_rows,
        })
    return results


def parse_selections(values: list[str]) -> dict[str, int]:
    selections = {}
    for value in values:
        asin, option = value.split("=", 1)
        selections[asin.strip().upper()] = int(option)
    return selections


def parse_skip_fields(values: list[str]) -> dict[str, set[str]]:
    valid = {"title", "item_name", "item_highlight", "description", "bullet", "search_terms"}
    result: dict[str, set[str]] = defaultdict(set)
    for value in values:
        asin, field = value.split("=", 1)
        field = field.strip()
        if field not in valid:
            raise ValueError(f"不支持的跳过字段：{field}")
        result[asin.strip().upper()].add(field)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    plan_parser = subparsers.add_parser("plan")
    plan_parser.add_argument("--reports", nargs="*", default=[])
    plan_parser.add_argument("--reports-dir", nargs="*", default=[])
    plan_parser.add_argument("--workbooks", nargs="*", default=[])
    plan_parser.add_argument("--workbooks-dir", nargs="*", default=[])
    plan_parser.add_argument("--title-option", action="append", default=[])
    plan_parser.add_argument("--short-title-policy", choices=["first", "require-selection", "long"], default="first", help="短标题方案默认策略：first 使用方案 1；require-selection 阻塞等待选择；long 使用标准长标题")
    plan_parser.add_argument("--skip-field", action="append", default=[], help="ASIN=field；仅用于用户明确要求不填充的字段")
    plan_parser.add_argument("--output", required=True)
    apply_parser = subparsers.add_parser("apply")
    apply_parser.add_argument("--plan", required=True)
    apply_parser.add_argument("--output-dir", default=None)
    args = parser.parse_args()
    if args.command == "plan":
        reports = collect_report_paths(args.reports, args.reports_dir)
        workbooks = collect_paths(args.workbooks, args.workbooks_dir, ("*.xlsm", "*.xlsx", "*.xls"))
        if not reports:
            raise ValueError("未找到标准 Listing 优化报告 JSON。")
        if not workbooks:
            raise ValueError("未找到用户提供或用户指定目录中的分类商品报告。")
        plan = build_plan(load_reports(reports, parse_selections(args.title_option), args.short_title_policy), [classify_workbook(path) for path in workbooks], parse_skip_fields(args.skip_field))
        Path(args.output).write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(plan, ensure_ascii=False, indent=2))
    else:
        plan = json.loads(Path(args.plan).read_text(encoding="utf-8"))
        output_dir = Path(args.output_dir) if args.output_dir else Path.cwd() / "outputs"
        print(json.dumps({"results": apply_plan(plan, output_dir)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
