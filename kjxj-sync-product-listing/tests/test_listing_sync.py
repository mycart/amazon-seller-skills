import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPT = Path(__file__).parents[1] / "scripts" / "listing_sync.py"
SPEC = importlib.util.spec_from_file_location("listing_sync", SCRIPT)
SYNC = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(SYNC)


def make_report(asin: str, options=None):
    report = {
        "asin": asin,
        "marketplace": "FR",
        "listing": {
            "title": "Long French Title",
            "bullets": [f"Bullet {index}" for index in range(1, 6)],
            "description": "First paragraph.\n\nSecond paragraph.",
            "backend_search_terms": "search words",
        },
    }
    if options is not None:
        report["title_options_2026"] = options
    return report


def make_template(path: Path, asin: str, include_item_name=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Modele"
    headers = [
        "SKU", "Title", "Nom de l'article", "Point fort de l’article",
        "Type d'identifiant de produit", "Identifiant du produit", "Description du produit",
        "Puce", "Puce", "Puce", "Puce", "Puce", "Mots-clés de recherche",
    ]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column).value = header
    sheet.cell(7, 1).value = "sku-1"
    sheet.cell(7, 2).value = "Old title"
    sheet.cell(7, 3).value = "Old name" if include_item_name else None
    sheet.cell(7, 4).value = "Old highlight"
    sheet.cell(7, 5).value = "ASIN"
    sheet.cell(7, 6).value = asin
    for column in range(8, 13):
        sheet.column_dimensions[sheet.cell(4, column).column_letter].hidden = True
    workbook.save(path)


def add_current_template(workbook, title: str, asins: list[str], hidden_bullet=False):
    sheet = workbook.create_sheet(title)
    headers = [
        "Titolo", "Nome dell’articolo", "Caratteristiche principali articolo",
        "Tipo ID di prodotto", "ID prodotto", "Descrizione del prodotto",
        "Punto elenco", "Punto elenco", "Punto elenco", "Punto elenco", "Punto elenco", "Chiavi di ricerca",
    ]
    attributes = [
        "title[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#1.value",
        "item_name[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#1.value",
        "title_differentiation[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#1.value",
        "amzn1.volt.ca.product_id_type", "amzn1.volt.ca.product_id_value",
        "product_description[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#1.value",
        *[f"bullet_point[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#{index}.value" for index in range(1, 6)],
        "generic_keyword[marketplace_id=APJ6JRA9NG5V4][language_tag=it_IT]#1.value",
    ]
    for column, (header, attribute) in enumerate(zip(headers, attributes), 1):
        sheet.cell(4, column).value = header
        sheet.cell(5, column).value = attribute
    if hidden_bullet:
        sheet.column_dimensions["H"].hidden = True
    for row, asin in enumerate(asins, 7):
        sheet.cell(row, 4).value = "ASIN"
        sheet.cell(row, 5).value = asin
    return sheet


def make_parent_child_template(path: Path, use_parent_asin=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Modele"
    headers = [
        "SKU", "Parent ASIN" if use_parent_asin else "Parent SKU", "Title", "Nom de l'article",
        "Type d'identifiant de produit", "Identifiant du produit", "Description du produit",
        "Puce", "Puce", "Puce", "Puce", "Puce", "Mots-clés de recherche",
    ]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column).value = header
    for row in range(1, 7):
        sheet.cell(row, 15).value = f"protected-{row}"
    records = [
        (7, "parent-sku", "", "B0PARENT01"),
        (8, "child-sku-1", "B0PARENT01" if use_parent_asin else "parent-sku", "B0CHILD001"),
        (9, "unrelated", "", "B0OTHER001"),
        (10, "child-sku-2", "parent-sku", "B0CHILD002"),
    ]
    for row, sku, parent, asin in records:
        sheet.cell(row, 1).value = sku
        sheet.cell(row, 2).value = parent
        sheet.cell(row, 5).value = "ASIN"
        sheet.cell(row, 6).value = asin
    workbook.save(path)


class ListingSyncTests(unittest.TestCase):
    def test_ireland_template_is_identified_by_marketplace_id_filename_and_language(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "爱尔兰-分类商品报告.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Template"
            headers = [
                "Title", "Item Name", "Item Highlight", "Product ID Type",
                "Product ID", "Product Description", *(["Bullet Point"] * 5), "Search Terms",
            ]
            attributes = [
                "title[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#1.value",
                "item_name[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#1.value",
                "title_differentiation[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#1.value",
                "amzn1.volt.ca.product_id_type", "amzn1.volt.ca.product_id_value",
                "product_description[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#1.value",
                *[f"bullet_point[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#{index}.value" for index in range(1, 6)],
                "generic_keyword[marketplace_id=A28R8C7NBKEWEA][language_tag=en_IE]#1.value",
            ]
            for column, (header, attribute) in enumerate(zip(headers, attributes), 1):
                sheet.cell(4, column).value = header
                sheet.cell(5, column).value = attribute
            sheet.cell(7, 4).value = "ASIN"
            sheet.cell(7, 5).value = "B0TESTIE01"
            workbook.save(source)
            result = SYNC.classify_workbook(source)
            self.assertEqual(result["status"], "ready")
            self.assertEqual(result["marketplace"], "IE")
            self.assertIn("marketplace_id", result["station_evidence"])
            self.assertIn("filename", result["station_evidence"])

    def test_localized_header_matrix(self):
        cases = [
            ["Titel", "Artikelname", "Art der Produkt-ID", "Produkt-ID", "Beschreibung des Produkts", "Aufzählungspunkt", "Suchbegriffe"],
            ["Title", "Item Name", "Product ID Type", "Product ID", "Product Description", "Bullet Point", "Search Terms"],
            ["Titre", "Nom de l'article", "Type d'identifiant de produit", "Identifiant du produit", "Description du produit", "Puce", "Mots-clés de recherche"],
            ["Titolo", "Nome articolo", "Tipo di ID prodotto", "ID prodotto", "Descrizione del prodotto", "Punto elenco", "Termini di ricerca"],
            ["Título", "Nombre del producto", "Tipo de ID de producto", "ID de producto", "Descripción del producto", "Viñeta", "Términos de búsqueda"],
            ["Título do produto", "Nome do item", "Tipo de ID do produto", "ID do produto", "Descrição do produto", "Marcador", "Termos de pesquisa"],
            ["Titel van het product", "Artikelnaam", "Typ produkt-ID", "Product-id", "Productbeschrijving", "Opsommingsteken", "Zoektermen"],
            ["Produktnamn", "Produktnamn", "Typ produkt-ID", "Produkt-ID", "Produktbeskrivning", "Punkt", "Söktermer"],
            ["Tytuł", "Nazwa produktu", "Typ identyfikatora produktu", "Identyfikator produktu", "Opis produktu", "Punkt wypunktowania", "Wyszukiwane hasła"],
            ["商品名", "商品名", "商品IDのタイプ", "商品ID", "商品の説明", "箇条書き", "検索キーワード"],
            ["Ürün Başlığı", "Ürün adı", "Ürün Kimliği Türü", "Ürün Kimliği", "Ürün açıklaması", "Madde işareti", "Arama terimleri"],
            ["عنوان المنتج", "اسم المنتج", "نوع معرف المنتج", "معرف المنتج", "وصف المنتج", "نقطة تعداد", "مصطلحات البحث"],
        ]
        for headers in cases:
            workbook = Workbook()
            sheet = workbook.active
            for column, value in enumerate(headers, 1):
                sheet.cell(4, column).value = value
            fields = SYNC.header_map(sheet)
            self.assertTrue({"title", "item_name", "product_id_type", "product_id", "description", "bullet", "search_terms"}.issubset(fields), headers)

    def test_attribute_ids_are_language_independent_and_marketplace_aliases_normalize(self):
        language_tags = ["de_DE", "en_US", "fr_FR", "it_IT", "es_ES", "pt_BR", "nl_NL", "sv_SE", "pl_PL", "ja_JP", "tr_TR", "ar_AE"]
        attributes = ["title", "item_name", "title_differentiation", "product_id_type", "product_id_value", "product_description", "bullet_point", "generic_keyword"]
        for language in language_tags:
            workbook = Workbook()
            sheet = workbook.active
            for column, attribute in enumerate(attributes, 1):
                sheet.cell(4, column).value = f"Localized field {column}"
                sheet.cell(5, column).value = f"{attribute}[marketplace_id=APJ6JRA9NG5V4][language_tag={language}]#1.value"
            fields = SYNC.header_map(sheet)
            self.assertTrue({"title", "item_name", "item_highlight", "product_id_type", "product_id", "description", "bullet", "search_terms"}.issubset(fields), language)
        for marketplace, aliases in SYNC.MARKETPLACE_ALIASES.items():
            for alias in aliases:
                self.assertEqual(SYNC.normalize_marketplace(alias), marketplace)

    def test_french_hidden_bullets_empty_fields_and_short_title(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "cloud" / "AmazonFrance.xlsx"
            source.parent.mkdir()
            output_dir = root / "project" / "outputs"
            asin = "B0TEST0001"
            make_template(source, asin)
            report_path = root / f"listing-optimization-report-{asin}.json"
            report_path.write_text(json.dumps(make_report(asin, [{"option": 1, "title": "Short Title", "item_highlights": "New highlight"}]), ensure_ascii=False), encoding="utf-8")
            workbook = SYNC.classify_workbook(source)
            self.assertEqual(workbook["status"], "ready")
            self.assertEqual(workbook["marketplace"], "FR")
            self.assertEqual(workbook["fields"]["bullet"], [8, 9, 10, 11, 12])
            reports = SYNC.load_reports([report_path], {asin: 1})
            plan = SYNC.build_plan(reports, [workbook], {})
            self.assertEqual(plan["status"], "ready")
            self.assertIn("空白，将填充", {item["source_state"] for item in plan["actions"][0]["planned_updates"]})
            result = SYNC.apply_plan(plan, output_dir)[0]
            output = Path(result["output"])
            updated = load_workbook(output, data_only=False)["Modele"]
            self.assertEqual(updated.cell(7, 2).value, "Short Title")
            self.assertEqual(updated.cell(7, 3).value, "Short Title")
            self.assertEqual(updated.cell(7, 4).value, "New highlight")
            self.assertEqual([updated.cell(7, column).value for column in range(8, 13)], [f"Bullet {index}" for index in range(1, 6)])
            self.assertEqual(result["written_cells"], 10)
            with zipfile.ZipFile(output) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            row = xml[xml.index('<row r="7"'):xml.index("</row>", xml.index('<row r="7"'))]
            self.assertLess(row.index('r="H7"'), row.index('r="M7"'))

    def test_long_title_keeps_existing_highlight_and_output_isolated(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "cloud" / "AmazonFrance.xlsx"
            source.parent.mkdir()
            asin = "B0TEST0002"
            make_template(source, asin, include_item_name=True)
            report_path = root / f"listing-optimization-report-{asin}.json"
            report_path.write_text(json.dumps(make_report(asin), ensure_ascii=False), encoding="utf-8")
            workbook = SYNC.classify_workbook(source)
            plan = SYNC.build_plan(SYNC.load_reports([report_path], {}), [workbook], {})
            with self.assertRaises(ValueError):
                SYNC.apply_plan(plan, source.parent)
            output_dir = root / "project" / "outputs"
            output = Path(SYNC.apply_plan(plan, output_dir)[0]["output"])
            updated = load_workbook(output, data_only=False)["Modele"]
            self.assertEqual(updated.cell(7, 2).value, "Long French Title")
            self.assertEqual(updated.cell(7, 4).value, "Old highlight")
            second = Path(SYNC.apply_plan(plan, output_dir)[0]["output"])
            self.assertTrue(second.name.endswith("_v2.xlsx"))

    def test_short_title_policy_defaults_to_option_one_and_can_require_selection(self):
        report = make_report("B0TEST0003", [{"option": 1, "title": "A", "item_highlights": "A"}, {"option": 2, "title": "B", "item_highlights": "B"}])
        value, error = SYNC.parse_title_options(report, {})
        self.assertEqual(value["title"], "A")
        self.assertIsNone(error)
        value, error = SYNC.parse_title_options(report, {}, "require-selection")
        self.assertIsNone(value)
        self.assertIn("需用户选择", error)

    def test_current_attribute_template_resolves_unique_target_sheet_and_hidden_columns(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "cloud" / "AmazonItalia.xlsx"
            source.parent.mkdir()
            workbook = Workbook()
            workbook.remove(workbook.active)
            add_current_template(workbook, "Partial", ["B0TESTIT01"])
            add_current_template(workbook, "Modello", ["B0TESTIT01", "B0TESTIT02"], hidden_bullet=True)
            workbook.save(source)
            reports = []
            for asin in ["B0TESTIT01", "B0TESTIT02"]:
                report = make_report(asin, [{"option": 1, "title": f"Short {asin}", "item_highlights": f"Highlight {asin}"}])
                report["marketplace"] = "Amazon IT"
                path = root / f"listing-optimization-report-{asin}.json"
                path.write_text(json.dumps(report), encoding="utf-8")
                reports.append(path)
            plan = SYNC.build_plan(SYNC.load_reports(reports, {}), [SYNC.classify_workbook(source)], {})
            self.assertEqual(plan["status"], "ready")
            self.assertEqual({action["sheet"] for action in plan["actions"]}, {"Modello"})
            self.assertIn("attribute_id", {item["source"] for item in plan["actions"][0]["field_metadata"]["item_name"]})
            self.assertTrue(any(item["hidden"] for item in plan["actions"][0]["planned_updates"] if item["field"] == "bullet"))
            output = Path(SYNC.apply_plan(plan, root / "project" / "outputs")[0]["output"])
            updated = load_workbook(output, data_only=False)["Modello"]
            self.assertEqual(updated.cell(7, 1).value, "Short B0TESTIT01")
            self.assertEqual(updated.cell(7, 3).value, "Highlight B0TESTIT01")
            self.assertEqual(updated.cell(7, 8).value, "Bullet 2")
            self.assertEqual(updated.cell(7, 12).value, "search words")
            self.assertEqual(updated.cell(7, 8).fill.fgColor.rgb, "FFC6EFCE")

    def test_equally_complete_template_sheets_remain_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "AmazonItalia.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            add_current_template(workbook, "Modello A", ["B0TESTIT03"])
            add_current_template(workbook, "Modello B", ["B0TESTIT03"])
            workbook.save(source)
            report = make_report("B0TESTIT03")
            report["marketplace"] = "Italia"
            report_path = Path(temp) / "listing-optimization-report-B0TESTIT03.json"
            report_path.write_text(json.dumps(report), encoding="utf-8")
            plan = SYNC.build_plan(SYNC.load_reports([report_path], {}), [SYNC.classify_workbook(source)], {})
            self.assertEqual(plan["status"], "review_required")
            self.assertIn("相同字段完整度", plan["actions"][0]["blockers"][0])

    def test_apply_keeps_modified_child_and_parent_sku_row_only(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "cloud" / "AmazonFrance.xlsx"
            source.parent.mkdir()
            make_parent_child_template(source)
            report_path = root / "listing-optimization-report-B0CHILD001.json"
            report_path.write_text(json.dumps(make_report("B0CHILD001")), encoding="utf-8")
            plan = SYNC.build_plan(SYNC.load_reports([report_path], {}), [SYNC.classify_workbook(source)], {})

            result = SYNC.apply_plan(plan, root / "outputs")[0]
            updated = load_workbook(result["output"], data_only=False)["Modele"]
            self.assertEqual([updated.cell(row, 15).value for row in range(1, 7)], [f"protected-{row}" for row in range(1, 7)])
            self.assertEqual(updated.cell(7, 6).value, "B0PARENT01")
            self.assertEqual(updated.cell(8, 6).value, "B0CHILD001")
            self.assertEqual(updated.max_row, 8)
            with zipfile.ZipFile(result["output"]) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertNotIn('<row r="9"', xml)
            self.assertNotIn('<row r="10"', xml)
            self.assertEqual(result["retained_parent_rows"], 1)
            self.assertEqual(result["deleted_data_rows"], 2)

    def test_apply_resolves_parent_asin_and_keeps_all_modified_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "cloud" / "AmazonFrance.xlsx"
            source.parent.mkdir()
            make_parent_child_template(source, use_parent_asin=True)
            reports = []
            for asin in ["B0CHILD001", "B0CHILD002"]:
                path = root / f"listing-optimization-report-{asin}.json"
                path.write_text(json.dumps(make_report(asin)), encoding="utf-8")
                reports.append(path)
            plan = SYNC.build_plan(SYNC.load_reports(reports, {}), [SYNC.classify_workbook(source)], {})

            result = SYNC.apply_plan(plan, root / "outputs")[0]
            updated = load_workbook(result["output"], data_only=False)["Modele"]
            self.assertEqual(updated.cell(7, 6).value, "B0PARENT01")
            self.assertEqual(updated.cell(8, 6).value, "B0CHILD001")
            self.assertEqual(updated.cell(9, 6).value, "B0CHILD002")
            self.assertEqual(updated.cell(9, 3).value, "Long French Title")
            self.assertEqual(updated.cell(9, 3).fill.fgColor.rgb, "FFC6EFCE")
            self.assertEqual(updated.max_row, 9)
            with zipfile.ZipFile(result["output"]) as archive:
                xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertIn('<row r="9"', xml)
            self.assertNotIn('<row r="10"', xml)
            self.assertEqual(result["retained_parent_rows"], 1)
            self.assertEqual(result["deleted_data_rows"], 1)


if __name__ == "__main__":
    unittest.main()
